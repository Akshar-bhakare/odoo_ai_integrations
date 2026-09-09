import os
import re
import hashlib
import sqlite3
import json
import secrets
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from datetime import datetime, date, timedelta
from pathlib import Path

import pandas as pd

# Excel styling (for report formatting)
from openpyxl.styles import Font

# PDF
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet


# =========================
# HARD RULES / PATHS
# =========================
SCHEME_START_DATE = date(2026, 2, 1)
REPEAT_WINDOW_DAYS = 90

SALARY_FILE = r"C:\Users\SALES INCENTIVES\SALESPERSON SALARIES.xlsx"
COMMISSIONS_FILE = r"C:\Users\SALES INCENTIVES\COMMISSIONS.xlsx"
CREDIT_NOTES_FILE = r"C:\\Users\\SALES INCENTIVES\\CREDIT_NOTES.xlsx"

# Local database + logo paths (safe defaults; app also checks same folder as EXE/PY)
APP_DATA_DIR = Path(os.environ.get("SUNLECTRIC_INCENTIVE_DATA", r"C:\\Users\\SALES INCENTIVES"))
LOCAL_DB_FILE = str(APP_DATA_DIR / "SUNLECTRIC_INCENTIVES_LOCAL.db")
LOGO_CANDIDATES = [
    Path(__file__).with_name("Sunlectric_logo.png") if "__file__" in globals() else Path("Sunlectric_logo.png"),
    Path(__file__).with_name("logo.png") if "__file__" in globals() else Path("logo.png"),
    Path(r"C:\\Users\\ankit\\Desktop\\Sunlectric_logo.png"),
    Path(r"C:\\Users\\ankit\\Desktop\\Sunlectric logo.png"),
]

INVOICE_PREFIX = "SL/FY"
INVOICE_PREFIXES = ("SL/FY", "SL/")  # supports Odoo invoice formats such as SL/FY... and SL/25-26/...
BILL_PREFIX = "BILL/"
CREDIT_NOTE_PREFIX = "RMHD/"  # ✅ FIX: was used but missing in your code

REQ_COLS = [
    "Date",
    "Number",
    "Account",
    "Partner",
    "Label",
    "Debit",
    "Credit",
    "Sales Order Lines/Salesperson",
]

SALES_KEYWORDS = ("sales income",)
COGS_KEYWORDS = ("cost of goods sold",)

# Incentive scheme
MIN_NEW_CUST_COUNT = 1
NEW_CUST_BILL_MIN = 100000.0
NEW_CUST_MIN_MARGIN_PCT = 3.5  # %
NEW_CUST_BONUS = 1500.0
REPEAT_BONUS = 1500.0  # paid ONCE within 90 days per qualifying new customer


# =========================
# Helpers
# =========================
def safe_str(x):
    return "" if pd.isna(x) else str(x).strip()


def is_invoice_number_value(x) -> bool:
    """True for customer invoices. Keeps credit notes and vendor bills out.
    Odoo exports may use SL/FY... or SL/25-26/... depending on configuration.
    """
    s = safe_str(x)
    if not s:
        return False
    if s.startswith(CREDIT_NOTE_PREFIX) or s.startswith(BILL_PREFIX):
        return False
    return any(s.startswith(pfx) for pfx in INVOICE_PREFIXES)


def fmt_money(x):
    try:
        return f"{float(x):,.2f}"
    except Exception:
        return ""


def to_num(x):
    if pd.isna(x) or x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", "")
    try:
        return float(s)
    except Exception:
        return 0.0


def parse_date_entry(s: str):
    s = (s or "").strip()
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except Exception:
        return None


def month_start_end(year: int, month: int):
    start = date(year, month, 1)
    if month == 12:
        end = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        end = date(year, month + 1, 1) - timedelta(days=1)
    return start, end


def quarter_start_end(year: int, q: int):
    m1 = 1 + (q - 1) * 3
    start = date(year, m1, 1)
    end = month_start_end(year, m1 + 2)[1]
    return start, end


def ym_from_date(d: date):
    return f"{d.year:04d}-{d.month:02d}"


def ensure_parent_dir(filepath: str):
    Path(filepath).parent.mkdir(parents=True, exist_ok=True)


def make_row_id(parts: list[str]) -> str:
    s = "||".join([p or "" for p in parts])
    return hashlib.sha1(s.encode("utf-8")).hexdigest()[:16]


def safe_excel_save(df: pd.DataFrame, path: str):
    ensure_parent_dir(path)
    try:
        df.to_excel(path, index=False)
        return path, None
    except PermissionError:
        p = Path(path)
        alt = str(p.with_name(p.stem + f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}" + p.suffix))
        df.to_excel(alt, index=False)
        warn = (
            f"Permission denied for:\n{path}\n\nSaved to:\n{alt}\n\n"
            "Close Excel if the file was open, then try saving again."
        )
        return alt, warn


def excel_format_2dp(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"


def pdf_cell(v):
    if v is None:
        return ""
    try:
        if isinstance(v, float) and pd.isna(v):
            return ""
    except Exception:
        pass

    if isinstance(v, (int, float)):
        return f"{float(v):.2f}"

    s = str(v)
    try:
        x = float(s.replace(",", ""))
        return f"{x:.2f}"
    except Exception:
        return s


def safe_pct(numer: float, denom: float) -> float:
    denom = float(denom)
    if abs(denom) < 1e-9:
        return 0.0
    return float(numer) / denom * 100.0


# =========================
# Main App
# =========================
class OdooSalesMarginApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Sales Incentives System (Odoo + Charges + Salaries + Commissions + Credit Notes)")
        self.root.geometry("1760x980")

        # Data
        self.df_raw = None
        self.invoice_summary = None
        self.bill_lines = None

        # Charges file (invoice wise)
        self.charges_map = {}  # invoice_no -> {"transport": x, "loading": y}

        # Salaries effective-dated:
        self.salary_steps = {}
        self.allowed_salespeople = []

        # Commissions DB
        self.commission_db = pd.DataFrame(columns=[
            "RowID", "Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "Salesperson"
        ])

        # Credit Notes DB
        self.creditnote_db = pd.DataFrame(columns=[
            "RowID", "Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "Salesperson"
        ])

        # Global customer first invoice (across all salespeople)
        self.global_customer_first = {}   # customer -> first_date (date)
        self.qual_new_customers_global = set()  # customers meeting 1,00,000 + 3.5% in first month
        self.qual_new_customer_first_dt = {}    # customer -> first_date

        # UI vars
        self.file_path_var = tk.StringVar(value="")
        self.charges_path_var = tk.StringVar(value="(not loaded)")
        self.salesperson_var = tk.StringVar(value="")

        self.filter_mode_var = tk.StringVar(value="All")
        self.month_choice_var = tk.StringVar(value="")
        self.quarter_choice_var = tk.StringVar(value="")
        self.custom_from_var = tk.StringVar(value="")
        self.custom_to_var = tk.StringVar(value="")

        self.use_net_for_incentive_var = tk.BooleanVar(value=True)

        # Product filter
        self.product_filter_var = tk.StringVar(value="")
        self.product_filter_active = ""   # confirmed/applied filter string
        self._ac_win = None               # autocomplete popup window

        # Customer-wise view vars
        self.customer_search_var = tk.StringVar(value="")
        self.customer_choice_var = tk.StringVar(value="")
        self.customer_all_only_var = tk.BooleanVar(value=True)

        self.current_range = (None, None)

        self._build_ui()

        # Ensure files exist and load
        self.ensure_salary_file_exists()
        self.ensure_commissions_file_exists()
        self.ensure_creditnotes_file_exists()

        self.load_salaries()
        self.load_commissions_db()
        self.load_creditnotes_db()

        # Run a small integrity check at startup (no data yet)
        self.system_integrity_check(show_popup=False)

    # =========================
    # System Integrity Check
    # =========================
    def system_integrity_check(self, show_popup=True):
        issues = []

        # Paths
        for pth, nm in [(SALARY_FILE, "Salaries"), (COMMISSIONS_FILE, "Commissions"), (CREDIT_NOTES_FILE, "Credit Notes")]:
            try:
                ensure_parent_dir(pth)
            except Exception as e:
                issues.append(f"{nm}: cannot create parent folder for path:\n{pth}\n{e}")

        # Prefix sanity
        if not INVOICE_PREFIX or not isinstance(INVOICE_PREFIX, str):
            issues.append("INVOICE_PREFIX is invalid.")
        if not BILL_PREFIX or not isinstance(BILL_PREFIX, str):
            issues.append("BILL_PREFIX is invalid.")
        if not CREDIT_NOTE_PREFIX or not isinstance(CREDIT_NOTE_PREFIX, str):
            issues.append("CREDIT_NOTE_PREFIX is invalid (required for RMHD/... credit notes).")

        # Scheme sanity
        if NEW_CUST_BILL_MIN <= 0:
            issues.append("NEW_CUST_BILL_MIN must be > 0.")
        if NEW_CUST_MIN_MARGIN_PCT < 0:
            issues.append("NEW_CUST_MIN_MARGIN_PCT must be >= 0.")
        if MIN_NEW_CUST_COUNT < 0:
            issues.append("MIN_NEW_CUST_COUNT must be >= 0.")
        if REPEAT_WINDOW_DAYS <= 0:
            issues.append("REPEAT_WINDOW_DAYS must be > 0.")

        # If data loaded, check required columns & types
        if self.df_raw is not None:
            missing = [c for c in REQ_COLS if c not in self.df_raw.columns]
            if missing:
                issues.append("Odoo file missing required columns: " + ", ".join(missing))

        msg = "✅ System integrity check OK.\n\n(No critical issues found.)"
        if issues:
            msg = "⚠️ System integrity check found issues:\n\n- " + "\n- ".join(issues)

        if show_popup:
            if issues:
                messagebox.showwarning("Integrity Check", msg)
            else:
                messagebox.showinfo("Integrity Check", msg)
        return issues

    # =========================
    # Ensure base files exist
    # =========================
    def ensure_salary_file_exists(self):
        ensure_parent_dir(SALARY_FILE)
        if not os.path.exists(SALARY_FILE):
            df = pd.DataFrame(columns=["Salesperson", "EffectiveFrom", "MonthlySalary"])
            safe_excel_save(df, SALARY_FILE)

    def ensure_commissions_file_exists(self):
        ensure_parent_dir(COMMISSIONS_FILE)
        if not os.path.exists(COMMISSIONS_FILE):
            df = pd.DataFrame(columns=["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"])
            safe_excel_save(df, COMMISSIONS_FILE)

    def ensure_creditnotes_file_exists(self):
        ensure_parent_dir(CREDIT_NOTES_FILE)
        if not os.path.exists(CREDIT_NOTES_FILE):
            df = pd.DataFrame(columns=["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"])
            safe_excel_save(df, CREDIT_NOTES_FILE)

    # =========================
    # Salaries (Effective-dated)
    # =========================
    def load_salaries(self):
        if not os.path.exists(SALARY_FILE):
            self.ensure_salary_file_exists()

        df = pd.read_excel(SALARY_FILE)
        df.columns = [str(c).strip() for c in df.columns]

        if "Salesperson" not in df.columns:
            if len(df.columns) >= 1:
                df = df.rename(columns={df.columns[0]: "Salesperson"})
            else:
                df["Salesperson"] = ""
        if "EffectiveFrom" not in df.columns:
            df["EffectiveFrom"] = "1900-01-01"
        if "MonthlySalary" not in df.columns:
            sal_cand = None
            for c in df.columns:
                if "salary" in c.lower():
                    sal_cand = c
                    break
            if sal_cand:
                df = df.rename(columns={sal_cand: "MonthlySalary"})
            else:
                df["MonthlySalary"] = 0.0

        df = df[["Salesperson", "EffectiveFrom", "MonthlySalary"]].copy()
        df["Salesperson"] = df["Salesperson"].apply(safe_str)
        df["MonthlySalary"] = df["MonthlySalary"].apply(to_num)
        df["EffectiveFrom"] = pd.to_datetime(df["EffectiveFrom"], errors="coerce")
        df.loc[df["EffectiveFrom"].isna(), "EffectiveFrom"] = pd.to_datetime("1900-01-01")

        df = df[(df["Salesperson"] != "") & (df["MonthlySalary"] > 0)]
        if df.empty:
            self.allowed_salespeople = []
            self.salary_steps = {}
            self._update_sp_listbox([])
            messagebox.showwarning(
                "Salaries Required",
                f"Please fill salaries file with columns:\nSalesperson, EffectiveFrom, MonthlySalary\n\n{SALARY_FILE}"
            )
            return

        steps = {}
        for sp, g in df.groupby("Salesperson"):
            g = g.sort_values("EffectiveFrom")
            steps[sp] = [(d.date(), float(s)) for d, s in zip(g["EffectiveFrom"], g["MonthlySalary"])]
        self.salary_steps = steps
        self.allowed_salespeople = sorted(list(steps.keys()))

        self._update_sp_listbox(self.allowed_salespeople)

    def salary_for_month(self, sp: str, yearmonth: str) -> float:
        if sp not in self.salary_steps:
            return 0.0
        try:
            y = int(yearmonth[:4]); m = int(yearmonth[5:7])
            target = date(y, m, 1)
        except Exception:
            return 0.0

        chosen = 0.0
        for eff, sal in self.salary_steps[sp]:
            if eff <= target:
                chosen = sal
            else:
                break
        return float(chosen)

    # =========================
    # Commissions DB (Load/Save)
    # =========================
    def load_commissions_db(self):
        self.ensure_commissions_file_exists()
        df = pd.read_excel(COMMISSIONS_FILE)
        df.columns = [str(c).strip() for c in df.columns]
        for c in ["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]:
            if c not in df.columns:
                df[c] = ""

        df["RowID"] = df["RowID"].apply(safe_str)
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        df["YearMonth"] = df["YearMonth"].apply(safe_str)
        df["Amount"] = df["Amount"].apply(to_num)
        df["Salesperson"] = df["Salesperson"].apply(safe_str)

        mask = (df["YearMonth"].astype(str).str.strip() == "") & (~df["Date"].isna())
        df.loc[mask, "YearMonth"] = df.loc[mask, "Date"].dt.strftime("%Y-%m")

        self.commission_db = df[["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]].copy()

    def save_commissions_db(self):
        out = self.commission_db.copy()
        out["Date"] = pd.to_datetime(out["Date"], errors="coerce").dt.date
        saved_path, warn = safe_excel_save(out, COMMISSIONS_FILE)
        if warn:
            messagebox.showwarning("Commissions Saved (Fallback)", warn)

    def commission_monthly_breakdown(self, sp: str) -> dict:
        if self.commission_db is None or self.commission_db.empty:
            return {}
        db = self.commission_db.copy()
        db["Salesperson"] = db["Salesperson"].apply(safe_str).str.strip()
        db = db[db["Salesperson"] == sp]
        if db.empty:
            return {}
        db["Date"] = pd.to_datetime(db["Date"], errors="coerce")
        db = db[~db["Date"].isna()]
        if db.empty:
            return {}

        start, end = self.current_range
        if start is not None:
            db = db[db["Date"].dt.date >= start]
        if end is not None:
            db = db[db["Date"].dt.date <= end]
        if db.empty:
            return {}

        db["YearMonth"] = db["Date"].dt.strftime("%Y-%m")
        g = db.groupby("YearMonth")["Amount"].sum()
        return {k: float(v) for k, v in g.items()}

    def commission_line_items_in_range(self, sp: str) -> pd.DataFrame:
        if self.commission_db is None or self.commission_db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID"])
        db = self.commission_db.copy()
        db["Salesperson"] = db["Salesperson"].apply(safe_str).str.strip()
        db = db[db["Salesperson"] == sp]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID"])
        db["Date"] = pd.to_datetime(db["Date"], errors="coerce")
        db = db[~db["Date"].isna()]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID"])

        start, end = self.current_range
        if start is not None:
            db = db[db["Date"].dt.date >= start]
        if end is not None:
            db = db[db["Date"].dt.date <= end]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID"])

        db["YearMonth"] = db["Date"].dt.strftime("%Y-%m")
        db = db.sort_values(["Date", "Number", "Partner"], na_position="last")
        return db[["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID"]].copy()

    # =========================
    # Credit Notes DB (Load/Save) + Edit/Open from GUI
    # =========================
    def load_creditnotes_db(self):
        self.ensure_creditnotes_file_exists()
        df = pd.read_excel(CREDIT_NOTES_FILE)
        df.columns = [str(c).strip() for c in df.columns]
        for c in ["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]:
            if c not in df.columns:
                df[c] = ""

        df["RowID"] = df["RowID"].apply(safe_str)
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        df["YearMonth"] = df["YearMonth"].apply(safe_str)
        df["Amount"] = df["Amount"].apply(to_num)
        df["Salesperson"] = df["Salesperson"].apply(safe_str)

        mask = (df["YearMonth"].astype(str).str.strip() == "") & (~df["Date"].isna())
        df.loc[mask, "YearMonth"] = df.loc[mask, "Date"].dt.strftime("%Y-%m")

        self.creditnote_db = df[["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]].copy()

    def save_creditnotes_db(self):
        out = self.creditnote_db.copy()
        out["Date"] = pd.to_datetime(out["Date"], errors="coerce").dt.date
        saved_path, warn = safe_excel_save(out, CREDIT_NOTES_FILE)
        if warn:
            messagebox.showwarning("Credit Notes Saved (Fallback)", warn)

    def open_creditnotes_file_external(self):
        try:
            self.ensure_creditnotes_file_exists()
            os.startfile(CREDIT_NOTES_FILE)
        except Exception as e:
            messagebox.showerror("Error", f"Could not open credit notes file:\n{e}")

    def open_creditnotes_editor(self):
        self.load_creditnotes_db()

        dlg = tk.Toplevel(self.root)
        dlg.title("Edit Credit Notes Assignments (RMHD/...)")
        dlg.geometry("1420x620")
        dlg.transient(self.root)
        dlg.grab_set()

        ttk.Label(
            dlg,
            text=f"Credit notes file: {CREDIT_NOTES_FILE}\n"
                 f"- You can correct wrong salesperson assignments\n"
                 f"- Assign unassigned rows\n"
                 f"- Delete rows (if you want)\n"
                 f"Note: These are line-items (for audit). Credit note impact in reports is computed from Odoo file, "
                 f"and salesperson mapping comes from your assignments here.",
            padding=10
        ).pack(anchor="w")

        frame = ttk.Frame(dlg, padding=10)
        frame.pack(fill="both", expand=True)

        cols = ["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]
        tree = ttk.Treeview(frame, columns=cols, show="headings", height=18)
        for c in cols:
            tree.heading(c, text=c)
            if c in ("Partner","Label","Account"):
                tree.column(c, width=240, anchor="w")
            elif c in ("Salesperson",):
                tree.column(c, width=160, anchor="w")
            elif c in ("RowID",):
                tree.column(c, width=140, anchor="w")
            else:
                tree.column(c, width=120, anchor="w" if c != "Amount" else "e")
        tree.pack(side="left", fill="both", expand=True)

        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        df = self.creditnote_db.copy()
        df = df.sort_values(["Date", "Number", "Partner"], na_position="last")
        for _, r in df.iterrows():
            dt = r.get("Date")
            dt_s = "" if pd.isna(dt) else pd.to_datetime(dt).strftime("%Y-%m-%d")
            tree.insert("", "end", values=(
                safe_str(r.get("RowID","")),
                dt_s,
                safe_str(r.get("YearMonth","")),
                fmt_money(float(r.get("Amount",0.0))),
                safe_str(r.get("Partner","")),
                safe_str(r.get("Label","")),
                safe_str(r.get("Number","")),
                safe_str(r.get("Account","")),
                safe_str(r.get("Salesperson","")),
            ))

        controls = ttk.Frame(dlg, padding=10)
        controls.pack(fill="x")

        ttk.Label(controls, text="Assign selected rows to:").pack(side="left")
        sp_var = tk.StringVar(value=self.allowed_salespeople[0] if self.allowed_salespeople else "")
        sp_combo = ttk.Combobox(controls, textvariable=sp_var, values=self.allowed_salespeople, state="readonly", width=28)
        sp_combo.pack(side="left", padx=10)

        def assign_selected():
            sp = sp_var.get().strip()
            if not sp:
                return
            for iid in tree.selection():
                vals = list(tree.item(iid, "values"))
                vals[-1] = sp
                tree.item(iid, values=vals)

        def clear_selected():
            for iid in tree.selection():
                vals = list(tree.item(iid, "values"))
                vals[-1] = ""
                tree.item(iid, values=vals)

        def delete_selected():
            for iid in tree.selection():
                tree.delete(iid)

        def save_and_close():
            rows = []
            for iid in tree.get_children():
                v = tree.item(iid, "values")
                rowid = safe_str(v[0])
                dt_s = safe_str(v[1])
                dt_parsed = pd.to_datetime(dt_s, errors="coerce")
                ymv = "" if pd.isna(dt_parsed) else dt_parsed.strftime("%Y-%m")
                amt = to_num(v[3])
                partner = safe_str(v[4])
                label = safe_str(v[5])
                number = safe_str(v[6])
                account = safe_str(v[7])
                sp = safe_str(v[8])

                rows.append({
                    "RowID": rowid,
                    "Date": dt_parsed,
                    "YearMonth": ymv,
                    "Amount": amt,
                    "Partner": partner,
                    "Label": label,
                    "Number": number,
                    "Account": account,
                    "Salesperson": sp
                })

            self.creditnote_db = pd.DataFrame(rows, columns=cols)
            self.save_creditnotes_db()
            self.load_creditnotes_db()
            dlg.destroy()
            self.refresh_view()
            messagebox.showinfo("Saved", "Credit notes assignments saved.")

        ttk.Button(controls, text="Assign Selected", command=assign_selected).pack(side="left")
        ttk.Button(controls, text="Clear Selected", command=clear_selected).pack(side="left", padx=8)
        ttk.Button(controls, text="Delete Selected", command=delete_selected).pack(side="left", padx=8)
        ttk.Button(controls, text="Save & Close", command=save_and_close).pack(side="right")

    # =========================
    # UI
    # =========================
    def _build_ui(self):
        # Company logo on GUI (safe: app continues even if image is missing)
        #try:
         #   self._logo_img_tk = None
          #  for _lp in LOGO_CANDIDATES:
           #     if _lp.exists():
            #        self._logo_img_tk = tk.PhotoImage(file=str(_lp))
             #       break
            #if# self._logo_img_tk is not None:
             #   logo_bar = ttk.Frame(self.root, padding=(10, 8, 10, 2))
              #  logo_bar.pack(fill="x")
               # ttk.Label(logo_bar, image=self._logo_img_tk).pack(anchor="center")
        #except Exception:
         #   pass

        top = ttk.Frame(self.root, padding=10)
        top.pack(fill="x")

        ttk.Button(top, text="Select Odoo Excel File", command=self.select_odoo_file).pack(side="left")
        ttk.Label(top, textvariable=self.file_path_var, width=48).pack(side="left", padx=10)

        ttk.Button(top, text="Load Charges File", command=self.select_charges_file).pack(side="left")
        ttk.Label(top, textvariable=self.charges_path_var, width=18).pack(side="left", padx=10)

        ttk.Button(top, text="Edit Salaries", command=self.open_salary_editor).pack(side="left", padx=(0, 10))
        ttk.Button(top, text="Open Salaries File", command=self.open_salary_file_external).pack(side="left", padx=(0, 10))

        ttk.Button(top, text="Edit Credit Notes", command=self.open_creditnotes_editor).pack(side="left", padx=(0, 10))
        ttk.Button(top, text="Open Credit Notes File", command=self.open_creditnotes_file_external).pack(side="left", padx=(0, 10))

        ttk.Button(top, text="Integrity Check", command=lambda: self.system_integrity_check(show_popup=True)).pack(side="left", padx=(0, 10))
        ttk.Button(top, text="Export (Excel / PDF)", command=self.export_report).pack(side="right")

        tools_row = ttk.Frame(self.root, padding=(10, 0, 10, 6))
        tools_row.pack(fill="x")
        ttk.Label(tools_row, text="Management:", font=("Segoe UI", 9, "bold")).pack(side="left", padx=(0, 8))
        ttk.Button(tools_row, text="Change Password", command=self.open_change_password).pack(side="left", padx=(0, 8))
        ttk.Button(tools_row, text="Incentive Rates / Scheme", command=self.open_scheme_settings).pack(side="left", padx=(0, 8))
        ttk.Button(tools_row, text="Manual Add/Deduct", command=self.open_manual_adjustments).pack(side="left", padx=(0, 8))
        ttk.Button(tools_row, text="Record Payment", command=self.open_payment_recorder).pack(side="left", padx=(0, 8))

        row2 = ttk.Frame(self.root, padding=(10, 0, 10, 0))
        row2.pack(fill="x")

        ttk.Label(row2, text="Salesperson:").pack(side="left")
        sp_frame = ttk.Frame(row2)
        sp_frame.pack(side="left", padx=(8, 4))
        self.sp_listbox = tk.Listbox(sp_frame, selectmode=tk.EXTENDED, height=4, width=26,
                                     exportselection=False, font=("Segoe UI", 9))
        sp_vsb = ttk.Scrollbar(sp_frame, orient="vertical", command=self.sp_listbox.yview)
        self.sp_listbox.configure(yscrollcommand=sp_vsb.set)
        self.sp_listbox.pack(side="left")
        sp_vsb.pack(side="right", fill="y")
        self.sp_listbox.bind("<<ListboxSelect>>", lambda e: self.refresh_view())
        # keep a dummy salesperson_combo attribute so old code paths don't crash
        self.salesperson_combo = self.sp_listbox

        ttk.Label(row2, text="Ctrl/Shift\n+click multi",
                  font=("Segoe UI", 7), foreground="grey").pack(side="left", padx=(2, 8))

        ttk.Button(row2, text="📊 Team Summary", command=self.open_team_summary).pack(side="left", padx=(0, 14))

        ttk.Checkbutton(
            row2,
            text="Use Net base for incentive (AdjGM - EmpExp - Commission)",
            variable=self.use_net_for_incentive_var,
            command=self.refresh_view,
        ).pack(side="left", padx=(0, 0))

        # Filters (date + product in one LabelFrame)
        filt = ttk.LabelFrame(self.root, text="Filters", padding=8)
        filt.pack(fill="x", padx=10, pady=(4, 8))

        # Date filter row
        filt_d = ttk.Frame(filt)
        filt_d.pack(fill="x")
        modes = ["All", "Month", "Quarter", "Custom"]
        for m in modes:
            ttk.Radiobutton(filt_d, text=m, value=m, variable=self.filter_mode_var,
                            command=self.on_filter_change).pack(side="left", padx=(0, 12))

        ttk.Label(filt_d, text="Month:").pack(side="left", padx=(10, 4))
        self.month_combo = ttk.Combobox(filt_d, textvariable=self.month_choice_var, state="disabled", width=12)
        self.month_combo.pack(side="left", padx=(0, 12))
        self.month_combo.bind("<<ComboboxSelected>>", lambda e: self.apply_filter())

        ttk.Label(filt_d, text="Quarter:").pack(side="left", padx=(10, 4))
        self.quarter_combo = ttk.Combobox(filt_d, textvariable=self.quarter_choice_var, state="disabled", width=10)
        self.quarter_combo.pack(side="left", padx=(0, 12))
        self.quarter_combo.bind("<<ComboboxSelected>>", lambda e: self.apply_filter())

        ttk.Label(filt_d, text="From (YYYY-MM-DD):").pack(side="left", padx=(10, 4))
        self.from_entry = ttk.Entry(filt_d, textvariable=self.custom_from_var, width=12, state="disabled")
        self.from_entry.pack(side="left", padx=(0, 10))
        ttk.Label(filt_d, text="To:").pack(side="left", padx=(0, 4))
        self.to_entry = ttk.Entry(filt_d, textvariable=self.custom_to_var, width=12, state="disabled")
        self.to_entry.pack(side="left", padx=(0, 10))
        ttk.Button(filt_d, text="Apply", command=self.apply_filter).pack(side="left")

        # Product filter row
        filt_p = ttk.Frame(filt)
        filt_p.pack(fill="x", pady=(6, 0))
        ttk.Label(filt_p, text="Product Filter:").pack(side="left")
        self.product_filter_entry = ttk.Entry(filt_p, textvariable=self.product_filter_var, width=55)
        self.product_filter_entry.pack(side="left", padx=(6, 4))
        self.product_filter_entry.bind("<KeyRelease>", self._on_product_key)
        self.product_filter_entry.bind("<Return>", lambda e: self._apply_typed_product_filter())
        self.product_filter_entry.bind("<FocusOut>",
                                       lambda e: self.root.after(250, self._hide_product_ac))
        ttk.Button(filt_p, text="🔍 Apply", command=self._apply_typed_product_filter).pack(side="left", padx=(0, 6))
        ttk.Button(filt_p, text="✕ Clear", command=self._clear_product_filter).pack(side="left", padx=(0, 10))
        self.product_active_lbl = ttk.Label(filt_p, text="", foreground="blue",
                                            font=("Segoe UI", 9, "italic"))
        self.product_active_lbl.pack(side="left")

        # Summary bar
        summary = ttk.Frame(self.root, padding=(10, 0, 10, 10))
        summary.pack(fill="x")

        self.lbl_sales = ttk.Label(summary, text="Sales: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_cogs = ttk.Label(summary, text="COGS: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_chg = ttk.Label(summary, text="Charges (T+L): 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_adjgm = ttk.Label(summary, text="Adjusted GM: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_exp = ttk.Label(summary, text="Employee Exp: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_comm = ttk.Label(summary, text="Commission: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_net = ttk.Label(summary, text="Net: 0.00", font=("Segoe UI", 11, "bold"))
        self.lbl_incent = ttk.Label(summary, text="Incentive (₹): 0.00", font=("Segoe UI", 11, "bold"))

        for w in [self.lbl_sales, self.lbl_cogs, self.lbl_chg, self.lbl_adjgm, self.lbl_exp, self.lbl_comm, self.lbl_net, self.lbl_incent]:
            w.pack(side="left", padx=(0, 14))

        # Notebook
        self.nb = ttk.Notebook(self.root)
        self.nb.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        self.tab_txn = ttk.Frame(self.nb)
        self.tab_inc = ttk.Frame(self.nb)
        self.tab_comm = ttk.Frame(self.nb)
        self.tab_cn = ttk.Frame(self.nb)
        self.tab_cust = ttk.Frame(self.nb)  # ✅ NEW: customer-wise verification

        self.nb.add(self.tab_txn, text="Transactions")
        self.nb.add(self.tab_inc, text="Incentives")
        self.nb.add(self.tab_comm, text="Commission Details")
        self.nb.add(self.tab_cn, text="Credit Note Details")
        self.nb.add(self.tab_cust, text="Customer Wise (Verify)")

        # Line items toolbar (above Transactions tree)
        txn_topbar = ttk.Frame(self.tab_txn)
        txn_topbar.pack(fill="x", padx=10, pady=(5, 0))
        ttk.Button(txn_topbar, text="📋 Show Line Items for Selected Invoice",
                   command=self.show_line_items_popup).pack(side="left")

        self.tree_txn = self._make_tree(self.tab_txn, kind="txn")
        self.tree_inc = self._make_tree(self.tab_inc, kind="inc")
        self.tree_comm = self._make_tree(self.tab_comm, kind="comm")
        self.tree_cn = self._make_tree(self.tab_cn, kind="cn")
        self.tree_cust = self._make_tree(self.tab_cust, kind="cust")

        # ✅ Color/Font tags (GUI)
        # Global new customers: green + bold
        self.tree_txn.tag_configure("GLOBAL_NEW", foreground="green", font=("Segoe UI", 9, "bold"))
        # Qualified repeats: blue
        self.tree_txn.tag_configure("QUAL_REPEAT", foreground="blue")
        # Qualified new customers (in practice subset of GLOBAL_NEW, but kept separate)
        self.tree_txn.tag_configure("QUAL_NEW", font=("Segoe UI", 9, "bold"))

        # Customer wise tree tags too
        self.tree_cust.tag_configure("GLOBAL_NEW", foreground="green", font=("Segoe UI", 9, "bold"))
        self.tree_cust.tag_configure("QUAL_REPEAT", foreground="blue")
        self.tree_cust.tag_configure("QUAL_NEW", font=("Segoe UI", 9, "bold"))

        # Multi-SP section header tag
        for _t in [self.tree_txn, self.tree_inc, self.tree_comm, self.tree_cn]:
            _t.tag_configure("SP_HEADER", background="#DDEEFF", font=("Segoe UI", 9, "bold"))

        # Customer-wise controls
        self._build_customer_controls()

        bottom = ttk.LabelFrame(self.root, text="Incentive Scheme Calculations (Selected Range)", padding=10)
        bottom.pack(fill="x", padx=10, pady=(0, 10))
        self.incent_text = tk.Text(bottom, height=8, wrap="word")
        self.incent_text.pack(fill="both", expand=True)

        self.status_var = tk.StringVar(value="Ready. Load Odoo file to begin.")
        ttk.Label(self.root, textvariable=self.status_var, anchor="w", padding=8).pack(fill="x")

    def _build_customer_controls(self):
        top = ttk.Frame(self.tab_cust, padding=(10, 10, 10, 0))
        top.pack(fill="x")

        ttk.Label(top, text="Search Customer:").pack(side="left")
        e = ttk.Entry(top, textvariable=self.customer_search_var, width=40)
        e.pack(side="left", padx=(8, 12))
        e.bind("<KeyRelease>", lambda _e: self.populate_customer_dropdown())

        ttk.Button(top, text="Refresh List", command=self.populate_customer_dropdown).pack(side="left", padx=(0, 12))

        ttk.Label(top, text="Customer:").pack(side="left")
        self.customer_combo = ttk.Combobox(top, textvariable=self.customer_choice_var, state="readonly", width=55)
        self.customer_combo.pack(side="left", padx=(8, 12))
        self.customer_combo.bind("<<ComboboxSelected>>", lambda e2: self.refresh_customer_view())

        ttk.Checkbutton(
            top,
            text="Show All Customers (ignore search)",
            variable=self.customer_all_only_var,
            command=self.populate_customer_dropdown,
        ).pack(side="left", padx=(0, 12))

        ttk.Button(top, text="View Selected Customer", command=self.refresh_customer_view).pack(side="left")
        ttk.Button(top, text="📥 Export Customer Report (Excel)",
                   command=self.export_customer_report).pack(side="left", padx=(12, 0))

        ttk.Label(
            self.tab_cust,
            text="This tab shows ALL sales (all salespeople) customer-wise for verification, within selected date range.",
            padding=(10, 6, 10, 6)
        ).pack(anchor="w")

    def _make_tree(self, parent, kind="txn"):
        frame = ttk.Frame(parent, padding=10)
        frame.pack(fill="both", expand=True)

        if kind == "txn":
            cols = ["Type", "Date", "Number", "Partner/Info", "Sales", "COGS", "Transport", "Loading",
                    "Adj GM", "GM %", "Emp Exp", "Commission", "Net"]
        elif kind == "inc":
            cols = ["Month", "Salary Used", "Base", "Threshold", "Carry In", "Met?", "10%", "New", "Repeat", "Total", "Notice"]
        elif kind == "comm":
            cols = ["Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "RowID"]
        elif kind == "cust":
            cols = ["Type", "Date", "Number", "Salesperson", "Partner", "Sales", "COGS", "Adj GM", "GM %"]
        else:  # cn
            cols = ["Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "RowID", "Salesperson"]

        tree = ttk.Treeview(frame, columns=cols, show="headings", height=22)
        for c in cols:
            tree.heading(c, text=c)

            if kind == "txn":
                if c in ("Type",):
                    tree.column(c, width=110, anchor="w")
                elif c in ("Date",):
                    tree.column(c, width=110, anchor="w")
                elif c in ("Number",):
                    tree.column(c, width=180, anchor="w")
                elif c in ("Partner/Info",):
                    tree.column(c, width=340, anchor="w")
                elif c in ("GM %",):
                    tree.column(c, width=90, anchor="e")
                else:
                    tree.column(c, width=120, anchor="e")

            elif kind == "inc":
                if c in ("Month", "Met?", "Notice"):
                    tree.column(c, width=110, anchor="w")
                else:
                    tree.column(c, width=150, anchor="e")

            elif kind == "cust":
                if c in ("Type", "Date", "Number", "Salesperson"):
                    tree.column(c, width=140, anchor="w")
                elif c in ("Partner",):
                    tree.column(c, width=360, anchor="w")
                else:
                    tree.column(c, width=140, anchor="e")

            else:
                if c in ("Partner", "Label", "Account"):
                    tree.column(c, width=280, anchor="w")
                elif c == "RowID":
                    tree.column(c, width=170, anchor="w")
                elif c in ("Date", "YearMonth", "Number"):
                    tree.column(c, width=120, anchor="w")
                elif c == "Salesperson":
                    tree.column(c, width=160, anchor="w")
                else:
                    tree.column(c, width=120, anchor="e")

        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        hsb.pack(side="bottom", fill="x")
        return tree

    # =========================
    # File selection
    # =========================
    def select_odoo_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Odoo Excel Export",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if not file_path:
            return

        if not self.allowed_salespeople:
            messagebox.showerror("Salaries Required", f"Fill salaries file first:\n{SALARY_FILE}")
            return

        self.file_path_var.set(file_path)
        try:
            self.load_and_prepare_odoo(file_path)
            self.populate_months_quarters()
            self.on_filter_change()

            # assign new commissions + credit notes if any
            self.assign_new_commissions_if_any()
            self.assign_new_creditnotes_if_any()

            self.refresh_view()
            self.populate_customer_dropdown()

            self.status_var.set("Odoo file loaded. Charges optional. Commissions/Credit Notes loaded/assigned.")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to process Odoo file:\n\n{e}")
            self.status_var.set("Error loading Odoo file.")

    def select_charges_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Charges Excel (Invoice loading/transport)",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if not file_path:
            return
        try:
            self.load_charges_file(file_path)
            self.charges_path_var.set(Path(file_path).name)
            self.refresh_view()
            self.status_var.set("Charges file loaded and applied (missing invoices assumed 0).")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load charges file:\n\n{e}")

    # =========================
    # Charges file
    # =========================
    def load_charges_file(self, file_path):
        df = pd.read_excel(file_path)
        df.columns = [str(c).strip() for c in df.columns]

        inv_col = None
        for cand in ["Invoice No", "Invoice", "Invoice Number", "InvoiceNo", "Number"]:
            if cand in df.columns:
                inv_col = cand
                break
        if not inv_col:
            raise ValueError("Could not find Invoice/Number column in charges file.")

        t_col = next((c for c in df.columns if "transport" in c.lower()), None)
        l_col = next((c for c in df.columns if "loading" in c.lower()), None)

        charges_map = {}
        for _, r in df.iterrows():
            inv = safe_str(r.get(inv_col, ""))
            if not inv:
                continue
            transport = to_num(r.get(t_col, 0.0)) if t_col else 0.0
            loading = to_num(r.get(l_col, 0.0)) if l_col else 0.0
            charges_map[inv] = {"transport": transport, "loading": loading}

        self.charges_map = charges_map
        if self.invoice_summary is not None:
            self.apply_charges_to_invoices()
            self.rebuild_global_customer_maps()  # ✅ recompute qualification because AdjGM changes

    # =========================
    # Data prep
    # =========================
    def load_and_prepare_odoo(self, file_path):
        df = pd.read_excel(file_path)
        df.columns = [str(c).strip() for c in df.columns]

        missing = [c for c in REQ_COLS if c not in df.columns]
        if missing:
            raise ValueError("Missing required columns:\n- " + "\n- ".join(missing))

        df = df.copy()
        df["Number"] = df["Number"].astype(str).str.strip()
        df["Account"] = df["Account"].astype(str).str.strip()
        df["Partner"] = df["Partner"].astype(str).str.strip()
        df["Label"] = df["Label"].astype(str).fillna("").str.strip()
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        df["Debit"] = pd.to_numeric(df["Debit"], errors="coerce").fillna(0.0)
        df["Credit"] = pd.to_numeric(df["Credit"], errors="coerce").fillna(0.0)
        df["Salesperson"] = df["Sales Order Lines/Salesperson"].apply(safe_str)
        self.df_raw = df

        inv = df[df["Number"].apply(is_invoice_number_value)].copy()
        bill = df[df["Number"].str.startswith(BILL_PREFIX)].copy()

        inv_sp = (
            inv.groupby("Number")["Salesperson"]
            .apply(lambda s: next((x for x in s if x), ""))
            .rename("Salesperson")
        )

        inv["acct_l"] = inv["Account"].str.lower()
        is_sales_kw = inv["acct_l"].apply(lambda x: any(k in x for k in SALES_KEYWORDS))
        is_cogs_kw = inv["acct_l"].apply(lambda x: any(k in x for k in COGS_KEYWORDS))

        sales_by_inv = inv[is_sales_kw].groupby("Number")["Credit"].sum()
        cogs_by_inv = inv[is_cogs_kw].groupby("Number")["Debit"].sum()

        credit_fallback = inv.groupby("Number")["Credit"].sum()
        debit_fallback = inv.groupby("Number")["Debit"].sum()

        inv_date = inv.groupby("Number")["Date"].min()
        inv_partner = inv.groupby("Number")["Partner"].apply(lambda s: next(iter(s), ""))

        summary = pd.DataFrame({"Date": inv_date, "Number": inv_date.index, "Partner": inv_partner}).set_index("Number", drop=False)
        summary = summary.join(inv_sp, how="left")

        summary["Sales"] = sales_by_inv.reindex(summary.index).fillna(0.0)
        summary["COGS"] = cogs_by_inv.reindex(summary.index).fillna(0.0)

        need_sales_fb = summary["Sales"].abs() < 1e-9
        summary.loc[need_sales_fb, "Sales"] = credit_fallback.reindex(summary.index).fillna(0.0)[need_sales_fb]

        need_cogs_fb = summary["COGS"].abs() < 1e-9
        summary.loc[need_cogs_fb, "COGS"] = debit_fallback.reindex(summary.index).fillna(0.0)[need_cogs_fb]

        summary["Gross Margin"] = summary["Sales"] - summary["COGS"]
        summary["YearMonth"] = summary["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        summary["Customer"] = summary["Partner"].apply(safe_str)

        summary["Transport"] = 0.0
        summary["Loading"] = 0.0
        summary["Adjusted GM"] = summary["Gross Margin"]

        self.invoice_summary = summary.reset_index(drop=True)

        bill["ExpensePerson"] = bill["Partner"].apply(safe_str)
        bill["ExpenseAmount"] = (bill["Debit"] - bill["Credit"]).astype(float)
        bill["YearMonth"] = bill["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        self.bill_lines = bill

        if self.charges_map:
            self.apply_charges_to_invoices()

        # ✅ Build global firsts + qualification maps
        self.rebuild_global_customer_maps()

        self.salesperson_combo = self.sp_listbox  # keep alias valid
        self._update_sp_listbox(self.allowed_salespeople)

    def apply_charges_to_invoices(self):
        inv = self.invoice_summary.copy()
        transport = []
        loading = []
        for _, r in inv.iterrows():
            num = safe_str(r["Number"])
            ch = self.charges_map.get(num)
            if ch:
                transport.append(float(ch.get("transport", 0.0)))
                loading.append(float(ch.get("loading", 0.0)))
            else:
                transport.append(0.0)
                loading.append(0.0)
        inv["Transport"] = transport
        inv["Loading"] = loading
        inv["Adjusted GM"] = inv["Gross Margin"] - inv["Transport"] - inv["Loading"]
        self.invoice_summary = inv

    def rebuild_global_customer_maps(self):
        self.global_customer_first = {}
        self.qual_new_customers_global = set()
        self.qual_new_customer_first_dt = {}

        if self.invoice_summary is None or self.invoice_summary.empty:
            return

        inv_all = self.invoice_summary.copy()
        inv_all["DateOnly"] = inv_all["Date"].dt.date
        inv_all = inv_all[inv_all["Customer"].astype(str).str.strip() != ""].copy()

        if inv_all.empty:
            return

        firsts = inv_all.groupby("Customer")["DateOnly"].min().reset_index()
        self.global_customer_first = {r["Customer"]: r["DateOnly"] for _, r in firsts.iterrows()}

        # ✅ Qualification (Requirement #2): only if Sales >= 1,00,000 and GM% >= 3.5 in their first month
        for cust, first_dt in self.global_customer_first.items():
            if not first_dt or first_dt < SCHEME_START_DATE:
                continue
            first_m = ym_from_date(first_dt)
            inv_first = inv_all[(inv_all["Customer"] == cust) & (inv_all["YearMonth"] == first_m)]
            s_val = float(inv_first["Sales"].sum())
            gm_val = float(inv_first["Adjusted GM"].sum())
            gm_pct = safe_pct(gm_val, s_val)
            if (s_val >= NEW_CUST_BILL_MIN - 1e-9) and (gm_pct >= NEW_CUST_MIN_MARGIN_PCT - 1e-9):
                self.qual_new_customers_global.add(cust)
                self.qual_new_customer_first_dt[cust] = first_dt

    # =========================
    # Filters
    # =========================
    def on_filter_change(self):
        mode = self.filter_mode_var.get()
        self.month_combo.configure(state="readonly" if mode == "Month" else "disabled")
        self.quarter_combo.configure(state="readonly" if mode == "Quarter" else "disabled")
        self.from_entry.configure(state="normal" if mode == "Custom" else "disabled")
        self.to_entry.configure(state="normal" if mode == "Custom" else "disabled")
        self.apply_filter()

    def apply_filter(self):
        mode = self.filter_mode_var.get()
        start = None
        end = None

        if mode == "All":
            start, end = None, None
        elif mode == "Month":
            m = self.month_choice_var.get().strip()
            if re.match(r"^\d{4}-\d{2}$", m):
                y = int(m[:4]); mo = int(m[5:7])
                start, end = month_start_end(y, mo)
        elif mode == "Quarter":
            q = self.quarter_choice_var.get().strip()
            mm = re.match(r"^(\d{4})-Q([1-4])$", q)
            if mm:
                y = int(mm.group(1)); qq = int(mm.group(2))
                start, end = quarter_start_end(y, qq)
        elif mode == "Custom":
            d1 = parse_date_entry(self.custom_from_var.get())
            d2 = parse_date_entry(self.custom_to_var.get())
            if not d1 or not d2:
                messagebox.showwarning("Custom Dates", "Enter valid dates (YYYY-MM-DD) for From and To.")
                return
            if d2 < d1:
                messagebox.showwarning("Custom Dates", "'To' date must be >= 'From' date.")
                return
            start, end = d1, d2

        self.current_range = (start, end)
        self.refresh_view()
        self.refresh_customer_view()

    def _filter_by_range(self, df: pd.DataFrame, date_col="Date"):
        start, end = self.current_range
        if start is None and end is None:
            return df
        d = df[date_col].dt.date
        mask = True
        if start is not None:
            mask = (d >= start)
        if end is not None:
            mask = mask & (d <= end)
        return df[mask].copy()

    def populate_months_quarters(self):
        dates = pd.concat(
            [self.invoice_summary["Date"].dropna(), self.bill_lines["Date"].dropna()],
            ignore_index=True,
        )
        if dates.empty:
            self.month_combo["values"] = []
            self.quarter_combo["values"] = []
            return

        dmin = dates.min().date()
        dmax = dates.max().date()

        months = []
        cur = date(dmin.year, dmin.month, 1)
        endm = date(dmax.year, dmax.month, 1)
        while cur <= endm:
            months.append(f"{cur.year:04d}-{cur.month:02d}")
            cur = date(cur.year + (cur.month == 12), 1 if cur.month == 12 else cur.month + 1, 1)
        self.month_combo["values"] = months
        if months:
            self.month_choice_var.set(months[-1])

        quarters = sorted({f"{m[:4]}-Q{(int(m[5:7])-1)//3+1}" for m in months})
        self.quarter_combo["values"] = quarters
        if quarters:
            self.quarter_choice_var.set(quarters[-1])

    # =========================
    # Customer flags for styling / incentives
    # =========================
    def is_global_new_customer_month(self, cust: str, inv_date: date) -> bool:
        first_dt = self.global_customer_first.get(cust)
        if not first_dt:
            return False
        if first_dt < SCHEME_START_DATE:
            return False
        return (first_dt.year == inv_date.year and first_dt.month == inv_date.month)

    def is_qualified_new_customer_month(self, cust: str, inv_date: date) -> bool:
        if cust not in self.qual_new_customers_global:
            return False
        first_dt = self.qual_new_customer_first_dt.get(cust)
        if not first_dt:
            return False
        return (first_dt.year == inv_date.year and first_dt.month == inv_date.month)

    def is_qualified_repeat_invoice(self, cust: str, inv_date: date) -> bool:
        # ✅ Requirement #3: only repeats for QUALIFIED new customers count / highlight
        if cust not in self.qual_new_customers_global:
            return False
        first_dt = self.qual_new_customer_first_dt.get(cust)
        if not first_dt:
            return False
        if inv_date <= first_dt:
            return False
        if inv_date > (first_dt + timedelta(days=REPEAT_WINDOW_DAYS)):
            return False
        return True

    # =========================
    # Customer-wise tab helpers
    # =========================
    def populate_customer_dropdown(self):
        if self.invoice_summary is None or self.invoice_summary.empty:
            self.customer_combo["values"] = []
            self.customer_choice_var.set("")
            return

        customers = sorted(set(self.invoice_summary["Customer"].astype(str).map(str.strip).tolist()))
        customers = [c for c in customers if c]

        if not self.customer_all_only_var.get():
            q = self.customer_search_var.get().strip().lower()
            if q:
                customers = [c for c in customers if q in c.lower()]

        self.customer_combo["values"] = customers
        cur = self.customer_choice_var.get().strip()
        if cur not in customers:
            self.customer_choice_var.set(customers[0] if customers else "")

    def refresh_customer_view(self):
        if self.invoice_summary is None or self.invoice_summary.empty:
            return

        cust = self.customer_choice_var.get().strip()
        for iid in self.tree_cust.get_children():
            self.tree_cust.delete(iid)

        if not cust:
            return

        inv_all = self.invoice_summary.copy()
        inv_all = self._filter_by_range(inv_all, "Date")
        inv_c = inv_all[inv_all["Customer"].astype(str).str.strip() == cust].copy()
        inv_c = inv_c.sort_values(["Date", "Number"], na_position="last")

        for _, r in inv_c.iterrows():
            dt = r["Date"]
            inv_date = dt.date() if not pd.isna(dt) else None
            date_s = "" if pd.isna(dt) else dt.strftime("%Y-%m-%d")

            sales = float(r.get("Sales", 0.0))
            adjgm = float(r.get("Adjusted GM", 0.0))
            gm_pct = safe_pct(adjgm, sales)

            tags = ()
            if inv_date:
                if self.is_global_new_customer_month(cust, inv_date):
                    tags = ("GLOBAL_NEW",)
                if self.is_qualified_repeat_invoice(cust, inv_date):
                    tags = ("QUAL_REPEAT",)
                # qualified new month bold (global-new already bold green)
                if self.is_qualified_new_customer_month(cust, inv_date):
                    tags = ("GLOBAL_NEW", "QUAL_NEW")

            self.tree_cust.insert(
                "",
                "end",
                values=(
                    "Invoice",
                    date_s,
                    safe_str(r.get("Number", "")),
                    safe_str(r.get("Salesperson", "")),
                    safe_str(r.get("Partner", "")),
                    fmt_money(sales),
                    fmt_money(float(r.get("COGS", 0.0))),
                    fmt_money(adjgm),
                    f"{gm_pct:.2f}",
                ),
                tags=tags
            )

        # Total
        if not inv_c.empty:
            self.tree_cust.insert("", "end", values=("", "", "", "", "", "", "", "", ""))
            s_tot = float(inv_c["Sales"].sum())
            c_tot = float(inv_c["COGS"].sum())
            g_tot = float(inv_c["Adjusted GM"].sum())
            self.tree_cust.insert(
                "",
                "end",
                values=("TOTAL", "", "", "", cust, fmt_money(s_tot), fmt_money(c_tot), fmt_money(g_tot), f"{safe_pct(g_tot, s_tot):.2f}")
            )

    # =========================
    # Generic assignment window
    # =========================
    def _open_assignment_window(self, title: str, info_text: str, prompt_df: pd.DataFrame, save_callback):
        if not self.allowed_salespeople:
            messagebox.showwarning("Salaries", "Load salaries first. Only salaried salespeople can be assigned.")
            return

        dlg = tk.Toplevel(self.root)
        dlg.title(title)
        dlg.geometry("1320x560")
        dlg.transient(self.root)
        dlg.grab_set()

        ttk.Label(dlg, text=info_text, padding=10).pack(anchor="w")

        frame = ttk.Frame(dlg, padding=10)
        frame.pack(fill="both", expand=True)

        cols = ["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"]
        tree = ttk.Treeview(frame, columns=cols, show="headings", height=18)
        for c in cols:
            tree.heading(c, text=c)
            if c in ("Label","Partner","Account"):
                tree.column(c, width=240, anchor="w")
            elif c == "RowID":
                tree.column(c, width=140, anchor="w")
            elif c == "Salesperson":
                tree.column(c, width=160, anchor="w")
            else:
                tree.column(c, width=120, anchor="w")
        tree.pack(side="left", fill="both", expand=True)

        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        for _, r in prompt_df.iterrows():
            dt = r.get("Date")
            dt_s = "" if pd.isna(dt) else pd.to_datetime(dt).strftime("%Y-%m-%d")
            ymv = safe_str(r.get("YearMonth", ""))
            amt = float(r.get("Amount", 0.0))
            tree.insert("", "end", values=(
                safe_str(r.get("RowID","")), dt_s, ymv, fmt_money(amt),
                safe_str(r.get("Partner","")), safe_str(r.get("Label","")),
                safe_str(r.get("Number","")), safe_str(r.get("Account","")), ""
            ))

        controls = ttk.Frame(dlg, padding=10)
        controls.pack(fill="x")

        ttk.Label(controls, text="Assign selected rows to:").pack(side="left")
        sp_var = tk.StringVar(value=self.allowed_salespeople[0] if self.allowed_salespeople else "")
        sp_combo = ttk.Combobox(controls, textvariable=sp_var, values=self.allowed_salespeople, state="readonly", width=28)
        sp_combo.pack(side="left", padx=10)

        def assign_selected():
            sp = sp_var.get().strip()
            if not sp:
                return
            for iid in tree.selection():
                vals = list(tree.item(iid, "values"))
                vals[-1] = sp
                tree.item(iid, values=vals)

        def save_and_close():
            new_rows = []
            for iid in tree.get_children():
                v = tree.item(iid, "values")
                rowid = safe_str(v[0])
                dt_s = safe_str(v[1])
                dt_parsed = pd.to_datetime(dt_s, errors="coerce")
                ymv = "" if pd.isna(dt_parsed) else dt_parsed.strftime("%Y-%m")
                amt = to_num(v[3])
                partner = safe_str(v[4])
                label = safe_str(v[5])
                number = safe_str(v[6])
                account = safe_str(v[7])
                sp = safe_str(v[8])

                new_rows.append({
                    "RowID": rowid,
                    "Date": dt_parsed,
                    "YearMonth": ymv,
                    "Amount": amt,
                    "Partner": partner,
                    "Label": label,
                    "Number": number,
                    "Account": account,
                    "Salesperson": sp
                })

            new_df = pd.DataFrame(new_rows)
            save_callback(new_df)
            dlg.destroy()
            self.refresh_view()
            messagebox.showinfo("Saved", "Assignments saved and loaded into system.")

        ttk.Button(controls, text="Assign Selected", command=assign_selected).pack(side="left")
        ttk.Button(controls, text="Save & Close", command=save_and_close).pack(side="right")

    # =========================
    # Commissions extraction + assignment
    # =========================
    def extract_commission_lines_from_odoo(self):
        if self.df_raw is None:
            return pd.DataFrame()

        df = self.df_raw.copy()
        df["AccountStr"] = df["Account"].astype(str)
        is_comm = df["AccountStr"].str.contains(r"\b211810\b", regex=True) | df["AccountStr"].str.startswith("211810")
        comm = df[is_comm].copy()
        if comm.empty:
            return comm

        comm["Date"] = pd.to_datetime(comm["Date"], errors="coerce")
        comm["YearMonth"] = comm["Date"].dt.strftime("%Y-%m")
        comm["Amount"] = (pd.to_numeric(comm["Debit"], errors="coerce").fillna(0.0)
                          - pd.to_numeric(comm["Credit"], errors="coerce").fillna(0.0)).astype(float)

        def build_id(r):
            parts = [
                safe_str(r.get("Number","")),
                safe_str(r.get("Account","")),
                safe_str(r.get("Partner","")),
                safe_str(r.get("Label","")),
                "" if pd.isna(r.get("Date")) else r["Date"].strftime("%Y-%m-%d"),
                f"{float(r.get('Debit',0.0)):.2f}",
                f"{float(r.get('Credit',0.0)):.2f}",
            ]
            return make_row_id(parts)

        comm["RowID"] = comm.apply(build_id, axis=1)
        comm["Partner"] = comm["Partner"].apply(safe_str)
        comm["Label"] = comm["Label"].apply(safe_str)
        comm["Number"] = comm["Number"].apply(safe_str)
        comm["Account"] = comm["Account"].apply(safe_str)

        return comm[["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account"]].copy()

    def assign_new_commissions_if_any(self):
        comm_lines = self.extract_commission_lines_from_odoo()
        if comm_lines.empty:
            return

        existing_ids = set(self.commission_db["RowID"].astype(str).map(str.strip).tolist()) if not self.commission_db.empty else set()
        new_rows = comm_lines[~comm_lines["RowID"].isin(existing_ids)].copy()

        if new_rows.empty:
            return

        def _save(new_df: pd.DataFrame):
            if self.commission_db.empty:
                self.commission_db = new_df
            else:
                self.commission_db = self.commission_db[~self.commission_db["RowID"].isin(new_df["RowID"])]
                self.commission_db = pd.concat([self.commission_db, new_df], ignore_index=True)
            self.save_commissions_db()
            self.load_commissions_db()

        self._open_assignment_window(
            title="Assign Salesperson to NEW Commission Entries (Account 211810)",
            info_text=f"Assign salesperson to NEW commission entries. Saved to:\n{COMMISSIONS_FILE}",
            prompt_df=new_rows,
            save_callback=_save
        )

    # =========================
    # Credit Notes extraction + assignment
    # =========================
    def extract_creditnote_lines_from_odoo(self):
        if self.df_raw is None:
            return pd.DataFrame()

        df = self.df_raw.copy()
        cn = df[df["Number"].astype(str).str.strip().str.startswith(CREDIT_NOTE_PREFIX)].copy()
        if cn.empty:
            return cn

        cn["Date"] = pd.to_datetime(cn["Date"], errors="coerce")
        cn["YearMonth"] = cn["Date"].dt.strftime("%Y-%m")
        cn["Amount"] = (pd.to_numeric(cn["Debit"], errors="coerce").fillna(0.0)
                        - pd.to_numeric(cn["Credit"], errors="coerce").fillna(0.0)).astype(float).abs()

        def build_id(r):
            parts = [
                safe_str(r.get("Number","")),
                safe_str(r.get("Account","")),
                safe_str(r.get("Partner","")),
                safe_str(r.get("Label","")),
                "" if pd.isna(r.get("Date")) else r["Date"].strftime("%Y-%m-%d"),
                f"{float(r.get('Debit',0.0)):.2f}",
                f"{float(r.get('Credit',0.0)):.2f}",
            ]
            return make_row_id(parts)

        cn["RowID"] = cn.apply(build_id, axis=1)
        cn["Partner"] = cn["Partner"].apply(safe_str)
        cn["Label"] = cn["Label"].apply(safe_str)
        cn["Number"] = cn["Number"].apply(safe_str)
        cn["Account"] = cn["Account"].apply(safe_str)

        return cn[["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account"]].copy()

    def assign_new_creditnotes_if_any(self):
        cn_lines = self.extract_creditnote_lines_from_odoo()
        if cn_lines.empty:
            return

        existing_ids = set(self.creditnote_db["RowID"].astype(str).map(str.strip).tolist()) if not self.creditnote_db.empty else set()
        new_rows = cn_lines[~cn_lines["RowID"].isin(existing_ids)].copy()

        if new_rows.empty:
            return

        def _save(new_df: pd.DataFrame):
            if self.creditnote_db.empty:
                self.creditnote_db = new_df
            else:
                self.creditnote_db = self.creditnote_db[~self.creditnote_db["RowID"].isin(new_df["RowID"])]
                self.creditnote_db = pd.concat([self.creditnote_db, new_df], ignore_index=True)
            self.save_creditnotes_db()
            self.load_creditnotes_db()

        self._open_assignment_window(
            title="Assign Salesperson to NEW Credit Notes (RMHD/...)",
            info_text=f"Assign salesperson to NEW credit note entries. Saved to:\n{CREDIT_NOTES_FILE}\n\n"
                      f"You can later edit/correct from 'Edit Credit Notes'.",
            prompt_df=new_rows,
            save_callback=_save
        )

    def creditnote_line_items_in_range(self, sp: str) -> pd.DataFrame:
        if self.creditnote_db is None or self.creditnote_db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID","Salesperson"])

        db = self.creditnote_db.copy()
        db["Salesperson"] = db["Salesperson"].apply(safe_str).str.strip()
        db = db[db["Salesperson"] == sp]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID","Salesperson"])

        db["Date"] = pd.to_datetime(db["Date"], errors="coerce")
        db = db[~db["Date"].isna()]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID","Salesperson"])

        start, end = self.current_range
        if start is not None:
            db = db[db["Date"].dt.date >= start]
        if end is not None:
            db = db[db["Date"].dt.date <= end]
        if db.empty:
            return pd.DataFrame(columns=["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID","Salesperson"])

        db["YearMonth"] = db["Date"].dt.strftime("%Y-%m")
        db = db.sort_values(["Date", "Number", "Partner"], na_position="last")
        return db[["Date","YearMonth","Amount","Partner","Label","Number","Account","RowID","Salesperson"]].copy()

    # =========================
    # Credit Note impact (computed from Odoo file; assigned salesperson from DB)
    # =========================
    def creditnote_number_to_salesperson(self) -> dict:
        if self.creditnote_db is None or self.creditnote_db.empty:
            return {}
        db = self.creditnote_db.copy()
        db["Salesperson"] = db["Salesperson"].apply(safe_str).str.strip()
        db = db[db["Salesperson"] != ""]
        if db.empty:
            return {}
        m = {}
        for num, g in db.groupby("Number"):
            sp = next((x for x in g["Salesperson"].tolist() if x), "")
            if sp:
                m[str(num).strip()] = sp
        return m

    def build_creditnote_summary(self) -> pd.DataFrame:
        if self.df_raw is None:
            return pd.DataFrame(columns=[
                "Number","Date","YearMonth","Partner","SalesEffect","COGSEffect","GMEffect","Salesperson"
            ])

        df = self.df_raw.copy()
        cn = df[df["Number"].astype(str).str.strip().str.startswith(CREDIT_NOTE_PREFIX)].copy()
        if cn.empty:
            return pd.DataFrame(columns=[
                "Number","Date","YearMonth","Partner","SalesEffect","COGSEffect","GMEffect","Salesperson"
            ])

        cn["Date"] = pd.to_datetime(cn["Date"], errors="coerce")
        cn = cn[~cn["Date"].isna()].copy()
        if cn.empty:
            return pd.DataFrame(columns=[
                "Number","Date","YearMonth","Partner","SalesEffect","COGSEffect","GMEffect","Salesperson"
            ])

        cn["acct_l"] = cn["Account"].astype(str).str.lower()
        is_sales = cn["acct_l"].apply(lambda x: any(k in x for k in SALES_KEYWORDS))
        is_cogs = cn["acct_l"].apply(lambda x: any(k in x for k in COGS_KEYWORDS))

        cn["SalesSigned"] = 0.0
        cn.loc[is_sales, "SalesSigned"] = (cn.loc[is_sales, "Credit"] - cn.loc[is_sales, "Debit"]).astype(float)

        cn["COGSSigned"] = 0.0
        cn.loc[is_cogs, "COGSSigned"] = (cn.loc[is_cogs, "Debit"] - cn.loc[is_cogs, "Credit"]).astype(float)

        def pick_partner(g):
            gs = g[g["SalesSigned"].abs() > 1e-9]
            if not gs.empty:
                return safe_str(gs.iloc[0].get("Partner", ""))
            return safe_str(g.iloc[0].get("Partner", ""))

        g = cn.groupby("Number")
        out = pd.DataFrame({
            "Number": list(g.groups.keys()),
            "Date": g["Date"].min().values,
            "SalesEffect": g["SalesSigned"].sum().values,
            "COGSEffect": g["COGSSigned"].sum().values,
        })
        out["Partner"] = [pick_partner(cn[cn["Number"] == n]) for n in out["Number"].tolist()]
        out["YearMonth"] = pd.to_datetime(out["Date"]).dt.strftime("%Y-%m")
        out["GMEffect"] = out["SalesEffect"] - out["COGSEffect"]

        num_to_sp = self.creditnote_number_to_salesperson()
        out["Salesperson"] = out["Number"].astype(str).str.strip().map(lambda x: num_to_sp.get(x, ""))

        out = out[(out["SalesEffect"].abs() > 1e-9) | (out["COGSEffect"].abs() > 1e-9)].copy()
        return out

    # =========================
    # Incentives
    # =========================
    def compute_incentives(self, sp: str, cn_summary_for_sp: pd.DataFrame):
        # FIX: never apply _filter_by_range here — carry chain needs ALL history from SCHEME_START_DATE
        inv = self.invoice_summary[self.invoice_summary["Salesperson"].astype(str).str.strip() == sp].copy()
        bills = self.bill_lines[self.bill_lines["ExpensePerson"].astype(str).str.strip() == sp].copy()

        inv["DateOnly"] = inv["Date"].dt.date
        bills["DateOnly"] = bills["Date"].dt.date

        inv_scheme = inv[inv["DateOnly"] >= SCHEME_START_DATE].copy()
        bills_scheme = bills[bills["DateOnly"] >= SCHEME_START_DATE].copy()

        cn_effect_scheme = (cn_summary_for_sp.copy()
                            if (cn_summary_for_sp is not None and not cn_summary_for_sp.empty)
                            else pd.DataFrame())
        if not cn_effect_scheme.empty:
            cn_effect_scheme["DateOnly"] = pd.to_datetime(cn_effect_scheme["Date"]).dt.date
            cn_effect_scheme = cn_effect_scheme[cn_effect_scheme["DateOnly"] >= SCHEME_START_DATE].copy()

        if inv_scheme.empty and bills_scheme.empty and cn_effect_scheme.empty:
            return [], {"note": f"No incentive-applicable data from {SCHEME_START_DATE}."}

        adjgm_month = inv_scheme.groupby("YearMonth")["Adjusted GM"].sum()
        cn_gm_month = (cn_effect_scheme.groupby("YearMonth")["GMEffect"].sum()
                       if not cn_effect_scheme.empty else pd.Series(dtype=float))

        all_months_for_gm = sorted(set(adjgm_month.index.tolist()) | set(cn_gm_month.index.tolist()))
        gm_by_month = {}
        for m in all_months_for_gm:
            gm_by_month[m] = float(adjgm_month.get(m, 0.0)) + float(cn_gm_month.get(m, 0.0))

        emp_exp_month = bills_scheme.groupby("YearMonth")["ExpenseAmount"].sum()

        # FIX: gap-free month sequence so zero-sales months still advance carry/miss_streak
        raw_months = [m for m in (set(gm_by_month.keys()) | set(emp_exp_month.index.tolist()))
                      if re.match(r"^\d{4}-\d{2}$", str(m))]
        if not raw_months:
            return [], {"note": f"No incentive-applicable data from {SCHEME_START_DATE}."}
        first_m, last_m = min(raw_months), max(raw_months)
        months, y, mo = [], int(first_m[:4]), int(first_m[5:7])
        ey, emo = int(last_m[:4]), int(last_m[5:7])
        while (y, mo) <= (ey, emo):
            months.append(f"{y:04d}-{mo:02d}")
            mo += 1
            if mo > 12:
                mo = 1; y += 1

        comm_by_month = self.commission_monthly_breakdown(sp)
        base_net = self.use_net_for_incentive_var.get()

        # ✅ qualification set already computed globally; used for incentive + repeat
        qual_new_customers = set(self.qual_new_customers_global)

        repeats_by_month = {}
        for cust in qual_new_customers:
            first_dt = self.qual_new_customer_first_dt.get(cust)
            if not first_dt:
                continue
            cutoff = first_dt + timedelta(days=REPEAT_WINDOW_DAYS)

            inv_c = self.invoice_summary[
                (self.invoice_summary["Salesperson"].astype(str).str.strip() == sp)
                & (self.invoice_summary["Customer"] == cust)
            ].copy()
            if inv_c.empty:
                continue

            inv_c["DateOnly"] = inv_c["Date"].dt.date
            rep = inv_c[(inv_c["DateOnly"] > first_dt) & (inv_c["DateOnly"] <= cutoff)]
            if rep.empty:
                continue

            start, end = self.current_range
            if start:
                rep = rep[rep["DateOnly"] >= start]
            if end:
                rep = rep[rep["DateOnly"] <= end]
            if rep.empty:
                continue

            first_repeat_dt = sorted(rep["DateOnly"].tolist())[0]
            m = ym_from_date(first_repeat_dt)
            repeats_by_month[m] = repeats_by_month.get(m, 0) + 1

        monthly_rows = []
        carry = 0.0
        miss_streak = 0
        total_paid = 0.0

        for m in months:
            salary = self.salary_for_month(sp, m)
            if salary <= 0:
                monthly_rows.append({
                    "Month": m, "Salary": 0.0, "Base": 0.0, "Threshold": 0.0, "CarryIn": carry,
                    "Met": "NO", "Inc10": 0.0, "NewBonus": 0.0, "RepeatBonus": 0.0, "Total": 0.0, "Notice": ""
                })
                continue

            adjgm_effective = float(gm_by_month.get(m, 0.0))
            emp_exp = float(emp_exp_month.get(m, 0.0))
            comm = float(comm_by_month.get(m, 0.0))

            base_val = (adjgm_effective - emp_exp - comm) if base_net else adjgm_effective
            threshold_base = float(_db_get_setting("THRESHOLD_SALARY_MULTIPLE", "6.0")) * salary
            threshold = threshold_base + carry

            met = base_val >= threshold - 1e-9
            inc_10 = 0.0
            if met:
                inc_10 = (float(_db_get_setting("INCENTIVE_PERCENTAGE", "10.0")) / 100.0) * base_val
                carry = 0.0
                miss_streak = 0
            else:
                carry = (threshold - base_val)
                miss_streak += 1

            notice_flag = "YES" if miss_streak >= 4 else ""

            # ✅ New customer bonus: only QUALIFIED customers count as "new customer" now
            inv_m = inv_scheme[inv_scheme["YearMonth"] == m].copy()
            new_bonus = 0.0
            if not inv_m.empty:
                custs = set(inv_m["Customer"].astype(str).map(str.strip).tolist())
                qualifying_this_month = []
                for cust2 in custs:
                    if cust2 in self.qual_new_customers_global:
                        first_dt2 = self.qual_new_customer_first_dt.get(cust2)
                        if first_dt2 and ym_from_date(first_dt2) == m:
                            qualifying_this_month.append(cust2)

                if len(qualifying_this_month) >= MIN_NEW_CUST_COUNT:
                    new_bonus = NEW_CUST_BONUS * float(len(qualifying_this_month))

            repeat_bonus = REPEAT_BONUS * float(repeats_by_month.get(m, 0))
            total = inc_10 + new_bonus + repeat_bonus
            total_paid += total

            monthly_rows.append({
                "Month": m,
                "Salary": salary,
                "Base": base_val,
                "Threshold": threshold,
                "CarryIn": (threshold - threshold_base),
                "Met": "YES" if met else "NO",
                "Inc10": inc_10,
                "NewBonus": new_bonus,
                "RepeatBonus": repeat_bonus,
                "Total": total,
                "Notice": notice_flag,
            })

        return monthly_rows, {"note": "", "total_paid": total_paid, "base_is_net": base_net}

    # =========================
    # Refresh view (GUI)
    # =========================
    def refresh_view(self):
        selected_sps = self.get_selected_salespeople()
        if not selected_sps or self.invoice_summary is None:
            return

        # keep salesperson_var in sync for export / other methods
        sp = selected_sps[0]
        self.salesperson_var.set(sp)

        for t in [self.tree_txn, self.tree_inc, self.tree_comm, self.tree_cn]:
            for item in t.get_children():
                t.delete(item)

        if len(selected_sps) > 1:
            self._refresh_multi_sp(selected_sps)
            return

        inv = self.invoice_summary[self.invoice_summary["Salesperson"].astype(str).str.strip() == sp].copy()
        bills = self.bill_lines[self.bill_lines["ExpensePerson"].astype(str).str.strip() == sp].copy()

        inv = self._filter_by_range(inv, "Date")
        bills = self._filter_by_range(bills, "Date")

        # Apply product filter
        if self.product_filter_active and self.df_raw is not None:
            mask = self.df_raw["Label"].str.contains(self.product_filter_active, case=False, na=False)
            matching_invs = set(
                self.df_raw[mask & self.df_raw["Number"].apply(is_invoice_number_value)]["Number"].tolist()
            )
            inv = inv[inv["Number"].isin(matching_invs)]

        inv = inv.sort_values(["Date", "Number"], na_position="last")
        bills = bills.sort_values(["Date", "Number"], na_position="last")

        cn_summary = self.build_creditnote_summary()
        if not cn_summary.empty:
            cn_summary = cn_summary[cn_summary["Salesperson"].astype(str).str.strip() == sp].copy()
            cn_summary["Date"] = pd.to_datetime(cn_summary["Date"], errors="coerce")
            cn_summary = cn_summary[~cn_summary["Date"].isna()]
            cn_summary = self._filter_by_range(cn_summary, "Date")
            cn_summary = cn_summary.sort_values(["Date", "Number"], na_position="last")

        sales_total = float(inv["Sales"].sum()) if not inv.empty else 0.0
        cogs_total = float(inv["COGS"].sum()) if not inv.empty else 0.0
        transport_total = float(inv["Transport"].sum()) if not inv.empty else 0.0
        loading_total = float(inv["Loading"].sum()) if not inv.empty else 0.0
        adjgm_total = float(inv["Adjusted GM"].sum()) if not inv.empty else 0.0

        cn_sales_total = float(cn_summary["SalesEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0
        cn_cogs_total = float(cn_summary["COGSEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0
        cn_gm_total = float(cn_summary["GMEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0

        sales_total += cn_sales_total
        cogs_total += cn_cogs_total
        adjgm_total += cn_gm_total

        emp_exp_total = float(bills["ExpenseAmount"].sum()) if not bills.empty else 0.0

        comm_by_month = self.commission_monthly_breakdown(sp)
        comm_total = float(sum(comm_by_month.values())) if comm_by_month else 0.0

        net_total = adjgm_total - emp_exp_total - comm_total
        total_gm_pct = safe_pct(adjgm_total, sales_total)

        # Invoices
        for _, r in inv.iterrows():
            dt = r["Date"]
            date_s = "" if pd.isna(dt) else dt.strftime("%Y-%m-%d")
            inv_date = dt.date() if not pd.isna(dt) else None
            cust = safe_str(r.get("Customer", r.get("Partner", "")))

            sales = float(r.get("Sales", 0.0))
            adjgm = float(r.get("Adjusted GM", 0.0))
            gm_pct = safe_pct(adjgm, sales)

            tags = ()
            if inv_date:
                # ✅ Requirement #6: global new customers green+bold
                if self.is_global_new_customer_month(cust, inv_date):
                    tags = ("GLOBAL_NEW",)

                # ✅ Requirement #5: qualified new customers bold (already green+bold in GUI)
                if self.is_qualified_new_customer_month(cust, inv_date):
                    tags = tuple(set(tags + ("QUAL_NEW",)))

                # ✅ Requirement #6: qualified repeat customers blue
                if self.is_qualified_repeat_invoice(cust, inv_date):
                    tags = ("QUAL_REPEAT",)

            self.tree_txn.insert(
                "",
                "end",
                values=(
                    "Invoice",
                    date_s,
                    safe_str(r["Number"]),
                    safe_str(r["Partner"]),
                    fmt_money(sales),
                    fmt_money(float(r.get("COGS", 0.0))),
                    fmt_money(float(r.get("Transport", 0.0))),
                    fmt_money(float(r.get("Loading", 0.0))),
                    fmt_money(adjgm),
                    f"{gm_pct:.2f}",   # ✅ Requirement #1
                    "",
                    "",
                    fmt_money(adjgm),
                ),
                tags=tags
            )

        # Credit Notes
        if cn_summary is not None and not cn_summary.empty:
            if not inv.empty:
                self.tree_txn.insert("", "end", values=("", "", "", "", "", "", "", "", "", "", "", "", ""))
            for _, r in cn_summary.iterrows():
                dt = r["Date"]
                date_s = "" if pd.isna(dt) else pd.to_datetime(dt).strftime("%Y-%m-%d")
                partner = safe_str(r.get("Partner", ""))
                sales_eff = float(r.get("SalesEffect", 0.0))
                cogs_eff = float(r.get("COGSEffect", 0.0))
                gm_eff = float(r.get("GMEffect", 0.0))
                gm_pct_cn = safe_pct(gm_eff, sales_eff) if abs(sales_eff) > 1e-9 else 0.0

                self.tree_txn.insert(
                    "",
                    "end",
                    values=(
                        "Credit Note",
                        date_s,
                        safe_str(r.get("Number", "")),
                        partner,
                        fmt_money(sales_eff),
                        fmt_money(cogs_eff),
                        fmt_money(0.0),
                        fmt_money(0.0),
                        fmt_money(gm_eff),
                        f"{gm_pct_cn:.2f}",
                        "",
                        "",
                        fmt_money(gm_eff),
                    ),
                )

        # Employee expenses
        if (not inv.empty or (cn_summary is not None and not cn_summary.empty)) and (not bills.empty):
            self.tree_txn.insert("", "end", values=("", "", "", "", "", "", "", "", "", "", "", "", ""))

        for _, r in bills.iterrows():
            date_s = "" if pd.isna(r["Date"]) else r["Date"].strftime("%Y-%m-%d")
            info = safe_str(r["Label"]) if safe_str(r["Label"]) else safe_str(r["Account"])
            amt = float(r["ExpenseAmount"])
            self.tree_txn.insert(
                "",
                "end",
                values=(
                    "Emp Exp",
                    date_s,
                    safe_str(r["Number"]),
                    info,
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                    fmt_money(-amt),
                    "",
                    fmt_money(-amt),
                ),
            )

        # Commission monthly rows
        if comm_by_month:
            self.tree_txn.insert("", "end", values=("", "", "", "", "", "", "", "", "", "", "", "", ""))
            for m in sorted(comm_by_month.keys()):
                v = float(comm_by_month[m])
                if abs(v) > 1e-9:
                    self.tree_txn.insert(
                        "",
                        "end",
                        values=(
                            "Commission",
                            m,
                            "",
                            "Account 211810 (assigned)",
                            "",
                            "",
                            "",
                            "",
                            "",
                            "",
                            "",
                            fmt_money(-v),
                            fmt_money(-v),
                        ),
                    )

        # Totals (✅ Requirement #1 total GM%)
        self.tree_txn.insert("", "end", values=("", "", "", "", "", "", "", "", "", "", "", "", ""))
        self.tree_txn.insert(
            "",
            "end",
            values=(
                "TOTAL",
                "",
                "",
                sp,
                fmt_money(sales_total),
                fmt_money(cogs_total),
                fmt_money(transport_total),
                fmt_money(loading_total),
                fmt_money(adjgm_total),
                f"{total_gm_pct:.2f}",
                fmt_money(-emp_exp_total),
                fmt_money(-comm_total),
                fmt_money(net_total),
            ),
        )

        self.lbl_sales.config(text=f"Sales: {sales_total:,.2f}")
        self.lbl_cogs.config(text=f"COGS: {cogs_total:,.2f}")
        self.lbl_chg.config(text=f"Charges (T+L): {(transport_total + loading_total):,.2f}")
        self.lbl_adjgm.config(text=f"Adjusted GM: {adjgm_total:,.2f}")
        self.lbl_exp.config(text=f"Employee Exp: {emp_exp_total:,.2f}")
        self.lbl_comm.config(text=f"Commission: {comm_total:,.2f}")
        self.lbl_net.config(text=f"Net: {net_total:,.2f}")

        # FIX: pass full unfiltered CN data so carry chain is correct
        cn_for_incentive = self.build_creditnote_summary()
        if not cn_for_incentive.empty:
            cn_for_incentive = cn_for_incentive[cn_for_incentive["Salesperson"].astype(str).str.strip() == sp].copy()
            cn_for_incentive["Date"] = pd.to_datetime(cn_for_incentive["Date"], errors="coerce")
            cn_for_incentive = cn_for_incentive[~cn_for_incentive["Date"].isna()]
        else:
            cn_for_incentive = pd.DataFrame()
        monthly_rows, summary = self.compute_incentives(sp, cn_summary_for_sp=cn_for_incentive)
        if summary.get("note"):
            self.lbl_incent.config(text="Incentive (₹): 0.00")
        else:
            total_paid = float(summary.get("total_paid", 0.0))
            self.lbl_incent.config(text=f"Incentive (₹): {total_paid:,.2f}")

        for r in monthly_rows:
            self.tree_inc.insert(
                "",
                "end",
                values=(
                    r["Month"],
                    fmt_money(r.get("Salary", 0.0)),
                    fmt_money(r.get("Base", 0.0)),
                    fmt_money(r.get("Threshold", 0.0)),
                    fmt_money(r.get("CarryIn", 0.0)),
                    r.get("Met", ""),
                    fmt_money(r.get("Inc10", 0.0)),
                    fmt_money(r.get("NewBonus", 0.0)),
                    fmt_money(r.get("RepeatBonus", 0.0)),
                    fmt_money(r.get("Total", 0.0)),
                    r.get("Notice", ""),
                ),
            )

        # Commission Details tab
        comm_items = self.commission_line_items_in_range(sp)
        if not comm_items.empty:
            for _, r in comm_items.iterrows():
                dt_s = "" if pd.isna(r["Date"]) else pd.to_datetime(r["Date"]).strftime("%Y-%m-%d")
                self.tree_comm.insert(
                    "",
                    "end",
                    values=(
                        dt_s,
                        safe_str(r["YearMonth"]),
                        fmt_money(float(r["Amount"])),
                        safe_str(r["Partner"]),
                        safe_str(r["Label"]),
                        safe_str(r["Number"]),
                        safe_str(r["Account"]),
                        safe_str(r["RowID"]),
                    )
                )
            self.tree_comm.insert("", "end", values=("", "", "", "", "", "", "", ""))
            self.tree_comm.insert("", "end", values=("TOTAL", "", fmt_money(float(comm_items["Amount"].sum())), "", "", "", "", ""))
        else:
            self.tree_comm.insert("", "end", values=("No commission entries in selected range", "", "", "", "", "", "", ""))

        # Credit Note Details tab
        cn_items = self.creditnote_line_items_in_range(sp)
        if not cn_items.empty:
            for _, r in cn_items.iterrows():
                dt_s = "" if pd.isna(r["Date"]) else pd.to_datetime(r["Date"]).strftime("%Y-%m-%d")
                self.tree_cn.insert(
                    "",
                    "end",
                    values=(
                        dt_s,
                        safe_str(r["YearMonth"]),
                        fmt_money(float(r["Amount"])),
                        safe_str(r["Partner"]),
                        safe_str(r["Label"]),
                        safe_str(r["Number"]),
                        safe_str(r["Account"]),
                        safe_str(r["RowID"]),
                        safe_str(r["Salesperson"]),
                    )
                )
            self.tree_cn.insert("", "end", values=("", "", "", "", "", "", "", "", ""))
            self.tree_cn.insert("", "end", values=("TOTAL", "", fmt_money(float(cn_items["Amount"].sum())), "", "", "", "", "", ""))
        else:
            self.tree_cn.insert("", "end", values=("No credit note entries in selected range", "", "", "", "", "", "", "", ""))

        # Incentive text
        self.incent_text.delete("1.0", "end")
        start, end = self.current_range
        range_label = "All Dates" if (start is None and end is None) else f"{start} to {end}"
        base_txt = "Net (AdjGM incl CN - EmpExp - Commission)" if self.use_net_for_incentive_var.get() else "Adjusted GM incl CN"
        self.incent_text.insert("end", f"Salesperson: {sp}\n")
        self.incent_text.insert("end", f"Range: {range_label}\n")
        self.incent_text.insert("end", f"Scheme Start: {SCHEME_START_DATE}\n")
        self.incent_text.insert("end", f"Base Used: {base_txt}\n")
        self.incent_text.insert("end", f"New customer QUALIFICATION: Sales >= ₹{NEW_CUST_BILL_MIN:,.0f} AND GM% >= {NEW_CUST_MIN_MARGIN_PCT:.2f}% in first month.\n")
        self.incent_text.insert("end", "Repeat bonus is considered ONLY for QUALIFIED new customers (and only once within 90 days).\n")
        self.incent_text.insert("end", "Credit notes reduce Sales/COGS correctly and therefore reduce Adjusted GM.\n")

        pf_txt = f" | Product: {self.product_filter_active}" if self.product_filter_active else ""
        self.status_var.set(
            f"Showing: {sp}{pf_txt} | Range: {range_label} | Salaries: {len(self.allowed_salespeople)} | "
            f"Commission rows: {len(self.commission_db)} | Credit note rows: {len(self.creditnote_db)} | "
            f"Qualified new customers (global): {len(self.qual_new_customers_global)}"
        )

    # =========================
    # Export (Excel / PDF)
    # =========================
    def export_report(self):
        selected_sps = self.get_selected_salespeople()
        if not selected_sps:
            messagebox.showwarning("Salesperson", "Please select a salesperson first.")
            return
        if self.invoice_summary is None:
            messagebox.showwarning("No data", "Please load an Odoo Excel file first.")
            return
        sp = selected_sps[0]
        if len(selected_sps) > 1:
            messagebox.showinfo("Export Note",
                f"Detailed export is per salesperson.\nExporting for: {sp}\n\n"
                "Use '📊 Team Summary → Export to Excel' for all selected.")

        dlg = tk.Toplevel(self.root)
        dlg.title("Export Format")
        dlg.geometry("360x230")
        dlg.transient(self.root)
        dlg.grab_set()

        fmt_var = tk.StringVar(value="excel")
        ttk.Label(dlg, text="Choose export format:", padding=10).pack(anchor="w")
        ttk.Radiobutton(dlg, text="Excel (.xlsx)", value="excel", variable=fmt_var).pack(anchor="w", padx=12)
        ttk.Radiobutton(dlg, text="PDF (.pdf)", value="pdf", variable=fmt_var).pack(anchor="w", padx=12)

        include_lines_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(dlg, text="Include Line Items sheet (Excel only)",
                        variable=include_lines_var).pack(anchor="w", padx=12, pady=(6, 0))

        def txn_row_tag_for_invoice(customer: str, inv_date: date) -> str:
            if not inv_date:
                return ""
            if self.is_qualified_repeat_invoice(customer, inv_date):
                return "QUAL_REPEAT"
            if self.is_qualified_new_customer_month(customer, inv_date):
                return "QUAL_NEW"
            if self.is_global_new_customer_month(customer, inv_date):
                return "GLOBAL_NEW"
            return ""

        def build_export_df_txn():
            inv = self.invoice_summary[self.invoice_summary["Salesperson"].astype(str).str.strip() == sp].copy()
            bills = self.bill_lines[self.bill_lines["ExpensePerson"].astype(str).str.strip() == sp].copy()
            inv = self._filter_by_range(inv, "Date")
            bills = self._filter_by_range(bills, "Date")
            inv = inv.sort_values(["Date", "Number"], na_position="last")
            bills = bills.sort_values(["Date", "Number"], na_position="last")

            cn_summary = self.build_creditnote_summary()
            if not cn_summary.empty:
                cn_summary = cn_summary[cn_summary["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_summary["Date"] = pd.to_datetime(cn_summary["Date"], errors="coerce")
                cn_summary = cn_summary[~cn_summary["Date"].isna()]
                cn_summary = self._filter_by_range(cn_summary, "Date")
                cn_summary = cn_summary.sort_values(["Date", "Number"], na_position="last")

            comm_by_month = self.commission_monthly_breakdown(sp)

            rows = []
            for _, r in inv.iterrows():
                inv_date = None
                if not pd.isna(r["Date"]):
                    inv_date = r["Date"].date()
                cust = safe_str(r.get("Customer", ""))
                sales = float(r["Sales"])
                adjgm = float(r["Adjusted GM"])
                gm_pct = safe_pct(adjgm, sales)
                tag = txn_row_tag_for_invoice(cust, inv_date)

                rows.append({
                    "Type": "Invoice",
                    "Date": "" if pd.isna(r["Date"]) else r["Date"].strftime("%Y-%m-%d"),
                    "Number": safe_str(r["Number"]),
                    "Partner/Info": safe_str(r["Partner"]),
                    "Sales": float(r["Sales"]),
                    "COGS": float(r["COGS"]),
                    "Transport": float(r["Transport"]),
                    "Loading": float(r["Loading"]),
                    "Adjusted GM": float(r["Adjusted GM"]),
                    "GM %": gm_pct,
                    "Employee Exp": 0.0,
                    "Commission": 0.0,
                    "Net": float(r["Adjusted GM"]),
                    "_TAG": tag,
                })

            if cn_summary is not None and not cn_summary.empty:
                for _, r in cn_summary.iterrows():
                    sales_eff = float(r["SalesEffect"])
                    gm_eff = float(r["GMEffect"])
                    gm_pct_cn = safe_pct(gm_eff, sales_eff) if abs(sales_eff) > 1e-9 else 0.0
                    rows.append({
                        "Type": "Credit Note",
                        "Date": "" if pd.isna(r["Date"]) else pd.to_datetime(r["Date"]).strftime("%Y-%m-%d"),
                        "Number": safe_str(r["Number"]),
                        "Partner/Info": safe_str(r["Partner"]),
                        "Sales": float(r["SalesEffect"]),
                        "COGS": float(r["COGSEffect"]),
                        "Transport": 0.0,
                        "Loading": 0.0,
                        "Adjusted GM": float(r["GMEffect"]),
                        "GM %": gm_pct_cn,
                        "Employee Exp": 0.0,
                        "Commission": 0.0,
                        "Net": float(r["GMEffect"]),
                        "_TAG": "",
                    })

            for _, r in bills.iterrows():
                info = safe_str(r["Label"]) if safe_str(r["Label"]) else safe_str(r["Account"])
                amt = float(r["ExpenseAmount"])
                rows.append({
                    "Type": "Emp Exp",
                    "Date": "" if pd.isna(r["Date"]) else r["Date"].strftime("%Y-%m-%d"),
                    "Number": safe_str(r["Number"]),
                    "Partner/Info": info,
                    "Sales": 0.0,
                    "COGS": 0.0,
                    "Transport": 0.0,
                    "Loading": 0.0,
                    "Adjusted GM": 0.0,
                    "GM %": 0.0,
                    "Employee Exp": -amt,
                    "Commission": 0.0,
                    "Net": -amt,
                    "_TAG": "",
                })

            for m in sorted(comm_by_month.keys()):
                v = float(comm_by_month[m])
                if abs(v) > 1e-9:
                    rows.append({
                        "Type": "Commission",
                        "Date": m,
                        "Number": "",
                        "Partner/Info": "Account 211810 (assigned)",
                        "Sales": 0.0,
                        "COGS": 0.0,
                        "Transport": 0.0,
                        "Loading": 0.0,
                        "Adjusted GM": 0.0,
                        "GM %": 0.0,
                        "Employee Exp": 0.0,
                        "Commission": -v,
                        "Net": -v,
                        "_TAG": "",
                    })

            df = pd.DataFrame(rows)

            sales_total = float(inv["Sales"].sum()) if not inv.empty else 0.0
            cogs_total = float(inv["COGS"].sum()) if not inv.empty else 0.0
            transport_total = float(inv["Transport"].sum()) if not inv.empty else 0.0
            loading_total = float(inv["Loading"].sum()) if not inv.empty else 0.0
            adjgm_total = float(inv["Adjusted GM"].sum()) if not inv.empty else 0.0

            cn_sales_total = float(cn_summary["SalesEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0
            cn_cogs_total = float(cn_summary["COGSEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0
            cn_gm_total = float(cn_summary["GMEffect"].sum()) if (cn_summary is not None and not cn_summary.empty) else 0.0

            sales_total += cn_sales_total
            cogs_total += cn_cogs_total
            adjgm_total += cn_gm_total

            emp_exp_total = float(bills["ExpenseAmount"].sum()) if not bills.empty else 0.0
            comm_total = float(sum(comm_by_month.values())) if comm_by_month else 0.0
            net_total = adjgm_total - emp_exp_total - comm_total

            total_row = pd.DataFrame([{
                "Type": "TOTAL",
                "Date": "",
                "Number": "",
                "Partner/Info": sp,
                "Sales": sales_total,
                "COGS": cogs_total,
                "Transport": transport_total,
                "Loading": loading_total,
                "Adjusted GM": adjgm_total,
                "GM %": safe_pct(adjgm_total, sales_total),
                "Employee Exp": -emp_exp_total,
                "Commission": -comm_total,
                "Net": net_total,
                "_TAG": "",
            }])

            if not df.empty:
                df = pd.concat([df, pd.DataFrame([{}]), total_row], ignore_index=True)
            else:
                df = total_row

            return df

        def build_export_df_incentives(cn_summary_for_sp: pd.DataFrame):
            rows, summary = self.compute_incentives(sp, cn_summary_for_sp=cn_summary_for_sp)
            if summary.get("note"):
                return pd.DataFrame([{"Note": summary["note"], "Scheme Start": str(SCHEME_START_DATE)}]), None

            df = pd.DataFrame(rows)
            s = pd.DataFrame([
                {"Metric": "Scheme Start", "Value": str(SCHEME_START_DATE)},
                {"Metric": "Base Used", "Value": "Net (AdjGM incl CN - EmpExp - Commission)" if self.use_net_for_incentive_var.get() else "Adjusted GM incl CN"},
                {"Metric": "Total Incentive Payable", "Value": float(summary.get("total_paid", 0.0))},
            ])
            return df, s

        def build_export_df_comm_details():
            df = self.commission_line_items_in_range(sp).copy()
            if df.empty:
                return pd.DataFrame([{"Note": "No commission entries in selected range."}])
            df["Date"] = df["Date"].dt.strftime("%Y-%m-%d")
            df = df.rename(columns={"Amount": "Amount (Positive=Expense)"})
            return df

        def build_export_df_cn_details():
            df = self.creditnote_line_items_in_range(sp).copy()
            if df.empty:
                return pd.DataFrame([{"Note": "No credit note entries in selected range."}])
            df["Date"] = df["Date"].dt.strftime("%Y-%m-%d")
            df = df.rename(columns={"Amount": "Amount (Positive=Line Magnitude)"})
            return df

        def apply_excel_styles(writer, df_txn):
            # style "Transactions" sheet based on _TAG column
            ws = writer.book["Transactions"]

            # find _TAG column index
            headers = [cell.value for cell in ws[1]]
            if "_TAG" not in headers:
                return
            tag_col = headers.index("_TAG") + 1

            green_bold = Font(color="008000", bold=True)  # green
            blue_font = Font(color="0000FF")              # blue
            bold_font = Font(bold=True)

            # apply for each row (start at row=2)
            max_row = ws.max_row
            max_col = ws.max_column
            for r in range(2, max_row + 1):
                tag = ws.cell(row=r, column=tag_col).value
                if not tag:
                    continue

                # apply across row excluding _TAG col
                for c in range(1, max_col + 1):
                    if c == tag_col:
                        continue
                    cell = ws.cell(row=r, column=c)
                    if tag == "QUAL_REPEAT":
                        cell.font = blue_font
                    elif tag == "GLOBAL_NEW":
                        cell.font = green_bold
                    elif tag == "QUAL_NEW":
                        cell.font = bold_font

            # hide _TAG column
            ws.column_dimensions[ws.cell(row=1, column=tag_col).column_letter].hidden = True

        def export_pdf(sp_name, out_path: Path, cn_summary_for_sp: pd.DataFrame):
            df_txn = build_export_df_txn()
            df_inc, df_sum = build_export_df_incentives(cn_summary_for_sp=cn_summary_for_sp)
            df_comm = build_export_df_comm_details()
            df_cn = build_export_df_cn_details()

            # styling needs tags
            tag_list = []
            if "_TAG" in df_txn.columns:
                tag_list = df_txn["_TAG"].fillna("").tolist()
                df_txn = df_txn.drop(columns=["_TAG"])

            df_txn = df_txn.round(2)
            df_inc = df_inc.round(2)
            if df_sum is not None:
                df_sum = df_sum.round(2)
            try:
                df_comm = df_comm.round(2)
            except Exception:
                pass
            try:
                df_cn = df_cn.round(2)
            except Exception:
                pass

            start, end = self.current_range
            range_label = "All Dates" if (start is None and end is None) else f"{start} to {end}"

            styles = getSampleStyleSheet()
            story = []
            try:
                for _lp in LOGO_CANDIDATES:
                    if _lp.exists():
                        story.append(Image(str(_lp), width=180, height=45))
                        story.append(Spacer(1, 6))
                        break
            except Exception:
                pass
            story.append(Paragraph("<b>Salesperson Report</b>", styles["Title"]))
            story.append(Paragraph(f"<b>Salesperson:</b> {sp_name}", styles["Normal"]))
            story.append(Paragraph(f"<b>Range:</b> {range_label}", styles["Normal"]))
            story.append(Paragraph(f"<b>Scheme Start:</b> {SCHEME_START_DATE}", styles["Normal"]))
            story.append(Paragraph(f"<b>New customer qualification:</b> Sales >= ₹{NEW_CUST_BILL_MIN:,.0f} and GM% >= {NEW_CUST_MIN_MARGIN_PCT:.2f}% (first month).", styles["Normal"]))
            story.append(Spacer(1, 10))

            # Transactions
            story.append(Paragraph("<b>Transactions</b>", styles["Heading2"]))
            cols = list(df_txn.columns)
            data = [cols] + [[pdf_cell(r.get(c, "")) for c in cols] for _, r in df_txn.fillna("").iterrows()]
            t = Table(data, repeatRows=1)

            base_style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("ALIGN", (4, 1), (-1, -1), "RIGHT"),
            ]

            # apply row styles using saved tag_list (align with df rows)
            # df_txn has blank separator rows too; tags list same length
            for i, tag in enumerate(tag_list, start=1):  # +1 because header is row 0 in Table
                if tag == "QUAL_REPEAT":
                    base_style.append(("TEXTCOLOR", (0, i), (-1, i), colors.blue))
                elif tag == "GLOBAL_NEW":
                    base_style.append(("TEXTCOLOR", (0, i), (-1, i), colors.green))
                    base_style.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))
                elif tag == "QUAL_NEW":
                    base_style.append(("FONTNAME", (0, i), (-1, i), "Helvetica-Bold"))

            t.setStyle(TableStyle(base_style))
            story.append(t)
            story.append(PageBreak())

            # Incentives
            story.append(Paragraph("<b>Incentives</b>", styles["Heading2"]))
            cols2 = list(df_inc.columns)
            data2 = [cols2] + [[pdf_cell(r.get(c, "")) for c in cols2] for _, r in df_inc.fillna("").iterrows()]
            t2 = Table(data2, repeatRows=1)
            t2.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t2)

            if df_sum is not None:
                story.append(Spacer(1, 10))
                story.append(Paragraph("<b>Incentive Summary</b>", styles["Heading3"]))
                cols3 = list(df_sum.columns)
                data3 = [cols3] + [[pdf_cell(r.get(c, "")) for c in cols3] for _, r in df_sum.fillna("").iterrows()]
                t3 = Table(data3, repeatRows=1)
                t3.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                ]))
                story.append(t3)

            story.append(PageBreak())

            # Commission Details
            story.append(Paragraph("<b>Commission Details</b>", styles["Heading2"]))
            cols4 = list(df_comm.columns)
            data4 = [cols4] + [[pdf_cell(r.get(c, "")) for c in cols4] for _, r in df_comm.fillna("").iterrows()]
            t4 = Table(data4, repeatRows=1)
            t4.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t4)

            story.append(PageBreak())

            # Credit Note Details
            story.append(Paragraph("<b>Credit Note Details (Assignments)</b>", styles["Heading2"]))
            cols5 = list(df_cn.columns)
            data5 = [cols5] + [[pdf_cell(r.get(c, "")) for c in cols5] for _, r in df_cn.fillna("").iterrows()]
            t5 = Table(data5, repeatRows=1)
            t5.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            story.append(t5)

            doc = SimpleDocTemplate(
                str(out_path),
                pagesize=landscape(A4),
                rightMargin=18, leftMargin=18, topMargin=18, bottomMargin=18
            )
            doc.build(story)

        def do_export():
            fmt = fmt_var.get()
            dlg.destroy()

            folder = filedialog.askdirectory(title="Select folder to save export")
            if not folder:
                return

            safe_name = re.sub(r"[^\w\- ]+", "", sp).strip().replace(" ", "_")
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")

            start, end = self.current_range
            range_tag = "ALL" if (start is None and end is None) else f"{start}_to_{end}"
            range_tag = range_tag.replace("-", "")

            cn_summary_for_sp = self.build_creditnote_summary()
            if not cn_summary_for_sp.empty:
                cn_summary_for_sp = cn_summary_for_sp[cn_summary_for_sp["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_summary_for_sp["Date"] = pd.to_datetime(cn_summary_for_sp["Date"], errors="coerce")
                cn_summary_for_sp = cn_summary_for_sp[~cn_summary_for_sp["Date"].isna()]
                # FIX: no date-range filter here — incentive engine needs full history

            if fmt == "excel":
                out_path = Path(folder) / f"Sales_Report_{safe_name}_{range_tag}_{ts}.xlsx"
                try:
                    df_txn = build_export_df_txn()
                    df_inc, df_sum = build_export_df_incentives(cn_summary_for_sp=cn_summary_for_sp)
                    df_comm = build_export_df_comm_details()
                    df_cn = build_export_df_cn_details()

                    df_txn = df_txn.round(2)
                    df_inc = df_inc.round(2)
                    if df_sum is not None:
                        df_sum = df_sum.round(2)
                    try:
                        df_comm = df_comm.round(2)
                    except Exception:
                        pass
                    try:
                        df_cn = df_cn.round(2)
                    except Exception:
                        pass

                    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
                        df_txn.to_excel(writer, index=False, sheet_name="Transactions")
                        df_inc.to_excel(writer, index=False, sheet_name="Incentives")
                        df_comm.to_excel(writer, index=False, sheet_name="Commission_Details")
                        df_cn.to_excel(writer, index=False, sheet_name="CreditNote_Details")
                        if df_sum is not None:
                            df_sum.to_excel(writer, index=False, sheet_name="Incentive_Summary")

                        excel_format_2dp(writer.book["Transactions"])
                        excel_format_2dp(writer.book["Incentives"])
                        excel_format_2dp(writer.book["Commission_Details"])
                        excel_format_2dp(writer.book["CreditNote_Details"])
                        if df_sum is not None:
                            excel_format_2dp(writer.book["Incentive_Summary"])

                        # ✅ Requirement #5/#6: bold/blue/green formatting in report
                        apply_excel_styles(writer, df_txn)

                        # Optional line items sheet
                        if include_lines_var.get() and self.df_raw is not None:
                            txn_inv_numbers = set(
                                df_txn[df_txn.get("Type", pd.Series(dtype=str)) == "Invoice"]["Number"].dropna().tolist()
                            ) if "Number" in df_txn.columns else set()
                            raw_li = self.df_raw[self.df_raw["Number"].isin(txn_inv_numbers)].copy()
                            if not raw_li.empty:
                                raw_li = raw_li.copy()
                                raw_li["Date"] = raw_li["Date"].dt.strftime("%Y-%m-%d")
                                li_cols = [c for c in ["Date", "Number", "Account", "Partner", "Label",
                                                        "Debit", "Credit", "Salesperson"] if c in raw_li.columns]
                                raw_li[li_cols].sort_values(["Number", "Date"]).to_excel(
                                    writer, index=False, sheet_name="Line Items")
                                excel_format_2dp(writer.book["Line Items"])

                    messagebox.showinfo("Exported", f"Excel saved:\n{out_path}")
                except Exception as e:
                    messagebox.showerror("Export Failed", f"Excel export failed:\n\n{e}")
            else:
                out_path = Path(folder) / f"Sales_Report_{safe_name}_{range_tag}_{ts}.pdf"
                try:
                    export_pdf(sp, out_path, cn_summary_for_sp=cn_summary_for_sp)
                    messagebox.showinfo("Exported", f"PDF saved:\n{out_path}")
                except Exception as e:
                    messagebox.showerror("Export Failed", f"PDF export failed:\n\n{e}")

        ttk.Button(dlg, text="Export", command=do_export).pack(pady=16)

    # =========================
    # Salesperson listbox helpers
    # =========================
    def get_selected_salespeople(self) -> list:
        """Return list of currently selected salespeople from the listbox."""
        indices = self.sp_listbox.curselection()
        return [self.sp_listbox.get(i) for i in indices]

    def _update_sp_listbox(self, salespeople: list):
        """Populate/refresh the listbox; preserve existing selection if possible."""
        prev = set(self.get_selected_salespeople())
        self.sp_listbox.delete(0, tk.END)
        for sp in salespeople:
            self.sp_listbox.insert(tk.END, sp)
        # restore selection or default to first
        restored = False
        for i, sp in enumerate(salespeople):
            if sp in prev:
                self.sp_listbox.selection_set(i)
                restored = True
        if not restored and salespeople:
            self.sp_listbox.selection_set(0)
            self.salesperson_var.set(salespeople[0])

    # =========================
    # Product filter / autocomplete
    # =========================
    def _on_product_key(self, event=None):
        query = self.product_filter_var.get().strip().lower()
        if not query or self.df_raw is None:
            self._hide_product_ac()
            return
        inv_raw = self.df_raw[self.df_raw["Number"].str.startswith(INVOICE_PREFIX, na=False)]
        all_labels = inv_raw["Label"].dropna().unique().tolist()
        matches = sorted(set(lb for lb in all_labels if query in lb.lower() and lb.strip()))[:25]
        if not matches:
            self._hide_product_ac()
            return
        self._show_product_ac(matches)

    def _show_product_ac(self, matches: list):
        entry = self.product_filter_entry
        x = entry.winfo_rootx()
        y = entry.winfo_rooty() + entry.winfo_height()
        # create or reuse autocomplete window
        if self._ac_win is None or not self._ac_win.winfo_exists():
            self._ac_win = tk.Toplevel(self.root)
            self._ac_win.wm_overrideredirect(True)
            self._ac_win.attributes("-topmost", True)
            self._ac_lb = tk.Listbox(self._ac_win, selectmode="single",
                                     font=("Segoe UI", 9), width=70)
            self._ac_lb.pack(fill="both", expand=True)
            self._ac_lb.bind("<<ListboxSelect>>", self._on_product_ac_select)
            self._ac_lb.bind("<Return>", self._on_product_ac_select)
        h = min(10, len(matches))
        self._ac_win.geometry(f"+{x}+{y}")
        self._ac_lb.configure(height=h)
        self._ac_lb.delete(0, tk.END)
        for m in matches:
            self._ac_lb.insert(tk.END, m)
        self._ac_win.deiconify()
        self._ac_win.lift()

    def _hide_product_ac(self):
        if self._ac_win and self._ac_win.winfo_exists():
            self._ac_win.withdraw()

    def _on_product_ac_select(self, event=None):
        try:
            idx = self._ac_lb.curselection()
            if idx:
                val = self._ac_lb.get(idx[0])
                self.product_filter_var.set(val)
                self.product_filter_active = val
                self.product_active_lbl.config(text=f"Active: {val[:60]}{'…' if len(val)>60 else ''}")
                self._hide_product_ac()
                self.refresh_view()
        except Exception:
            pass

    def _apply_typed_product_filter(self):
        """Apply whatever text is currently typed — matches ALL products containing that text."""
        val = self.product_filter_var.get().strip()
        self._hide_product_ac()
        if not val:
            self._clear_product_filter()
            return
        self.product_filter_active = val
        short = val[:60] + ("…" if len(val) > 60 else "")
        self.product_active_lbl.config(text=f"Active: \"{short}\"  (all matching products)")
        self.refresh_view()

    def _clear_product_filter(self):
        self.product_filter_var.set("")
        self.product_filter_active = ""
        self.product_active_lbl.config(text="")
        self._hide_product_ac()
        self.refresh_view()

    # =========================
    # Invoice line items popup
    # =========================
    def show_line_items_popup(self):
        if self.df_raw is None:
            messagebox.showwarning("No data", "Load an Odoo file first.")
            return
        sel = self.tree_txn.selection()
        if not sel:
            messagebox.showinfo("Select a row", "Click an Invoice row in the Transactions tab, then press this button.")
            return
        vals = self.tree_txn.item(sel[0], "values")
        if not vals:
            return
        row_type = str(vals[0]).strip()
        inv_no = str(vals[2]).strip()
        if row_type != "Invoice" or not is_invoice_number_value(inv_no):
            messagebox.showinfo("Invoice rows only", "Please select an Invoice row (not a bill, CN, or total).")
            return
        lines = self.df_raw[self.df_raw["Number"] == inv_no].copy()
        if lines.empty:
            messagebox.showinfo("Not found", f"No raw line items found for {inv_no}.")
            return

        pop = tk.Toplevel(self.root)
        pop.title(f"Line Items  —  {inv_no}")
        pop.geometry("1150x380")
        pop.transient(self.root)

        ttk.Label(pop, text=f"Invoice: {inv_no}   |   {len(lines)} line(s)",
                  font=("Segoe UI", 10, "bold"), padding=8).pack(anchor="w")

        frm = ttk.Frame(pop)
        frm.pack(fill="both", expand=True, padx=10, pady=(0, 10))

        cols = ["Date", "Number", "Account", "Partner", "Label", "Debit", "Credit"]
        t2 = ttk.Treeview(frm, columns=cols, show="headings", height=14)
        for c in cols:
            t2.heading(c, text=c)
            if c == "Label":
                t2.column(c, width=380, anchor="w")
            elif c == "Account":
                t2.column(c, width=280, anchor="w")
            elif c == "Partner":
                t2.column(c, width=200, anchor="w")
            elif c == "Number":
                t2.column(c, width=160, anchor="w")
            elif c == "Date":
                t2.column(c, width=110, anchor="w")
            else:
                t2.column(c, width=120, anchor="e")

        vsb2 = ttk.Scrollbar(frm, orient="vertical", command=t2.yview)
        hsb2 = ttk.Scrollbar(frm, orient="horizontal", command=t2.xview)
        t2.configure(yscrollcommand=vsb2.set, xscrollcommand=hsb2.set)
        t2.pack(side="left", fill="both", expand=True)
        vsb2.pack(side="right", fill="y")
        hsb2.pack(side="bottom", fill="x")

        for _, r in lines.iterrows():
            dt = r["Date"]
            dt_s = "" if pd.isna(dt) else dt.strftime("%Y-%m-%d")
            t2.insert("", "end", values=(
                dt_s, safe_str(r["Number"]), safe_str(r["Account"]),
                safe_str(r["Partner"]), safe_str(r["Label"]),
                fmt_money(float(r["Debit"])), fmt_money(float(r["Credit"])),
            ))
        # Totals row
        t2.insert("", "end", values=("", "", "", "", "TOTAL",
            fmt_money(float(lines["Debit"].sum())),
            fmt_money(float(lines["Credit"].sum())),
        ))

    # =========================
    # Multi-SP combined view
    # =========================
    def _refresh_multi_sp(self, selected_sps: list):
        cn_all = self.build_creditnote_summary()
        grand_s = grand_c = grand_t = grand_l = grand_g = grand_e = grand_comm = grand_net = grand_inc = 0.0
        first_sp = True

        for sp in selected_sps:
            inv = self.invoice_summary[self.invoice_summary["Salesperson"].astype(str).str.strip() == sp].copy()
            bills = self.bill_lines[self.bill_lines["ExpensePerson"].astype(str).str.strip() == sp].copy()
            inv   = self._filter_by_range(inv, "Date")
            bills = self._filter_by_range(bills, "Date")

            # product filter
            if self.product_filter_active and self.df_raw is not None:
                mask = self.df_raw["Label"].str.contains(self.product_filter_active, case=False, na=False)
                matching_invs = set(
                    self.df_raw[mask & self.df_raw["Number"].apply(is_invoice_number_value)]["Number"].tolist()
                )
                inv = inv[inv["Number"].isin(matching_invs)]

            cn_sp = pd.DataFrame()
            if not cn_all.empty:
                cn_sp = cn_all[cn_all["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_sp["Date"] = pd.to_datetime(cn_sp["Date"], errors="coerce")
                cn_sp = cn_sp[~cn_sp["Date"].isna()]
                cn_sp = self._filter_by_range(cn_sp, "Date")

            s  = float(inv["Sales"].sum())        if not inv.empty else 0.0
            c  = float(inv["COGS"].sum())         if not inv.empty else 0.0
            t  = float(inv["Transport"].sum())    if not inv.empty else 0.0
            l_ = float(inv["Loading"].sum())      if not inv.empty else 0.0
            g  = float(inv["Adjusted GM"].sum())  if not inv.empty else 0.0
            if not cn_sp.empty:
                s += float(cn_sp["SalesEffect"].sum())
                c += float(cn_sp["COGSEffect"].sum())
                g += float(cn_sp["GMEffect"].sum())
            e    = float(bills["ExpenseAmount"].sum()) if not bills.empty else 0.0
            cb   = self.commission_monthly_breakdown(sp)
            comm = float(sum(cb.values())) if cb else 0.0
            net  = g - e - comm

            grand_s += s; grand_c += c; grand_t += t; grand_l += l_
            grand_g += g; grand_e += e; grand_comm += comm; grand_net += net

            E13 = ("","","","","","","","","","","","","")

            # ── Transactions tab ──
            if not first_sp:
                self.tree_txn.insert("", "end", values=E13)
            first_sp = False
            self.tree_txn.insert("", "end",
                values=(f"━━  {sp}  ━━","","","","","","","","","","","",""),
                tags=("SP_HEADER",))

            for _, r in inv.sort_values(["Date","Number"], na_position="last").iterrows():
                dt = r["Date"]; inv_date = dt.date() if not pd.isna(dt) else None
                date_s = "" if pd.isna(dt) else dt.strftime("%Y-%m-%d")
                cust = safe_str(r.get("Customer", r.get("Partner","")))
                sv = float(r.get("Sales",0.0)); gv = float(r.get("Adjusted GM",0.0))
                gp = safe_pct(gv, sv)
                tags = ()
                if inv_date:
                    if self.is_global_new_customer_month(cust, inv_date):      tags = ("GLOBAL_NEW",)
                    if self.is_qualified_new_customer_month(cust, inv_date):   tags = tuple(set(tags+("QUAL_NEW",)))
                    if self.is_qualified_repeat_invoice(cust, inv_date):       tags = ("QUAL_REPEAT",)
                self.tree_txn.insert("", "end", values=(
                    "Invoice", date_s, safe_str(r["Number"]), safe_str(r["Partner"]),
                    fmt_money(sv), fmt_money(float(r.get("COGS",0))),
                    fmt_money(float(r.get("Transport",0))), fmt_money(float(r.get("Loading",0))),
                    fmt_money(gv), f"{gp:.2f}", "", "", fmt_money(gv),
                ), tags=tags)

            if not cn_sp.empty:
                for _, r in cn_sp.sort_values(["Date","Number"],na_position="last").iterrows():
                    dt = r["Date"]; date_s = "" if pd.isna(dt) else pd.to_datetime(dt).strftime("%Y-%m-%d")
                    se = float(r.get("SalesEffect",0)); ge = float(r.get("GMEffect",0))
                    self.tree_txn.insert("", "end", values=(
                        "Credit Note", date_s, safe_str(r.get("Number","")), safe_str(r.get("Partner","")),
                        fmt_money(se), fmt_money(float(r.get("COGSEffect",0))),
                        "0.00","0.00", fmt_money(ge),
                        f"{safe_pct(ge,se):.2f}" if abs(se)>1e-9 else "0.00",
                        "","", fmt_money(ge),
                    ))

            for _, r in bills.sort_values(["Date","Number"],na_position="last").iterrows():
                date_s = "" if pd.isna(r["Date"]) else r["Date"].strftime("%Y-%m-%d")
                info = safe_str(r["Label"]) if safe_str(r["Label"]) else safe_str(r["Account"])
                amt = float(r["ExpenseAmount"])
                self.tree_txn.insert("", "end", values=(
                    "Emp Exp", date_s, safe_str(r["Number"]), info,
                    "","","","","","", fmt_money(-amt),"", fmt_money(-amt),
                ))

            for m_ in sorted(cb.keys()):
                v = float(cb[m_])
                if abs(v) > 1e-9:
                    self.tree_txn.insert("", "end", values=(
                        "Commission", m_, "", "Account 211810 (assigned)",
                        "","","","","","","", fmt_money(-v), fmt_money(-v),
                    ))

            self.tree_txn.insert("", "end", values=(
                f"  Subtotal — {sp}", "","", sp,
                fmt_money(s), fmt_money(c), fmt_money(t), fmt_money(l_),
                fmt_money(g), f"{safe_pct(g,s):.2f}",
                fmt_money(-e), fmt_money(-comm), fmt_money(net),
            ), tags=("SP_HEADER",))

            # ── Incentives tab ──
            cn_full = pd.DataFrame()
            if not cn_all.empty:
                cn_full = cn_all[cn_all["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_full["Date"] = pd.to_datetime(cn_full["Date"], errors="coerce")
                cn_full = cn_full[~cn_full["Date"].isna()]
            m_rows, inc_sum = self.compute_incentives(sp, cn_summary_for_sp=cn_full)
            sp_inc = float(inc_sum.get("total_paid",0.0)) if not inc_sum.get("note") else 0.0
            grand_inc += sp_inc

            self.tree_inc.insert("", "end",
                values=(f"━━  {sp}  ━━","","","","","","","","","",""),
                tags=("SP_HEADER",))
            for mr in m_rows:
                self.tree_inc.insert("", "end", values=(
                    mr["Month"], fmt_money(mr.get("Salary",0)), fmt_money(mr.get("Base",0)),
                    fmt_money(mr.get("Threshold",0)), fmt_money(mr.get("CarryIn",0)),
                    mr.get("Met",""), fmt_money(mr.get("Inc10",0)),
                    fmt_money(mr.get("NewBonus",0)), fmt_money(mr.get("RepeatBonus",0)),
                    fmt_money(mr.get("Total",0)), mr.get("Notice",""),
                ))

            # ── Commission Details tab ──
            ci = self.commission_line_items_in_range(sp)
            if not ci.empty:
                self.tree_comm.insert("", "end",
                    values=(f"━━  {sp}  ━━","","","","","","",""), tags=("SP_HEADER",))
                for _, r in ci.iterrows():
                    dt_s = "" if pd.isna(r["Date"]) else pd.to_datetime(r["Date"]).strftime("%Y-%m-%d")
                    self.tree_comm.insert("", "end", values=(
                        dt_s, safe_str(r["YearMonth"]), fmt_money(float(r["Amount"])),
                        safe_str(r["Partner"]), safe_str(r["Label"]),
                        safe_str(r["Number"]), safe_str(r["Account"]), safe_str(r["RowID"]),
                    ))

            # ── Credit Notes tab ──
            cni = self.creditnote_line_items_in_range(sp)
            if not cni.empty:
                self.tree_cn.insert("", "end",
                    values=(f"━━  {sp}  ━━","","","","","","","",""), tags=("SP_HEADER",))
                for _, r in cni.iterrows():
                    dt_s = "" if pd.isna(r["Date"]) else pd.to_datetime(r["Date"]).strftime("%Y-%m-%d")
                    self.tree_cn.insert("", "end", values=(
                        dt_s, safe_str(r["YearMonth"]), fmt_money(float(r["Amount"])),
                        safe_str(r["Partner"]), safe_str(r["Label"]),
                        safe_str(r["Number"]), safe_str(r["Account"]),
                        safe_str(r["RowID"]), safe_str(r["Salesperson"]),
                    ))

        # Grand total row in Transactions tab
        self.tree_txn.insert("", "end", values=E13)
        self.tree_txn.insert("", "end", values=(
            "GRAND TOTAL","","", f"{len(selected_sps)} salespeople",
            fmt_money(grand_s), fmt_money(grand_c), fmt_money(grand_t), fmt_money(grand_l),
            fmt_money(grand_g), f"{safe_pct(grand_g, grand_s):.2f}",
            fmt_money(-grand_e), fmt_money(-grand_comm), fmt_money(grand_net),
        ), tags=("SP_HEADER",))

        # Summary bar
        self.lbl_sales.config(text=f"Sales: {grand_s:,.2f}")
        self.lbl_cogs.config(text=f"COGS: {grand_c:,.2f}")
        self.lbl_chg.config(text=f"Charges (T+L): {(grand_t+grand_l):,.2f}")
        self.lbl_adjgm.config(text=f"Adjusted GM: {grand_g:,.2f}")
        self.lbl_exp.config(text=f"Employee Exp: {grand_e:,.2f}")
        self.lbl_comm.config(text=f"Commission: {grand_comm:,.2f}")
        self.lbl_net.config(text=f"Net: {grand_net:,.2f}")
        self.lbl_incent.config(text=f"Incentive (₹): {grand_inc:,.2f}")

        self.incent_text.delete("1.0","end")
        start, end = self.current_range
        range_label = "All Dates" if (start is None and end is None) else f"{start} to {end}"
        pf = f" | Product Filter: {self.product_filter_active}" if self.product_filter_active else ""
        self.incent_text.insert("end", f"Multi-SP View: {', '.join(selected_sps)}\n")
        self.incent_text.insert("end", f"Range: {range_label}{pf}\n")
        self.incent_text.insert("end", "Incentive tab shows monthly breakdown per salesperson (with carry-forward).\n")
        self.status_var.set(
            f"Multi-SP ({len(selected_sps)} selected){pf} | Range: {range_label} | "
            f"Qualified new customers: {len(self.qual_new_customers_global)}"
        )

    # =========================
    # Team Summary popup (keep alongside multi-select)
    # =========================
    def open_team_summary(self):
        if self.invoice_summary is None:
            messagebox.showwarning("No data", "Please load an Odoo Excel file first.")
            return
        if not self.allowed_salespeople:
            messagebox.showwarning("No salespeople", "No salaried salespeople found.")
            return

        sel_dlg = tk.Toplevel(self.root)
        sel_dlg.title("Select Salespeople for Team Summary")
        sel_dlg.geometry("340x420")
        sel_dlg.transient(self.root)
        sel_dlg.grab_set()

        ttk.Label(sel_dlg, text="Select salespeople to include:", padding=10).pack(anchor="w")

        all_var = tk.BooleanVar(value=True)
        check_vars = {}

        frm = ttk.Frame(sel_dlg, padding=(10, 0, 10, 0))
        frm.pack(fill="both", expand=True)
        canvas = tk.Canvas(frm, borderwidth=0)
        vsb = ttk.Scrollbar(frm, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        inner = ttk.Frame(canvas)
        canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))

        for sp in self.allowed_salespeople:
            v = tk.BooleanVar(value=True)
            check_vars[sp] = v
            ttk.Checkbutton(inner, text=sp, variable=v).pack(anchor="w", padx=6, pady=2)

        def toggle_all():
            for v in check_vars.values():
                v.set(all_var.get())

        ctrl = ttk.Frame(sel_dlg, padding=10)
        ctrl.pack(fill="x")
        ttk.Checkbutton(ctrl, text="Select / Deselect All", variable=all_var, command=toggle_all).pack(anchor="w")

        def show_report():
            selected = [sp for sp, v in check_vars.items() if v.get()]
            if not selected:
                messagebox.showwarning("Nothing selected", "Select at least one salesperson.")
                return
            sel_dlg.destroy()
            self._show_team_report(selected)

        ttk.Button(sel_dlg, text="Show Team Summary", command=show_report).pack(pady=10)

    def _show_team_report(self, salespeople: list):
        start, end = self.current_range
        range_label = "All Dates" if (start is None and end is None) else f"{start} to {end}"
        cn_all = self.build_creditnote_summary()

        rows = []
        for sp in salespeople:
            inv   = self.invoice_summary[self.invoice_summary["Salesperson"].astype(str).str.strip() == sp].copy()
            bills = self.bill_lines[self.bill_lines["ExpensePerson"].astype(str).str.strip() == sp].copy()
            inv   = self._filter_by_range(inv, "Date")
            bills = self._filter_by_range(bills, "Date")

            cn_sp = pd.DataFrame()
            if not cn_all.empty:
                cn_sp = cn_all[cn_all["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_sp["Date"] = pd.to_datetime(cn_sp["Date"], errors="coerce")
                cn_sp = cn_sp[~cn_sp["Date"].isna()]
                cn_sp = self._filter_by_range(cn_sp, "Date")

            s = float(inv["Sales"].sum()) if not inv.empty else 0.0
            c = float(inv["COGS"].sum())  if not inv.empty else 0.0
            t = float(inv["Transport"].sum()) if not inv.empty else 0.0
            l_= float(inv["Loading"].sum())   if not inv.empty else 0.0
            g = float(inv["Adjusted GM"].sum()) if not inv.empty else 0.0
            if not cn_sp.empty:
                s += float(cn_sp["SalesEffect"].sum())
                c += float(cn_sp["COGSEffect"].sum())
                g += float(cn_sp["GMEffect"].sum())
            e    = float(bills["ExpenseAmount"].sum()) if not bills.empty else 0.0
            cb   = self.commission_monthly_breakdown(sp)
            comm = float(sum(cb.values())) if cb else 0.0
            net  = g - e - comm

            cn_full = pd.DataFrame()
            if not cn_all.empty:
                cn_full = cn_all[cn_all["Salesperson"].astype(str).str.strip() == sp].copy()
                cn_full["Date"] = pd.to_datetime(cn_full["Date"], errors="coerce")
                cn_full = cn_full[~cn_full["Date"].isna()]
            _, inc_sum = self.compute_incentives(sp, cn_summary_for_sp=cn_full)
            incentive = float(inc_sum.get("total_paid",0.0)) if not inc_sum.get("note") else 0.0

            rows.append({"Salesperson":sp,"Sales":s,"COGS":c,"Transport":t,"Loading":l_,
                         "Adjusted GM":g,"GM %":safe_pct(g,s),"Emp Exp":e,
                         "Commission":comm,"Net":net,"Incentive":incentive})

        win = tk.Toplevel(self.root)
        win.title(f"Team Summary — {range_label}")
        win.geometry("1300x520")
        win.transient(self.root)

        ttk.Label(win, text=f"Team Summary  |  Range: {range_label}  |  {len(salespeople)} salespeople",
                  font=("Segoe UI", 11, "bold"), padding=10).pack(anchor="w")

        cols = ["Salesperson","Sales","COGS","Transport","Loading",
                "Adjusted GM","GM %","Emp Exp","Commission","Net","Incentive"]
        frm = ttk.Frame(win, padding=10)
        frm.pack(fill="both", expand=True)
        tree = ttk.Treeview(frm, columns=cols, show="headings", height=20)
        for c in cols:
            tree.heading(c, text=c)
            tree.column(c, width=200 if c=="Salesperson" else 80 if c=="GM %" else 120,
                        anchor="w" if c=="Salesperson" else "e")
        vsb2 = ttk.Scrollbar(frm, orient="vertical",  command=tree.yview)
        hsb2 = ttk.Scrollbar(frm, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vsb2.set, xscrollcommand=hsb2.set)
        tree.pack(side="left", fill="both", expand=True)
        vsb2.pack(side="right", fill="y")
        hsb2.pack(side="bottom", fill="x")
        tree.tag_configure("TOTAL", font=("Segoe UI", 9, "bold"))

        for r in rows:
            tree.insert("", "end", values=(
                r["Salesperson"], fmt_money(r["Sales"]), fmt_money(r["COGS"]),
                fmt_money(r["Transport"]), fmt_money(r["Loading"]),
                fmt_money(r["Adjusted GM"]), f"{r['GM %']:.2f}",
                fmt_money(r["Emp Exp"]), fmt_money(r["Commission"]),
                fmt_money(r["Net"]), fmt_money(r["Incentive"]),
            ))

        def _s(k): return sum(r[k] for r in rows)
        ts = _s("Sales")
        tree.insert("", "end", values=("","","","","","","","","","",""))
        tree.insert("", "end", values=(
            "GRAND TOTAL", fmt_money(ts), fmt_money(_s("COGS")),
            fmt_money(_s("Transport")), fmt_money(_s("Loading")),
            fmt_money(_s("Adjusted GM")), f"{safe_pct(_s('Adjusted GM'),ts):.2f}",
            fmt_money(_s("Emp Exp")), fmt_money(_s("Commission")),
            fmt_money(_s("Net")), fmt_money(_s("Incentive")),
        ), tags=("TOTAL",))

        def export_team_excel():
            folder = filedialog.askdirectory(title="Select folder to save Team Summary")
            if not folder: return
            ts2 = datetime.now().strftime("%Y%m%d_%H%M%S")
            out_path = Path(folder) / f"Team_Summary_{ts2}.xlsx"
            try:
                df_rows = [{k: round(v,2) if isinstance(v,float) else v
                            for k,v in r.items()} for r in rows]
                df_rows.append({})
                total_s = _s("Sales")
                df_rows.append({"Salesperson":"GRAND TOTAL",
                    "Sales":round(total_s,2),"COGS":round(_s("COGS"),2),
                    "Transport":round(_s("Transport"),2),"Loading":round(_s("Loading"),2),
                    "Adjusted GM":round(_s("Adjusted GM"),2),
                    "GM %":round(safe_pct(_s("Adjusted GM"),total_s),2),
                    "Emp Exp":round(_s("Emp Exp"),2),"Commission":round(_s("Commission"),2),
                    "Net":round(_s("Net"),2),"Incentive":round(_s("Incentive"),2)})
                df = pd.DataFrame(df_rows)

                # Incentive breakdown per SP
                inc_rows = []
                for sp in salespeople:
                    cn_f = pd.DataFrame()
                    if not cn_all.empty:
                        cn_f = cn_all[cn_all["Salesperson"].astype(str).str.strip()==sp].copy()
                        cn_f["Date"] = pd.to_datetime(cn_f["Date"], errors="coerce")
                        cn_f = cn_f[~cn_f["Date"].isna()]
                    mrows, _ = self.compute_incentives(sp, cn_summary_for_sp=cn_f)
                    for mr in mrows:
                        inc_rows.append({"Salesperson": sp, **mr})
                df_inc = pd.DataFrame(inc_rows)

                with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
                    df.to_excel(writer, index=False, sheet_name="Team Summary")
                    if not df_inc.empty:
                        df_inc.to_excel(writer, index=False, sheet_name="Incentive Breakdown")
                    excel_format_2dp(writer.book["Team Summary"])
                messagebox.showinfo("Exported", f"Saved:\n{out_path}")
            except Exception as ex:
                messagebox.showerror("Export Failed", str(ex))

        ttk.Button(win, text="📥 Export to Excel", command=export_team_excel).pack(pady=8)

    # =========================
    # Customer-wise Export
    # =========================
    def export_customer_report(self):
        if self.invoice_summary is None:
            messagebox.showwarning("No data", "Please load an Odoo Excel file first.")
            return
        folder = filedialog.askdirectory(title="Select folder to save Customer Report")
        if not folder: return

        cust_filter = self.customer_choice_var.get().strip()
        show_all = self.customer_all_only_var.get() or not cust_filter

        inv_all = self.invoice_summary.copy()
        inv_all = self._filter_by_range(inv_all, "Date")
        inv_all = inv_all[inv_all["Customer"].astype(str).str.strip() != ""].copy()
        if not show_all and cust_filter:
            inv_all = inv_all[inv_all["Customer"].astype(str).str.strip() == cust_filter].copy()
        if inv_all.empty:
            messagebox.showwarning("No data", "No customer data in selected range/filter.")
            return

        inv_all = inv_all.sort_values(["Customer","Date","Number"], na_position="last")
        rows = []
        for _, r in inv_all.iterrows():
            dt = r["Date"]; inv_date = dt.date() if not pd.isna(dt) else None
            cust = safe_str(r.get("Customer", r.get("Partner","")))
            sv = float(r.get("Sales",0.0)); gv = float(r.get("Adjusted GM",0.0))
            tag = ""
            if inv_date:
                if self.is_qualified_repeat_invoice(cust, inv_date):       tag = "Qualified Repeat"
                elif self.is_qualified_new_customer_month(cust, inv_date): tag = "Qualified New"
                elif self.is_global_new_customer_month(cust, inv_date):    tag = "New Customer"
            rows.append({"Customer":cust, "Salesperson":safe_str(r.get("Salesperson","")),
                "Date":"" if pd.isna(dt) else dt.strftime("%Y-%m-%d"),
                "Invoice No":safe_str(r.get("Number","")),
                "Sales":round(sv,2), "COGS":round(float(r.get("COGS",0)),2),
                "Transport":round(float(r.get("Transport",0)),2),
                "Loading":round(float(r.get("Loading",0)),2),
                "Adjusted GM":round(gv,2), "GM %":round(safe_pct(gv,sv),2), "Flag":tag})

        df = pd.DataFrame(rows)
        summary_rows = []
        for cust_name, grp in df.groupby("Customer", sort=True):
            s2 = float(grp["Sales"].sum()); g2 = float(grp["Adjusted GM"].sum())
            summary_rows.append({"Customer":cust_name,
                "Salesperson":grp["Salesperson"].iloc[0] if not grp.empty else "",
                "Invoices":len(grp), "Total Sales":round(s2,2),
                "Total COGS":round(float(grp["COGS"].sum()),2),
                "Total Adjusted GM":round(g2,2), "GM %":round(safe_pct(g2,s2),2),
                "Flag":grp["Flag"].iloc[0] if not grp.empty else ""})
        df_summary = pd.DataFrame(summary_rows)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        start, end = self.current_range
        range_tag = "ALL" if (start is None and end is None) else f"{start}_to_{end}"
        range_tag = range_tag.replace("-","")
        cust_tag = re.sub(r"[^\w\-]+","",cust_filter).strip()[:30] if (not show_all and cust_filter) else "ALL"
        out_path = Path(folder) / f"Customer_Report_{cust_tag}_{range_tag}_{ts}.xlsx"

        try:
            from openpyxl.styles import PatternFill
            with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
                df_summary.to_excel(writer, index=False, sheet_name="Customer Summary")
                df.to_excel(writer, index=False, sheet_name="Invoice Detail")
                excel_format_2dp(writer.book["Customer Summary"])
                excel_format_2dp(writer.book["Invoice Detail"])
                ws = writer.book["Invoice Detail"]
                headers = [c.value for c in ws[1]]
                if "Flag" in headers:
                    fcol = headers.index("Flag") + 1
                    green_fill = PatternFill("solid", fgColor="C6EFCE")
                    blue_fill  = PatternFill("solid", fgColor="BDD7EE")
                    for ri in range(2, ws.max_row + 1):
                        fv = ws.cell(row=ri, column=fcol).value or ""
                        fill = green_fill if "New" in fv else blue_fill if "Repeat" in fv else None
                        if fill:
                            for ci in range(1, ws.max_column + 1):
                                ws.cell(row=ri, column=ci).fill = fill
            messagebox.showinfo("Exported", f"Customer report saved:\n{out_path}")
        except Exception as ex:
            messagebox.showerror("Export Failed", str(ex))

    # =========================
    # Salaries editor (EffectiveFrom)
    # =========================
    def open_salary_file_external(self):
        try:
            self.ensure_salary_file_exists()
            os.startfile(SALARY_FILE)
        except Exception as e:
            messagebox.showerror("Error", f"Could not open file:\n{e}")

    def open_salary_editor(self):
        self.ensure_salary_file_exists()
        try:
            df = pd.read_excel(SALARY_FILE)
        except Exception:
            df = pd.DataFrame(columns=["Salesperson", "EffectiveFrom", "MonthlySalary"])

        df.columns = [str(c).strip() for c in df.columns]
        for c in ["Salesperson", "EffectiveFrom", "MonthlySalary"]:
            if c not in df.columns:
                df[c] = ""
        df = df[["Salesperson", "EffectiveFrom", "MonthlySalary"]].copy()

        dlg = tk.Toplevel(self.root)
        dlg.title("Edit Salaries (EffectiveFrom)")
        dlg.geometry("860x560")
        dlg.transient(self.root)
        dlg.grab_set()

        ttk.Label(dlg, text=f"Salaries file: {SALARY_FILE}", padding=10).pack(anchor="w")

        frame = ttk.Frame(dlg, padding=10)
        frame.pack(fill="both", expand=True)

        cols = ["Salesperson", "EffectiveFrom", "MonthlySalary"]
        tree = ttk.Treeview(frame, columns=cols, show="headings", height=16)
        for c in cols:
            tree.heading(c, text=c)
            tree.column(c, width=280 if c == "Salesperson" else 200, anchor="w" if c != "MonthlySalary" else "e")
        tree.pack(side="left", fill="both", expand=True)

        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")

        for _, r in df.fillna("").iterrows():
            tree.insert("", "end", values=(safe_str(r["Salesperson"]), safe_str(r["EffectiveFrom"]), safe_str(r["MonthlySalary"])))

        controls = ttk.Frame(dlg, padding=10)
        controls.pack(fill="x")

        sp_var = tk.StringVar()
        eff_var = tk.StringVar(value=SCHEME_START_DATE.isoformat())
        sal_var = tk.StringVar()

        ttk.Label(controls, text="Salesperson").grid(row=0, column=0, sticky="w")
        ttk.Entry(controls, textvariable=sp_var, width=22).grid(row=0, column=1, padx=6)
        ttk.Label(controls, text="EffectiveFrom (YYYY-MM-DD)").grid(row=0, column=2, sticky="w")
        ttk.Entry(controls, textvariable=eff_var, width=18).grid(row=0, column=3, padx=6)
        ttk.Label(controls, text="MonthlySalary").grid(row=0, column=4, sticky="w")
        ttk.Entry(controls, textvariable=sal_var, width=14).grid(row=0, column=5, padx=6)

        def add_row():
            spn = sp_var.get().strip()
            eff = eff_var.get().strip()
            sal = sal_var.get().strip()
            if not spn:
                return
            d = parse_date_entry(eff)
            if not d:
                messagebox.showerror("Invalid", "EffectiveFrom must be YYYY-MM-DD")
                return
            try:
                sval = float(sal.replace(",", ""))
                if sval <= 0:
                    raise ValueError()
            except Exception:
                messagebox.showerror("Invalid", "MonthlySalary must be a positive number")
                return

            tree.insert("", "end", values=(spn, d.isoformat(), f"{sval:.2f}"))
            sp_var.set("")
            sal_var.set("")

        def delete_selected():
            for iid in tree.selection():
                tree.delete(iid)

        def save_file():
            rows = []
            for iid in tree.get_children():
                v = tree.item(iid, "values")
                spn = safe_str(v[0])
                eff = safe_str(v[1])
                sal = to_num(v[2])
                if spn:
                    rows.append({"Salesperson": spn, "EffectiveFrom": eff, "MonthlySalary": sal})
            out = pd.DataFrame(rows, columns=["Salesperson","EffectiveFrom","MonthlySalary"])
            saved_path, warn = safe_excel_save(out, SALARY_FILE)
            if warn:
                messagebox.showwarning("Salaries Saved (Fallback)", warn)

            self.load_salaries()
            dlg.destroy()
            self.refresh_view()
            messagebox.showinfo("Saved", f"Salaries saved:\n{saved_path}")

        ttk.Button(controls, text="Add", command=add_row).grid(row=1, column=1, pady=10, sticky="w")
        ttk.Button(controls, text="Delete Selected", command=delete_selected).grid(row=1, column=2, pady=10, sticky="w")
        ttk.Button(controls, text="Save Salaries File", command=save_file).grid(row=1, column=5, pady=10, sticky="e")



# =========================
# Local SQLite DB / Security / Editable Settings Enhancements
# =========================
_DB_ORIGINAL_INIT = OdooSalesMarginApp.__init__
_DB_ORIGINAL_LOAD_AND_PREPARE_ODOO = OdooSalesMarginApp.load_and_prepare_odoo
_DB_ORIGINAL_LOAD_CHARGES_FILE = OdooSalesMarginApp.load_charges_file
_DB_ORIGINAL_LOAD_SALARIES = OdooSalesMarginApp.load_salaries
_DB_ORIGINAL_SAVE_COMMISSIONS_DB = OdooSalesMarginApp.save_commissions_db
_DB_ORIGINAL_LOAD_COMMISSIONS_DB = OdooSalesMarginApp.load_commissions_db
_DB_ORIGINAL_SAVE_CREDITNOTES_DB = OdooSalesMarginApp.save_creditnotes_db
_DB_ORIGINAL_LOAD_CREDITNOTES_DB = OdooSalesMarginApp.load_creditnotes_db
_DB_ORIGINAL_COMPUTE_INCENTIVES = OdooSalesMarginApp.compute_incentives


def _db_conn():
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(LOCAL_DB_FILE)
    con.execute("PRAGMA foreign_keys=ON")
    return con


def _password_hash(password: str, salt: str | None = None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.sha256((salt + "::" + password).encode("utf-8")).hexdigest()
    return salt, digest


def _db_init_schema():
    with _db_conn() as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS ledger_rows (
            row_id TEXT PRIMARY KEY,
            date TEXT, number TEXT, account TEXT, partner TEXT, label TEXT,
            debit REAL DEFAULT 0, credit REAL DEFAULT 0, salesperson TEXT,
            imported_at TEXT DEFAULT CURRENT_TIMESTAMP, source_file TEXT
        );
        CREATE TABLE IF NOT EXISTS charges (
            invoice_no TEXT PRIMARY KEY,
            transport REAL DEFAULT 0,
            loading REAL DEFAULT 0,
            source_file TEXT,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS salaries (
            salesperson TEXT NOT NULL,
            effective_from TEXT NOT NULL,
            monthly_salary REAL NOT NULL,
            PRIMARY KEY (salesperson, effective_from)
        );
        CREATE TABLE IF NOT EXISTS commissions (
            row_id TEXT PRIMARY KEY, date TEXT, yearmonth TEXT, amount REAL,
            partner TEXT, label TEXT, number TEXT, account TEXT, salesperson TEXT
        );
        CREATE TABLE IF NOT EXISTS credit_notes (
            row_id TEXT PRIMARY KEY, date TEXT, yearmonth TEXT, amount REAL,
            partner TEXT, label TEXT, number TEXT, account TEXT, salesperson TEXT
        );
        CREATE TABLE IF NOT EXISTS manual_adjustments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            salesperson TEXT NOT NULL,
            yearmonth TEXT NOT NULL,
            amount REAL NOT NULL,
            purpose TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        CREATE TABLE IF NOT EXISTS incentive_payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            salesperson TEXT NOT NULL,
            yearmonth TEXT NOT NULL,
            amount REAL NOT NULL,
            paid_on TEXT NOT NULL,
            mode TEXT,
            reference TEXT,
            notes TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE (salesperson, yearmonth)
        );
        """)
        # default password: admin123, user should change inside app
        if not con.execute("SELECT 1 FROM app_settings WHERE key='password_hash'").fetchone():
            salt, digest = _password_hash("admin123")
            con.execute("INSERT OR REPLACE INTO app_settings(key,value) VALUES('password_salt',?)", (salt,))
            con.execute("INSERT OR REPLACE INTO app_settings(key,value) VALUES('password_hash',?)", (digest,))
        defaults = {
            "MIN_NEW_CUST_COUNT": str(MIN_NEW_CUST_COUNT),
            "NEW_CUST_BILL_MIN": str(NEW_CUST_BILL_MIN),
            "NEW_CUST_MIN_MARGIN_PCT": str(NEW_CUST_MIN_MARGIN_PCT),
            "NEW_CUST_BONUS": str(NEW_CUST_BONUS),
            "REPEAT_BONUS": str(REPEAT_BONUS),
            "REPEAT_WINDOW_DAYS": str(REPEAT_WINDOW_DAYS),
            "SCHEME_START_DATE": str(SCHEME_START_DATE),
            "THRESHOLD_SALARY_MULTIPLE": "6.0",
            "INCENTIVE_PERCENTAGE": "10.0",
        }
        for k, v in defaults.items():
            con.execute("INSERT OR IGNORE INTO app_settings(key,value) VALUES(?,?)", (k, v))


def _db_get_setting(key, default=""):
    with _db_conn() as con:
        row = con.execute("SELECT value FROM app_settings WHERE key=?", (key,)).fetchone()
    return row[0] if row else default


def _db_set_setting(key, value):
    with _db_conn() as con:
        con.execute("INSERT OR REPLACE INTO app_settings(key,value) VALUES(?,?)", (key, str(value)))


def _db_load_scheme_globals():
    global MIN_NEW_CUST_COUNT, NEW_CUST_BILL_MIN, NEW_CUST_MIN_MARGIN_PCT, NEW_CUST_BONUS, REPEAT_BONUS, REPEAT_WINDOW_DAYS, SCHEME_START_DATE
    try:
        MIN_NEW_CUST_COUNT = int(float(_db_get_setting("MIN_NEW_CUST_COUNT", MIN_NEW_CUST_COUNT)))
        NEW_CUST_BILL_MIN = float(_db_get_setting("NEW_CUST_BILL_MIN", NEW_CUST_BILL_MIN))
        NEW_CUST_MIN_MARGIN_PCT = float(_db_get_setting("NEW_CUST_MIN_MARGIN_PCT", NEW_CUST_MIN_MARGIN_PCT))
        NEW_CUST_BONUS = float(_db_get_setting("NEW_CUST_BONUS", NEW_CUST_BONUS))
        REPEAT_BONUS = float(_db_get_setting("REPEAT_BONUS", REPEAT_BONUS))
        REPEAT_WINDOW_DAYS = int(float(_db_get_setting("REPEAT_WINDOW_DAYS", REPEAT_WINDOW_DAYS)))
        SCHEME_START_DATE = pd.to_datetime(_db_get_setting("SCHEME_START_DATE", SCHEME_START_DATE)).date()
    except Exception:
        pass


def _row_id_from_odoo_row(r):
    d = pd.to_datetime(r.get("Date"), errors="coerce")
    d_s = "" if pd.isna(d) else d.strftime("%Y-%m-%d")
    return make_row_id([
        safe_str(r.get("Number","")), safe_str(r.get("Account","")), safe_str(r.get("Partner","")),
        safe_str(r.get("Label","")), d_s, f"{to_num(r.get('Debit',0)):.2f}",
        f"{to_num(r.get('Credit',0)):.2f}", safe_str(r.get("Sales Order Lines/Salesperson", r.get("Salesperson","")))
    ])


def _db_import_ledger_excel(file_path):
    df = pd.read_excel(file_path)
    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in REQ_COLS if c not in df.columns]
    if missing:
        raise ValueError("Missing required columns:\n- " + "\n- ".join(missing))
    df = df.copy()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Debit"] = pd.to_numeric(df["Debit"], errors="coerce").fillna(0.0)
    df["Credit"] = pd.to_numeric(df["Credit"], errors="coerce").fillna(0.0)
    inserted = 0
    with _db_conn() as con:
        for _, r in df.iterrows():
            rid = _row_id_from_odoo_row(r)
            cur = con.execute("""
                INSERT OR IGNORE INTO ledger_rows
                (row_id,date,number,account,partner,label,debit,credit,salesperson,source_file)
                VALUES (?,?,?,?,?,?,?,?,?,?)
            """, (
                rid,
                "" if pd.isna(r["Date"]) else r["Date"].strftime("%Y-%m-%d"),
                safe_str(r["Number"]), safe_str(r["Account"]), safe_str(r["Partner"]), safe_str(r["Label"]),
                float(r["Debit"]), float(r["Credit"]), safe_str(r["Sales Order Lines/Salesperson"]),
                str(file_path)
            ))
            inserted += cur.rowcount
    return inserted, len(df)


def _db_ledger_to_excel_temp():
    with _db_conn() as con:
        rows = con.execute("""
            SELECT date as Date, number as Number, account as Account, partner as Partner, label as Label,
                   debit as Debit, credit as Credit, salesperson as [Sales Order Lines/Salesperson]
            FROM ledger_rows
            ORDER BY date, number, row_id
        """).fetchall()
        cols = ["Date","Number","Account","Partner","Label","Debit","Credit","Sales Order Lines/Salesperson"]
    df = pd.DataFrame(rows, columns=cols)
    if df.empty:
        raise ValueError("Local database has no Odoo rows yet.")
    tmp = APP_DATA_DIR / "_temp_all_odoo_rows_for_processing.xlsx"
    df.to_excel(tmp, index=False)
    return str(tmp)


def _db_init_app(self, root):
    _db_init_schema()
    _db_load_scheme_globals()
    self.local_db_path = LOCAL_DB_FILE
    _DB_ORIGINAL_INIT(self, root)
    try:
        self.status_var.set(f"Ready. Local database: {LOCAL_DB_FILE}")
    except Exception:
        pass


def _db_load_and_prepare_odoo(self, file_path):
    inserted, total = _db_import_ledger_excel(file_path)
    tmp = _db_ledger_to_excel_temp()
    _DB_ORIGINAL_LOAD_AND_PREPARE_ODOO(self, tmp)
    try:
        self.status_var.set(f"Odoo import complete: {inserted} fresh rows added, {total - inserted} duplicates ignored. Database retained all history.")
    except Exception:
        pass


def _db_load_charges_file(self, file_path):
    df = pd.read_excel(file_path)
    df.columns = [str(c).strip() for c in df.columns]
    inv_col = next((c for c in ["Invoice No", "Invoice", "Invoice Number", "InvoiceNo", "Number"] if c in df.columns), None)
    if not inv_col:
        raise ValueError("Could not find Invoice/Number column in charges file.")
    t_col = next((c for c in df.columns if "transport" in c.lower()), None)
    l_col = next((c for c in df.columns if "loading" in c.lower()), None)
    inserted = updated = 0
    with _db_conn() as con:
        for _, r in df.iterrows():
            inv = safe_str(r.get(inv_col, ""))
            if not inv:
                continue
            transport = to_num(r.get(t_col, 0.0)) if t_col else 0.0
            loading = to_num(r.get(l_col, 0.0)) if l_col else 0.0
            exists = con.execute("SELECT 1 FROM charges WHERE invoice_no=?", (inv,)).fetchone()
            con.execute("""
                INSERT INTO charges(invoice_no,transport,loading,source_file,updated_at)
                VALUES(?,?,?,?,CURRENT_TIMESTAMP)
                ON CONFLICT(invoice_no) DO UPDATE SET
                    transport=excluded.transport, loading=excluded.loading,
                    source_file=excluded.source_file, updated_at=CURRENT_TIMESTAMP
            """, (inv, transport, loading, str(file_path)))
            inserted += 0 if exists else 1
            updated += 1 if exists else 0
    with _db_conn() as con:
        all_rows = con.execute("SELECT invoice_no, transport, loading FROM charges").fetchall()
    self.charges_map = {r[0]: {"transport": float(r[1] or 0), "loading": float(r[2] or 0)} for r in all_rows}
    if self.invoice_summary is not None:
        self.apply_charges_to_invoices()
        self.rebuild_global_customer_maps()
    try:
        messagebox.showinfo("Charges Saved", f"Charges database updated.\nFresh invoices: {inserted}\nExisting invoices updated/kept: {updated}")
    except Exception:
        pass


def _db_load_salaries(self):
    # First migrate legacy Excel into SQLite if DB has no salaries.
    _db_init_schema()
    with _db_conn() as con:
        has_db = con.execute("SELECT 1 FROM salaries LIMIT 1").fetchone()
    if not has_db and os.path.exists(SALARY_FILE):
        try:
            df = pd.read_excel(SALARY_FILE)
            df.columns = [str(c).strip() for c in df.columns]
            if "Salesperson" in df.columns:
                if "EffectiveFrom" not in df.columns:
                    df["EffectiveFrom"] = "1900-01-01"
                if "MonthlySalary" not in df.columns:
                    sal_cand = next((c for c in df.columns if "salary" in c.lower()), None)
                    if sal_cand:
                        df = df.rename(columns={sal_cand: "MonthlySalary"})
                if "MonthlySalary" in df.columns:
                    with _db_conn() as con:
                        for _, r in df.iterrows():
                            sp = safe_str(r.get("Salesperson",""))
                            sal = to_num(r.get("MonthlySalary",0))
                            eff = pd.to_datetime(r.get("EffectiveFrom","1900-01-01"), errors="coerce")
                            if sp and sal > 0:
                                con.execute("INSERT OR REPLACE INTO salaries VALUES(?,?,?)", (sp, eff.strftime("%Y-%m-%d") if not pd.isna(eff) else "1900-01-01", sal))
        except Exception:
            pass

    with _db_conn() as con:
        rows = con.execute("SELECT salesperson,effective_from,monthly_salary FROM salaries ORDER BY salesperson,effective_from").fetchall()
    if rows:
        steps = {}
        for sp, eff_s, sal in rows:
            try:
                eff = pd.to_datetime(eff_s).date()
            except Exception:
                eff = date(1900,1,1)
            steps.setdefault(sp, []).append((eff, float(sal or 0)))
        self.salary_steps = steps
        self.allowed_salespeople = sorted(steps.keys())
        self._update_sp_listbox(self.allowed_salespeople)
        # Keep legacy Excel mirror for audit/backward compatibility.
        try:
            out = pd.DataFrame(rows, columns=["Salesperson","EffectiveFrom","MonthlySalary"])
            safe_excel_save(out, SALARY_FILE)
        except Exception:
            pass
        return
    return _DB_ORIGINAL_LOAD_SALARIES(self)


def _df_to_db_table(df, table):
    if df is None:
        return
    df = df.copy()
    with _db_conn() as con:
        con.execute(f"DELETE FROM {table}")
        for _, r in df.iterrows():
            con.execute(f"""INSERT OR REPLACE INTO {table}
            (row_id,date,yearmonth,amount,partner,label,number,account,salesperson)
            VALUES(?,?,?,?,?,?,?,?,?)""", (
                safe_str(r.get("RowID","")),
                "" if pd.isna(pd.to_datetime(r.get("Date"), errors="coerce")) else pd.to_datetime(r.get("Date")).strftime("%Y-%m-%d"),
                safe_str(r.get("YearMonth","")),
                to_num(r.get("Amount",0)),
                safe_str(r.get("Partner","")), safe_str(r.get("Label","")),
                safe_str(r.get("Number","")), safe_str(r.get("Account","")), safe_str(r.get("Salesperson",""))
            ))


def _db_table_to_df(table):
    with _db_conn() as con:
        rows = con.execute(f"SELECT row_id,date,yearmonth,amount,partner,label,number,account,salesperson FROM {table}").fetchall()
    df = pd.DataFrame(rows, columns=["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"])
    if not df.empty:
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        df["Amount"] = df["Amount"].apply(to_num)
    return df


def _db_load_commissions(self):
    _DB_ORIGINAL_LOAD_COMMISSIONS_DB(self)
    dbdf = _db_table_to_df("commissions")
    if not dbdf.empty:
        self.commission_db = dbdf
    elif self.commission_db is not None and not self.commission_db.empty:
        _df_to_db_table(self.commission_db, "commissions")


def _db_save_commissions(self):
    _DB_ORIGINAL_SAVE_COMMISSIONS_DB(self)
    _df_to_db_table(self.commission_db, "commissions")


def _db_load_creditnotes(self):
    _DB_ORIGINAL_LOAD_CREDITNOTES_DB(self)
    dbdf = _db_table_to_df("credit_notes")
    if not dbdf.empty:
        self.creditnote_db = dbdf
    elif self.creditnote_db is not None and not self.creditnote_db.empty:
        _df_to_db_table(self.creditnote_db, "credit_notes")


def _db_save_creditnotes(self):
    _DB_ORIGINAL_SAVE_CREDITNOTES_DB(self)
    _df_to_db_table(self.creditnote_db, "credit_notes")


def _get_manual_adjustments(sp):
    with _db_conn() as con:
        rows = con.execute("SELECT yearmonth, SUM(amount) FROM manual_adjustments WHERE salesperson=? GROUP BY yearmonth", (sp,)).fetchall()
    return {ym: float(a or 0) for ym, a in rows}


def _db_compute_incentives(self, sp, cn_summary_for_sp):
    rows, summary = _DB_ORIGINAL_COMPUTE_INCENTIVES(self, sp, cn_summary_for_sp)
    if not rows or summary.get("note"):
        return rows, summary
    adj = _get_manual_adjustments(sp)
    total_adj = 0.0
    for r in rows:
        m = r.get("Month")
        a = float(adj.get(m, 0.0))
        r["ManualAdjustment"] = a
        if abs(a) > 1e-9:
            r["Total"] = float(r.get("Total", 0.0)) + a
            r["Notice"] = (safe_str(r.get("Notice","")) + (" | " if safe_str(r.get("Notice","")) else "") + f"Manual {'+' if a>=0 else ''}{a:,.2f}").strip()
            total_adj += a
    summary["manual_adjustments"] = total_adj
    summary["total_paid"] = float(summary.get("total_paid", 0.0)) + total_adj
    return rows, summary


def _selected_salespeople(self):
    try:
        sel = [self.sp_listbox.get(i) for i in self.sp_listbox.curselection()]
        return [s for s in sel if safe_str(s)]
    except Exception:
        return []


def open_scheme_settings(self):
    _db_load_scheme_globals()
    dlg = tk.Toplevel(self.root); dlg.title("Editable Incentive Logic"); dlg.geometry("560x470"); dlg.transient(self.root); dlg.grab_set()
    fields = [
        ("Scheme Start Date (YYYY-MM-DD)", "SCHEME_START_DATE"),
        ("Threshold Salary Multiple", "THRESHOLD_SALARY_MULTIPLE"),
        ("Incentive Percentage", "INCENTIVE_PERCENTAGE"),
        ("Min New Customer Count", "MIN_NEW_CUST_COUNT"),
        ("New Customer Bill Minimum", "NEW_CUST_BILL_MIN"),
        ("New Customer Min GM %", "NEW_CUST_MIN_MARGIN_PCT"),
        ("New Customer Bonus", "NEW_CUST_BONUS"),
        ("Repeat Bonus", "REPEAT_BONUS"),
        ("Repeat Window Days", "REPEAT_WINDOW_DAYS"),
    ]
    vars_ = {}
    frm = ttk.Frame(dlg, padding=16); frm.pack(fill="both", expand=True)
    ttk.Label(frm, text="These settings are saved in the local database and applied immediately.", font=("Segoe UI", 10, "bold")).grid(row=0,column=0,columnspan=2,sticky="w",pady=(0,12))
    for i,(label,key) in enumerate(fields, start=1):
        ttk.Label(frm, text=label).grid(row=i,column=0,sticky="w",pady=5)
        v = tk.StringVar(value=_db_get_setting(key, ""))
        vars_[key]=v
        ttk.Entry(frm, textvariable=v, width=24).grid(row=i,column=1,sticky="w",pady=5)
    def save():
        for k,v in vars_.items():
            _db_set_setting(k, v.get().strip())
        _db_load_scheme_globals()
        try:
            self.rebuild_global_customer_maps()
            self.refresh_view()
        except Exception:
            pass
        dlg.destroy()
        messagebox.showinfo("Saved", "Incentive logic saved to local database.")
    ttk.Button(frm, text="Save Settings", command=save).grid(row=len(fields)+1,column=1,sticky="e",pady=18)


def open_manual_adjustments(self):
    dlg = tk.Toplevel(self.root); dlg.title("Manual Incentive Add/Deduct"); dlg.geometry("760x470"); dlg.transient(self.root); dlg.grab_set()
    frm = ttk.Frame(dlg, padding=12); frm.pack(fill="both", expand=True)
    spv = tk.StringVar(value=(_selected_salespeople(self)[0] if _selected_salespeople(self) else (self.allowed_salespeople[0] if self.allowed_salespeople else "")))
    ymv = tk.StringVar(value=self.month_choice_var.get() or datetime.now().strftime("%Y-%m"))
    amtv = tk.StringVar(value="0")
    purpv = tk.StringVar(value="")
    for row, (lbl, var) in enumerate([("Salesperson", spv), ("YearMonth", ymv), ("Amount (+ add / - deduct)", amtv), ("Purpose", purpv)]):
        ttk.Label(frm, text=lbl).grid(row=row,column=0,sticky="w",pady=5)
        if lbl=="Salesperson":
            ttk.Combobox(frm, textvariable=var, values=self.allowed_salespeople, width=32, state="readonly").grid(row=row,column=1,sticky="w",pady=5)
        else:
            ttk.Entry(frm, textvariable=var, width=50).grid(row=row,column=1,sticky="w",pady=5)
    cols=("ID","Salesperson","YearMonth","Amount","Purpose","Created")
    tree=ttk.Treeview(frm, columns=cols, show="headings", height=12)
    for c in cols:
        tree.heading(c,text=c); tree.column(c,width=110 if c!="Purpose" else 280, anchor="w")
    tree.grid(row=5,column=0,columnspan=3,sticky="nsew",pady=(12,6))
    frm.rowconfigure(5, weight=1); frm.columnconfigure(1, weight=1)
    def reload():
        for i in tree.get_children(): tree.delete(i)
        with _db_conn() as con:
            rows=con.execute("SELECT id,salesperson,yearmonth,amount,purpose,created_at FROM manual_adjustments ORDER BY yearmonth DESC, salesperson").fetchall()
        for r in rows: tree.insert("", "end", values=(r[0],r[1],r[2],fmt_money(r[3]),r[4],r[5]))
    def add():
        if not spv.get().strip() or not re.match(r"^\d{4}-\d{2}$", ymv.get().strip()):
            messagebox.showwarning("Missing", "Select salesperson and enter YearMonth as YYYY-MM."); return
        with _db_conn() as con:
            con.execute("INSERT INTO manual_adjustments(salesperson,yearmonth,amount,purpose) VALUES(?,?,?,?)", (spv.get().strip(), ymv.get().strip(), to_num(amtv.get()), purpv.get().strip() or "Manual adjustment"))
        reload(); self.refresh_view()
    def delete():
        for iid in tree.selection():
            rid=tree.item(iid,"values")[0]
            with _db_conn() as con: con.execute("DELETE FROM manual_adjustments WHERE id=?", (rid,))
        reload(); self.refresh_view()
    btn=ttk.Frame(frm); btn.grid(row=4,column=0,columnspan=3,sticky="w",pady=8)
    ttk.Button(btn,text="Add Adjustment",command=add).pack(side="left")
    ttk.Button(btn,text="Delete Selected",command=delete).pack(side="left",padx=8)
    reload()


def open_payment_recorder(self):
    sp_list = _selected_salespeople(self)
    sp = sp_list[0] if sp_list else (self.allowed_salespeople[0] if self.allowed_salespeople else "")
    dlg = tk.Toplevel(self.root); dlg.title("Record Final Incentive Paid"); dlg.geometry("620x360"); dlg.transient(self.root); dlg.grab_set()
    frm=ttk.Frame(dlg,padding=14); frm.pack(fill="both",expand=True)
    spv=tk.StringVar(value=sp); ymv=tk.StringVar(value=self.month_choice_var.get() or datetime.now().strftime("%Y-%m"))
    amtv=tk.StringVar(value=""); paidv=tk.StringVar(value=date.today().strftime("%Y-%m-%d")); modev=tk.StringVar(value=""); refv=tk.StringVar(value=""); notesv=tk.StringVar(value="")
    fields=[("Salesperson",spv),("YearMonth",ymv),("Paid Amount",amtv),("Paid On",paidv),("Mode",modev),("Reference",refv),("Notes",notesv)]
    for i,(lbl,var) in enumerate(fields):
        ttk.Label(frm,text=lbl).grid(row=i,column=0,sticky="w",pady=5)
        if lbl=="Salesperson":
            ttk.Combobox(frm,textvariable=var,values=self.allowed_salespeople,state="readonly",width=32).grid(row=i,column=1,sticky="w",pady=5)
        else:
            ttk.Entry(frm,textvariable=var,width=42).grid(row=i,column=1,sticky="w",pady=5)
    def compute_amt():
        if self.invoice_summary is None: return
        cn_all=self.build_creditnote_summary()
        cn_sp=cn_all[cn_all["Salesperson"].astype(str).str.strip()==spv.get().strip()].copy() if not cn_all.empty else pd.DataFrame()
        rows, summary=self.compute_incentives(spv.get().strip(), cn_sp)
        amount=sum(float(r.get("Total",0)) for r in rows if r.get("Month")==ymv.get().strip())
        amtv.set(f"{amount:.2f}")
    def save():
        if not spv.get().strip() or not re.match(r"^\d{4}-\d{2}$", ymv.get().strip()):
            messagebox.showwarning("Missing", "Select salesperson and YearMonth."); return
        with _db_conn() as con:
            con.execute("""INSERT INTO incentive_payments(salesperson,yearmonth,amount,paid_on,mode,reference,notes)
                           VALUES(?,?,?,?,?,?,?)
                           ON CONFLICT(salesperson,yearmonth) DO UPDATE SET amount=excluded.amount, paid_on=excluded.paid_on,
                           mode=excluded.mode, reference=excluded.reference, notes=excluded.notes""",
                        (spv.get().strip(), ymv.get().strip(), to_num(amtv.get()), paidv.get().strip(), modev.get().strip(), refv.get().strip(), notesv.get().strip()))
        dlg.destroy(); messagebox.showinfo("Saved", "Final incentive payment recorded in local database.")
    ttk.Button(frm,text="Auto-fill Current Payable",command=compute_amt).grid(row=8,column=0,sticky="w",pady=15)
    ttk.Button(frm,text="Save Payment Record",command=save).grid(row=8,column=1,sticky="e",pady=15)


def open_change_password(self):
    dlg=tk.Toplevel(self.root); dlg.title("Change Password"); dlg.geometry("430x230"); dlg.transient(self.root); dlg.grab_set()
    frm=ttk.Frame(dlg,padding=16); frm.pack(fill="both",expand=True)
    old=tk.StringVar(); new=tk.StringVar(); new2=tk.StringVar()
    for i,(lbl,var) in enumerate([("Current password",old),("New password",new),("Repeat new password",new2)]):
        ttk.Label(frm,text=lbl).grid(row=i,column=0,sticky="w",pady=8)
        ttk.Entry(frm,textvariable=var,show="*",width=28).grid(row=i,column=1,sticky="w",pady=8)
    def save():
        salt=_db_get_setting("password_salt",""); digest=_db_get_setting("password_hash","")
        _, entered=_password_hash(old.get(), salt)
        if entered != digest:
            messagebox.showerror("Wrong Password","Current password is not correct."); return
        if len(new.get()) < 4 or new.get()!=new2.get():
            messagebox.showwarning("Check Password","New password must match and be at least 4 characters."); return
        salt2,digest2=_password_hash(new.get())
        _db_set_setting("password_salt", salt2); _db_set_setting("password_hash", digest2)
        dlg.destroy(); messagebox.showinfo("Saved","Password changed.")
    ttk.Button(frm,text="Change Password",command=save).grid(row=4,column=1,sticky="e",pady=14)


def _login_then_start(root):
    _db_init_schema()
    result={"ok":False}
    dlg=tk.Toplevel(root); dlg.title("Sunlectric Incentive Login"); dlg.geometry("380x190"); dlg.resizable(False,False); dlg.grab_set()
    frm=ttk.Frame(dlg,padding=18); frm.pack(fill="both",expand=True)
    ttk.Label(frm,text="Sunlectric Incentive Tool",font=("Segoe UI",12,"bold")).pack(anchor="center",pady=(0,12))
    pw=tk.StringVar()
    ttk.Label(frm,text="Password").pack(anchor="w")
    ent=ttk.Entry(frm,textvariable=pw,show="*"); ent.pack(fill="x",pady=(4,12)); ent.focus_set()
    ttk.Label(frm,text="Default first password: admin123 (change after login)",foreground="gray").pack(anchor="w")
    def attempt(*_):
        salt=_db_get_setting("password_salt",""); digest=_db_get_setting("password_hash","")
        _, entered=_password_hash(pw.get(), salt)
        if entered==digest:
            result["ok"]=True; dlg.destroy()
        else:
            messagebox.showerror("Login Failed","Wrong password.")
    ttk.Button(frm,text="Login",command=attempt).pack(anchor="e",pady=10)
    ent.bind("<Return>", attempt)
    root.wait_window(dlg)
    return result["ok"]


# Patch class with DB-backed behavior while preserving original workflow.
OdooSalesMarginApp.__init__ = _db_init_app
OdooSalesMarginApp.load_and_prepare_odoo = _db_load_and_prepare_odoo
OdooSalesMarginApp.load_charges_file = _db_load_charges_file
OdooSalesMarginApp.load_salaries = _db_load_salaries
OdooSalesMarginApp.load_commissions_db = _db_load_commissions
OdooSalesMarginApp.save_commissions_db = _db_save_commissions
OdooSalesMarginApp.load_creditnotes_db = _db_load_creditnotes
OdooSalesMarginApp.save_creditnotes_db = _db_save_creditnotes
OdooSalesMarginApp.compute_incentives = _db_compute_incentives
OdooSalesMarginApp.open_scheme_settings = open_scheme_settings
OdooSalesMarginApp.open_manual_adjustments = open_manual_adjustments
OdooSalesMarginApp.open_payment_recorder = open_payment_recorder
OdooSalesMarginApp.open_change_password = open_change_password


# =========================
# DB-ONLY STORAGE PATCH (v2)
# This replaces the earlier Excel-mirror behavior. Excel is now import/export only.
# On startup the working views are rebuilt from SQLite, so the app no longer opens blank.
# =========================

def _db_has_ledger_rows():
    with _db_conn() as con:
        return con.execute("SELECT 1 FROM ledger_rows LIMIT 1").fetchone() is not None


def _db_load_all_charges_map():
    with _db_conn() as con:
        rows = con.execute("SELECT invoice_no, transport, loading FROM charges").fetchall()
    return {safe_str(r[0]): {"transport": float(r[1] or 0), "loading": float(r[2] or 0)} for r in rows if safe_str(r[0])}


def _db_ledger_to_df():
    with _db_conn() as con:
        rows = con.execute("""
            SELECT date, number, account, partner, label, debit, credit, salesperson
            FROM ledger_rows
            ORDER BY date, number, row_id
        """).fetchall()
    cols = ["Date", "Number", "Account", "Partner", "Label", "Debit", "Credit", "Sales Order Lines/Salesperson"]
    df = pd.DataFrame(rows, columns=cols)
    if df.empty:
        return df
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Debit"] = pd.to_numeric(df["Debit"], errors="coerce").fillna(0.0)
    df["Credit"] = pd.to_numeric(df["Credit"], errors="coerce").fillna(0.0)
    return df


def _db_prepare_working_data_from_dataframe(self, df):
    """Same calculation workflow as the original loader, but source is SQLite DataFrame, not Excel."""
    if df is None or df.empty:
        self.df_raw = pd.DataFrame(columns=REQ_COLS + ["Salesperson"])
        self.invoice_summary = pd.DataFrame()
        self.bill_lines = pd.DataFrame()
        return

    missing = [c for c in REQ_COLS if c not in df.columns]
    if missing:
        raise ValueError("Database rows are missing required fields:\n- " + "\n- ".join(missing))

    df = df.copy()
    df["Number"] = df["Number"].astype(str).str.strip()
    df["Account"] = df["Account"].astype(str).str.strip()
    df["Partner"] = df["Partner"].astype(str).str.strip()
    df["Label"] = df["Label"].astype(str).fillna("").str.strip()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Debit"] = pd.to_numeric(df["Debit"], errors="coerce").fillna(0.0)
    df["Credit"] = pd.to_numeric(df["Credit"], errors="coerce").fillna(0.0)
    df["Salesperson"] = df["Sales Order Lines/Salesperson"].apply(safe_str)
    self.df_raw = df

    inv = df[df["Number"].apply(is_invoice_number_value)].copy()
    bill = df[df["Number"].str.startswith(BILL_PREFIX, na=False)].copy()

    if inv.empty:
        self.invoice_summary = pd.DataFrame(columns=["Date","Number","Partner","Salesperson","Sales","COGS","Gross Margin","YearMonth","Customer","Transport","Loading","Adjusted GM"])
    else:
        inv_sp = (
            inv.groupby("Number")["Salesperson"]
            .apply(lambda s: next((x for x in s if safe_str(x)), ""))
            .rename("Salesperson")
        )
        inv["acct_l"] = inv["Account"].str.lower()
        is_sales_kw = inv["acct_l"].apply(lambda x: any(k in x for k in SALES_KEYWORDS))
        is_cogs_kw = inv["acct_l"].apply(lambda x: any(k in x for k in COGS_KEYWORDS))
        sales_by_inv = inv[is_sales_kw].groupby("Number")["Credit"].sum()
        cogs_by_inv = inv[is_cogs_kw].groupby("Number")["Debit"].sum()
        credit_fallback = inv.groupby("Number")["Credit"].sum()
        debit_fallback = inv.groupby("Number")["Debit"].sum()
        inv_date = inv.groupby("Number")["Date"].min()
        inv_partner = inv.groupby("Number")["Partner"].apply(lambda s: next(iter(s), ""))

        summary = pd.DataFrame({"Date": inv_date, "Number": inv_date.index, "Partner": inv_partner}).set_index("Number", drop=False)
        summary = summary.join(inv_sp, how="left")
        summary["Sales"] = sales_by_inv.reindex(summary.index).fillna(0.0)
        summary["COGS"] = cogs_by_inv.reindex(summary.index).fillna(0.0)
        need_sales_fb = summary["Sales"].abs() < 1e-9
        summary.loc[need_sales_fb, "Sales"] = credit_fallback.reindex(summary.index).fillna(0.0)[need_sales_fb]
        need_cogs_fb = summary["COGS"].abs() < 1e-9
        summary.loc[need_cogs_fb, "COGS"] = debit_fallback.reindex(summary.index).fillna(0.0)[need_cogs_fb]
        summary["Gross Margin"] = summary["Sales"] - summary["COGS"]
        summary["YearMonth"] = summary["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        summary["Customer"] = summary["Partner"].apply(safe_str)
        summary["Transport"] = 0.0
        summary["Loading"] = 0.0
        summary["Adjusted GM"] = summary["Gross Margin"]
        self.invoice_summary = summary.reset_index(drop=True)

    if bill.empty:
        self.bill_lines = pd.DataFrame(columns=list(df.columns) + ["ExpensePerson","ExpenseAmount","YearMonth"])
    else:
        bill["ExpensePerson"] = bill["Partner"].apply(safe_str)
        bill["ExpenseAmount"] = (bill["Debit"] - bill["Credit"]).astype(float)
        bill["YearMonth"] = bill["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        self.bill_lines = bill

    self.charges_map = _db_load_all_charges_map()
    if self.charges_map and self.invoice_summary is not None and not self.invoice_summary.empty:
        self.apply_charges_to_invoices()
    self.rebuild_global_customer_maps()
    self._update_sp_listbox(self.allowed_salespeople)


def _db_rebuild_working_data_from_database(self):
    df = _db_ledger_to_df()
    if df.empty:
        return False
    _db_prepare_working_data_from_dataframe(self, df)
    try:
        self.populate_months_quarters()
        self.on_filter_change()
        self.refresh_view()
        self.populate_customer_dropdown()
        self.status_var.set(f"Loaded data from local database: {LOCAL_DB_FILE}")
    except Exception:
        pass
    return True


def _db_load_and_prepare_odoo_v2(self, file_path):
    inserted, total = _db_import_ledger_excel(file_path)
    _db_rebuild_working_data_from_database(self)
    try:
        self.status_var.set(f"Odoo import complete: {inserted} fresh rows added, {total - inserted} duplicates ignored. All working data is now from SQLite DB only.")
        messagebox.showinfo("Import Complete", f"Fresh rows added: {inserted}\nDuplicates ignored: {total - inserted}\n\nStorage used: SQLite database only. Excel was used only as the import source.")
    except Exception:
        pass


def _db_load_salaries_v2(self):
    _db_init_schema()
    with _db_conn() as con:
        rows = con.execute("SELECT salesperson,effective_from,monthly_salary FROM salaries ORDER BY salesperson,effective_from").fetchall()
    if rows:
        steps = {}
        for sp, eff_s, sal in rows:
            try:
                eff = pd.to_datetime(eff_s).date()
            except Exception:
                eff = date(1900, 1, 1)
            steps.setdefault(safe_str(sp), []).append((eff, float(sal or 0)))
        self.salary_steps = steps
        self.allowed_salespeople = sorted([x for x in steps.keys() if x])
        self._update_sp_listbox(self.allowed_salespeople)
    else:
        self.salary_steps = {}
        self.allowed_salespeople = []
        try:
            self._update_sp_listbox([])
        except Exception:
            pass
        try:
            messagebox.showwarning("Salaries Required", "No salaries are saved in the local database. Use 'Edit Salaries' to add salary effective dates and amounts.")
        except Exception:
            pass


def _db_save_commissions_v2(self):
    _df_to_db_table(self.commission_db, "commissions")


def _db_load_commissions_v2(self):
    self.commission_db = _db_table_to_df("commissions")


def _db_save_creditnotes_v2(self):
    _df_to_db_table(self.creditnote_db, "credit_notes")


def _db_load_creditnotes_v2(self):
    self.creditnote_db = _db_table_to_df("credit_notes")


def _db_open_salary_editor_v2(self):
    self.load_salaries()
    rows = []
    with _db_conn() as con:
        rows = con.execute("SELECT salesperson,effective_from,monthly_salary FROM salaries ORDER BY salesperson,effective_from").fetchall()
    df = pd.DataFrame(rows, columns=["Salesperson", "EffectiveFrom", "MonthlySalary"])

    dlg = tk.Toplevel(self.root)
    dlg.title("Edit Salaries (Database)")
    dlg.geometry("860x560")
    dlg.transient(self.root)
    dlg.grab_set()
    ttk.Label(dlg, text=f"Salaries are stored in local SQLite database only:\n{LOCAL_DB_FILE}", padding=10).pack(anchor="w")
    frame = ttk.Frame(dlg, padding=10); frame.pack(fill="both", expand=True)
    cols = ["Salesperson", "EffectiveFrom", "MonthlySalary"]
    tree = ttk.Treeview(frame, columns=cols, show="headings", height=16)
    for c in cols:
        tree.heading(c, text=c)
        tree.column(c, width=280 if c == "Salesperson" else 200, anchor="w" if c != "MonthlySalary" else "e")
    tree.pack(side="left", fill="both", expand=True)
    vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview); tree.configure(yscrollcommand=vsb.set); vsb.pack(side="right", fill="y")
    for _, r in df.fillna("").iterrows():
        tree.insert("", "end", values=(safe_str(r["Salesperson"]), safe_str(r["EffectiveFrom"]), safe_str(r["MonthlySalary"])))

    controls = ttk.Frame(dlg, padding=10); controls.pack(fill="x")
    sp_var = tk.StringVar(); eff_var = tk.StringVar(value=SCHEME_START_DATE.isoformat()); sal_var = tk.StringVar()
    ttk.Label(controls, text="Salesperson").grid(row=0, column=0, sticky="w")
    ttk.Entry(controls, textvariable=sp_var, width=22).grid(row=0, column=1, padx=6)
    ttk.Label(controls, text="EffectiveFrom (YYYY-MM-DD)").grid(row=0, column=2, sticky="w")
    ttk.Entry(controls, textvariable=eff_var, width=18).grid(row=0, column=3, padx=6)
    ttk.Label(controls, text="MonthlySalary").grid(row=0, column=4, sticky="w")
    ttk.Entry(controls, textvariable=sal_var, width=14).grid(row=0, column=5, padx=6)

    def add_row():
        spn = sp_var.get().strip(); eff = eff_var.get().strip(); sal = sal_var.get().strip()
        if not spn: return
        d = parse_date_entry(eff)
        if not d:
            messagebox.showerror("Invalid", "EffectiveFrom must be YYYY-MM-DD"); return
        sval = to_num(sal)
        if sval <= 0:
            messagebox.showerror("Invalid", "MonthlySalary must be a positive number"); return
        tree.insert("", "end", values=(spn, d.isoformat(), f"{sval:.2f}")); sp_var.set(""); sal_var.set("")

    def delete_selected():
        for iid in tree.selection(): tree.delete(iid)

    def save_db():
        with _db_conn() as con:
            con.execute("DELETE FROM salaries")
            for iid in tree.get_children():
                v = tree.item(iid, "values")
                spn = safe_str(v[0]); eff = safe_str(v[1]); sal = to_num(v[2])
                d = parse_date_entry(eff)
                if spn and d and sal > 0:
                    con.execute("INSERT OR REPLACE INTO salaries(salesperson,effective_from,monthly_salary) VALUES(?,?,?)", (spn, d.isoformat(), sal))
        self.load_salaries(); dlg.destroy(); self.refresh_view()
        messagebox.showinfo("Saved", "Salaries saved in local database only.")

    ttk.Button(controls, text="Add", command=add_row).grid(row=1, column=1, pady=10, sticky="w")
    ttk.Button(controls, text="Delete Selected", command=delete_selected).grid(row=1, column=2, pady=10, sticky="w")
    ttk.Button(controls, text="Save to Database", command=save_db).grid(row=1, column=5, pady=10, sticky="e")


def _db_open_salary_file_external_v2(self):
    messagebox.showinfo("Database Only", f"Salary master is now stored only in SQLite database:\n{LOCAL_DB_FILE}\n\nUse 'Edit Salaries' inside the app to add or change salary records.")


def _db_init_app_v2(self, root):
    _db_init_schema()
    _db_load_scheme_globals()
    self.local_db_path = LOCAL_DB_FILE
    _DB_ORIGINAL_INIT(self, root)
    # Re-apply DB-only loaders because original init calls class methods during construction.
    self.load_salaries()
    self.load_commissions_db()
    self.load_creditnotes_db()
    loaded = _db_rebuild_working_data_from_database(self)
    try:
        if loaded:
            self.status_var.set(f"Ready. Existing records loaded from SQLite DB: {LOCAL_DB_FILE}")
        else:
            self.status_var.set(f"Ready. SQLite DB is active. Import Odoo Excel to add fresh rows: {LOCAL_DB_FILE}")
    except Exception:
        pass


# Override earlier compatibility patch. From here onward Excel is import/export only, never storage.
OdooSalesMarginApp.__init__ = _db_init_app_v2
OdooSalesMarginApp.load_and_prepare_odoo = _db_load_and_prepare_odoo_v2
OdooSalesMarginApp.load_salaries = _db_load_salaries_v2
OdooSalesMarginApp.load_commissions_db = _db_load_commissions_v2
OdooSalesMarginApp.save_commissions_db = _db_save_commissions_v2
OdooSalesMarginApp.load_creditnotes_db = _db_load_creditnotes_v2
OdooSalesMarginApp.save_creditnotes_db = _db_save_creditnotes_v2
OdooSalesMarginApp.open_salary_editor = _db_open_salary_editor_v2
OdooSalesMarginApp.open_salary_file_external = _db_open_salary_file_external_v2



# =========================
# FINAL PATCH v4: DB-only raw data view + robust invoice/person matching
# =========================
_FINAL_ORIGINAL_BUILD_UI_V4 = OdooSalesMarginApp._build_ui
_FINAL_ORIGINAL_REFRESH_VIEW_V4 = OdooSalesMarginApp.refresh_view
_FINAL_ORIGINAL_LOAD_CHARGES_V4 = OdooSalesMarginApp.load_charges_file

def _norm_person_name_v4(x):
    return re.sub(r"\s+", " ", safe_str(x)).strip().casefold()

def _map_salesperson_to_master_v4(self, sp):
    s = safe_str(sp)
    if not s:
        return ""
    lookup = {_norm_person_name_v4(x): x for x in getattr(self, "allowed_salespeople", [])}
    return lookup.get(_norm_person_name_v4(s), s)

def is_invoice_number_value_v4(x) -> bool:
    s = safe_str(x).upper()
    if not s:
        return False
    if s.startswith(CREDIT_NOTE_PREFIX.upper()) or s.startswith(BILL_PREFIX.upper()):
        return False
    return (s.startswith("SL/") or s.startswith("SLFY") or s.startswith("SINV") or s.startswith("INV/") or s.startswith("INV-"))

is_invoice_number_value = is_invoice_number_value_v4

def _db_prepare_working_data_from_dataframe_v4(self, df):
    if df is None or df.empty:
        self.df_raw = pd.DataFrame(columns=REQ_COLS + ["Salesperson"])
        self.invoice_summary = pd.DataFrame(columns=["Date","Number","Partner","Salesperson","Sales","COGS","Gross Margin","YearMonth","Customer","Transport","Loading","Adjusted GM"])
        self.bill_lines = pd.DataFrame(columns=REQ_COLS + ["Salesperson", "ExpensePerson", "ExpenseAmount", "YearMonth"])
        return
    missing = [c for c in REQ_COLS if c not in df.columns]
    if missing:
        raise ValueError("Database rows are missing required fields:\n- " + "\n- ".join(missing))
    df = df.copy()
    df["Number"] = df["Number"].astype(str).str.strip()
    df["Account"] = df["Account"].astype(str).str.strip()
    df["Partner"] = df["Partner"].astype(str).str.strip()
    df["Label"] = df["Label"].astype(str).fillna("").str.strip()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["Debit"] = pd.to_numeric(df["Debit"], errors="coerce").fillna(0.0)
    df["Credit"] = pd.to_numeric(df["Credit"], errors="coerce").fillna(0.0)
    df["Salesperson"] = df["Sales Order Lines/Salesperson"].apply(lambda x: _map_salesperson_to_master_v4(self, x))
    df["Sales Order Lines/Salesperson"] = df["Salesperson"]
    self.df_raw = df
    inv = df[df["Number"].apply(is_invoice_number_value_v4)].copy()
    bill = df[df["Number"].str.upper().str.startswith(BILL_PREFIX.upper(), na=False)].copy()
    if inv.empty:
        self.invoice_summary = pd.DataFrame(columns=["Date","Number","Partner","Salesperson","Sales","COGS","Gross Margin","YearMonth","Customer","Transport","Loading","Adjusted GM"])
    else:
        inv_sp = inv.groupby("Number")["Salesperson"].apply(lambda s: next((safe_str(x) for x in s if safe_str(x)), "")).rename("Salesperson")
        inv["acct_l"] = inv["Account"].str.lower()
        is_sales_kw = inv["acct_l"].apply(lambda x: any(k in x for k in SALES_KEYWORDS))
        is_cogs_kw = inv["acct_l"].apply(lambda x: any(k in x for k in COGS_KEYWORDS))
        sales_by_inv = inv[is_sales_kw].groupby("Number")["Credit"].sum()
        cogs_by_inv = inv[is_cogs_kw].groupby("Number")["Debit"].sum()
        credit_fallback = inv.groupby("Number")["Credit"].sum()
        debit_fallback = inv.groupby("Number")["Debit"].sum()
        inv_date = inv.groupby("Number")["Date"].min()
        inv_partner = inv.groupby("Number")["Partner"].apply(lambda s: next(iter(s), ""))
        summary = pd.DataFrame({"Date": inv_date, "Number": inv_date.index, "Partner": inv_partner}).set_index("Number", drop=False)
        summary = summary.join(inv_sp, how="left")
        summary["Sales"] = sales_by_inv.reindex(summary.index).fillna(0.0)
        summary["COGS"] = cogs_by_inv.reindex(summary.index).fillna(0.0)
        need_sales_fb = summary["Sales"].abs() < 1e-9
        summary.loc[need_sales_fb, "Sales"] = credit_fallback.reindex(summary.index).fillna(0.0)[need_sales_fb]
        need_cogs_fb = summary["COGS"].abs() < 1e-9
        summary.loc[need_cogs_fb, "COGS"] = debit_fallback.reindex(summary.index).fillna(0.0)[need_cogs_fb]
        summary["Gross Margin"] = summary["Sales"] - summary["COGS"]
        summary["YearMonth"] = summary["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        summary["Customer"] = summary["Partner"].apply(safe_str)
        summary["Transport"] = 0.0
        summary["Loading"] = 0.0
        summary["Adjusted GM"] = summary["Gross Margin"]
        self.invoice_summary = summary.reset_index(drop=True)
    if bill.empty:
        self.bill_lines = pd.DataFrame(columns=list(df.columns) + ["ExpensePerson","ExpenseAmount","YearMonth"])
    else:
        bill["ExpensePerson"] = bill["Partner"].apply(lambda x: _map_salesperson_to_master_v4(self, x))
        bill["ExpenseAmount"] = (bill["Debit"] - bill["Credit"]).astype(float)
        bill["YearMonth"] = bill["Date"].apply(lambda d: "" if pd.isna(d) else f"{d.year:04d}-{d.month:02d}")
        self.bill_lines = bill
    self.charges_map = _db_load_all_charges_map()
    if self.charges_map and self.invoice_summary is not None and not self.invoice_summary.empty:
        self.apply_charges_to_invoices()
    self.rebuild_global_customer_maps()
    try:
        self._update_sp_listbox(self.allowed_salespeople)
    except Exception:
        pass

def _db_rebuild_working_data_from_database_v4(self):
    df = _db_ledger_to_df()
    if df.empty:
        return False
    _db_prepare_working_data_from_dataframe_v4(self, df)
    try:
        self.populate_months_quarters()
        self.on_filter_change()
        self.refresh_view()
        self.populate_customer_dropdown()
        self.refresh_raw_data_view()
        self.status_var.set(f"Loaded data from local SQLite database: {LOCAL_DB_FILE}")
    except Exception:
        pass
    return True

def _db_load_and_prepare_odoo_v4(self, file_path):
    inserted, total = _db_import_ledger_excel(file_path)
    _db_rebuild_working_data_from_database_v4(self)
    try:
        self.status_var.set(f"Odoo import complete: {inserted} fresh rows added, {total - inserted} duplicates ignored. Views are rebuilt from SQLite DB only.")
        messagebox.showinfo("Import Complete", f"Fresh rows added: {inserted}\nDuplicates ignored: {total - inserted}\n\nInvoices, credit notes, bills and raw data are now loaded from the SQLite database only.")
    except Exception:
        pass

def _db_load_charges_file_v4(self, file_path):
    _FINAL_ORIGINAL_LOAD_CHARGES_V4(self, file_path)
    try:
        self.refresh_raw_data_view()
    except Exception:
        pass

def _add_raw_data_tab_v4(self):
    self.tab_raw = ttk.Frame(self.nb)
    self.nb.add(self.tab_raw, text="Raw Data / Charges Verify")
    filters = ttk.LabelFrame(self.tab_raw, text="Raw DB Filters", padding=8)
    filters.pack(fill="x", padx=10, pady=(8, 4))
    self.raw_name_filter_var = tk.StringVar(value="")
    self.raw_emp_filter_var = tk.StringVar(value="All")
    self.raw_from_filter_var = tk.StringVar(value="")
    self.raw_to_filter_var = tk.StringVar(value="")
    self.raw_type_filter_var = tk.StringVar(value="All")
    ttk.Label(filters, text="Customer/Name/Number contains:").grid(row=0, column=0, sticky="w", padx=(0, 4), pady=3)
    ttk.Entry(filters, textvariable=self.raw_name_filter_var, width=34).grid(row=0, column=1, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="Employee:").grid(row=0, column=2, sticky="w", padx=(0, 4), pady=3)
    self.raw_emp_combo = ttk.Combobox(filters, textvariable=self.raw_emp_filter_var, width=28, state="readonly")
    self.raw_emp_combo.grid(row=0, column=3, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="Type:").grid(row=0, column=4, sticky="w", padx=(0, 4), pady=3)
    ttk.Combobox(filters, textvariable=self.raw_type_filter_var, values=["All", "Invoice", "Credit Note", "Bill", "Charge", "Other"], width=14, state="readonly").grid(row=0, column=5, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="From YYYY-MM-DD:").grid(row=1, column=0, sticky="w", padx=(0, 4), pady=3)
    ttk.Entry(filters, textvariable=self.raw_from_filter_var, width=16).grid(row=1, column=1, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="To YYYY-MM-DD:").grid(row=1, column=2, sticky="w", padx=(0, 4), pady=3)
    ttk.Entry(filters, textvariable=self.raw_to_filter_var, width=16).grid(row=1, column=3, sticky="w", padx=(0, 12), pady=3)
    ttk.Button(filters, text="Apply Raw Filters", command=self.refresh_raw_data_view).grid(row=1, column=4, sticky="w", padx=(0, 8), pady=3)
    ttk.Button(filters, text="Clear", command=self.clear_raw_data_filters).grid(row=1, column=5, sticky="w", pady=3)
    self.raw_count_var = tk.StringVar(value="")
    ttk.Label(self.tab_raw, textvariable=self.raw_count_var, padding=(10, 2, 10, 2), foreground="blue").pack(anchor="w")
    frame = ttk.Frame(self.tab_raw, padding=10)
    frame.pack(fill="both", expand=True)
    cols = ["Type", "Date", "Number/Invoice", "Employee", "Partner/Name", "Account", "Label", "Debit", "Credit", "Transport", "Loading", "Source"]
    self.tree_raw = ttk.Treeview(frame, columns=cols, show="headings", height=22)
    for c in cols:
        self.tree_raw.heading(c, text=c)
        if c in ("Partner/Name", "Account", "Label", "Source"):
            self.tree_raw.column(c, width=230, anchor="w")
        elif c in ("Debit", "Credit", "Transport", "Loading"):
            self.tree_raw.column(c, width=110, anchor="e")
        else:
            self.tree_raw.column(c, width=130, anchor="w")
    vsb = ttk.Scrollbar(frame, orient="vertical", command=self.tree_raw.yview)
    hsb = ttk.Scrollbar(frame, orient="horizontal", command=self.tree_raw.xview)
    self.tree_raw.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
    self.tree_raw.pack(side="left", fill="both", expand=True)
    vsb.pack(side="right", fill="y")
    hsb.pack(side="bottom", fill="x")

def _build_ui_v4(self):
    _FINAL_ORIGINAL_BUILD_UI_V4(self)
    try:
        _add_raw_data_tab_v4(self)
    except Exception as e:
        try:
            messagebox.showwarning("Raw Data Tab", f"Could not create Raw Data tab:\n{e}")
        except Exception:
            pass

def clear_raw_data_filters_v4(self):
    self.raw_name_filter_var.set("")
    self.raw_emp_filter_var.set("All")
    self.raw_from_filter_var.set("")
    self.raw_to_filter_var.set("")
    self.raw_type_filter_var.set("All")
    self.refresh_raw_data_view()

def _raw_data_rows_v4(self):
    rows = []
    if self.df_raw is not None and not self.df_raw.empty:
        for _, r in self.df_raw.copy().iterrows():
            number = safe_str(r.get("Number", ""))
            up = number.upper()
            typ = "Invoice" if is_invoice_number_value_v4(number) else ("Credit Note" if up.startswith(CREDIT_NOTE_PREFIX.upper()) else ("Bill" if up.startswith(BILL_PREFIX.upper()) else "Other"))
            rows.append({"Type": typ, "Date": "" if pd.isna(r.get("Date")) else pd.to_datetime(r.get("Date")).strftime("%Y-%m-%d"), "Number/Invoice": number, "Employee": safe_str(r.get("Salesperson", r.get("Sales Order Lines/Salesperson", ""))), "Partner/Name": safe_str(r.get("Partner", "")), "Account": safe_str(r.get("Account", "")), "Label": safe_str(r.get("Label", "")), "Debit": float(to_num(r.get("Debit", 0))), "Credit": float(to_num(r.get("Credit", 0))), "Transport": "", "Loading": "", "Source": "ledger_rows DB"})
    with _db_conn() as con:
        ch = con.execute("SELECT invoice_no, transport, loading, source_file, updated_at FROM charges ORDER BY invoice_no").fetchall()
    for invoice_no, transport, loading, source_file, updated_at in ch:
        sp = ""; partner = ""; dt = ""
        if self.invoice_summary is not None and not self.invoice_summary.empty:
            hit = self.invoice_summary[self.invoice_summary["Number"].astype(str).str.strip() == safe_str(invoice_no)]
            if not hit.empty:
                rr = hit.iloc[0]
                sp = safe_str(rr.get("Salesperson", ""))
                partner = safe_str(rr.get("Partner", ""))
                dt = "" if pd.isna(rr.get("Date")) else pd.to_datetime(rr.get("Date")).strftime("%Y-%m-%d")
        rows.append({"Type": "Charge", "Date": dt, "Number/Invoice": safe_str(invoice_no), "Employee": sp, "Partner/Name": partner, "Account": "Loading/Transportation Charges", "Label": "Stored in charges DB", "Debit": "", "Credit": "", "Transport": float(transport or 0), "Loading": float(loading or 0), "Source": safe_str(source_file) or "charges DB"})
    return rows

def refresh_raw_data_view_v4(self):
    if not hasattr(self, "tree_raw"):
        return
    for iid in self.tree_raw.get_children():
        self.tree_raw.delete(iid)
    try:
        employees = ["All"] + sorted(set([safe_str(x) for x in getattr(self, "allowed_salespeople", []) if safe_str(x)]))
        self.raw_emp_combo["values"] = employees
        if self.raw_emp_filter_var.get() not in employees:
            self.raw_emp_filter_var.set("All")
    except Exception:
        pass
    q = self.raw_name_filter_var.get().strip().casefold()
    emp = self.raw_emp_filter_var.get().strip()
    typ = self.raw_type_filter_var.get().strip() or "All"
    d1 = parse_date_entry(self.raw_from_filter_var.get()) if self.raw_from_filter_var.get().strip() else None
    d2 = parse_date_entry(self.raw_to_filter_var.get()) if self.raw_to_filter_var.get().strip() else None
    shown = 0; total_debit = total_credit = total_transport = total_loading = 0.0
    for r in _raw_data_rows_v4(self):
        if typ != "All" and r["Type"] != typ:
            continue
        if emp and emp != "All" and _norm_person_name_v4(r.get("Employee", "")) != _norm_person_name_v4(emp):
            continue
        if q:
            hay = " ".join([safe_str(r.get(k, "")) for k in ["Number/Invoice", "Employee", "Partner/Name", "Account", "Label"]]).casefold()
            if q not in hay:
                continue
        rd = parse_date_entry(r.get("Date", ""))
        if d1 and (not rd or rd < d1):
            continue
        if d2 and (not rd or rd > d2):
            continue
        values = []
        for c in ["Type", "Date", "Number/Invoice", "Employee", "Partner/Name", "Account", "Label", "Debit", "Credit", "Transport", "Loading", "Source"]:
            v = r.get(c, "")
            values.append(fmt_money(v) if c in ("Debit", "Credit", "Transport", "Loading") and v != "" else v)
        self.tree_raw.insert("", "end", values=values)
        shown += 1
        total_debit += to_num(r.get("Debit", 0)); total_credit += to_num(r.get("Credit", 0)); total_transport += to_num(r.get("Transport", 0)); total_loading += to_num(r.get("Loading", 0))
    self.raw_count_var.set(f"Showing {shown} DB rows | Debit {fmt_money(total_debit)} | Credit {fmt_money(total_credit)} | Transport {fmt_money(total_transport)} | Loading {fmt_money(total_loading)}")

def refresh_view_v4(self):
    _FINAL_ORIGINAL_REFRESH_VIEW_V4(self)
    try:
        self.refresh_raw_data_view()
    except Exception:
        pass

def _db_init_app_v4(self, root):
    _db_init_schema()
    _db_load_scheme_globals()
    self.local_db_path = LOCAL_DB_FILE
    _DB_ORIGINAL_INIT(self, root)
    self.load_salaries()
    self.load_commissions_db()
    self.load_creditnotes_db()
    loaded = _db_rebuild_working_data_from_database_v4(self)
    try:
        self.status_var.set(f"Ready. Existing records loaded from SQLite DB: {LOCAL_DB_FILE}" if loaded else f"Ready. SQLite DB is active. Import Odoo Excel to add fresh rows: {LOCAL_DB_FILE}")
        self.refresh_raw_data_view()
    except Exception:
        pass

OdooSalesMarginApp._build_ui = _build_ui_v4
OdooSalesMarginApp.__init__ = _db_init_app_v4
OdooSalesMarginApp.load_and_prepare_odoo = _db_load_and_prepare_odoo_v4
OdooSalesMarginApp.load_charges_file = _db_load_charges_file_v4
OdooSalesMarginApp.refresh_view = refresh_view_v4
OdooSalesMarginApp.refresh_raw_data_view = refresh_raw_data_view_v4
OdooSalesMarginApp.clear_raw_data_filters = clear_raw_data_filters_v4



# =========================
# FINAL PATCH v5: better horizontal scrolling + DB admin tools
# =========================
def _db_noop_file_storage(*args, **kwargs):
    return None
OdooSalesMarginApp.ensure_salary_file_exists = _db_noop_file_storage
OdooSalesMarginApp.ensure_commissions_file_exists = _db_noop_file_storage
OdooSalesMarginApp.ensure_creditnotes_file_exists = _db_noop_file_storage

def _verify_admin_password_dialog(parent, title="Password Required"):
    result = {"ok": False}
    dlg = tk.Toplevel(parent); dlg.title(title); dlg.geometry("390x170"); dlg.resizable(False, False); dlg.transient(parent); dlg.grab_set()
    frm = ttk.Frame(dlg, padding=16); frm.pack(fill="both", expand=True)
    ttk.Label(frm, text="Enter app password to continue", font=("Segoe UI", 10, "bold")).pack(anchor="w", pady=(0, 10))
    pw = tk.StringVar(); ent = ttk.Entry(frm, textvariable=pw, show="*", width=32); ent.pack(fill="x"); ent.focus_set()
    def check(*_):
        salt = _db_get_setting("password_salt", ""); digest = _db_get_setting("password_hash", "")
        _, entered = _password_hash(pw.get(), salt)
        if entered == digest:
            result["ok"] = True; dlg.destroy()
        else:
            messagebox.showerror("Wrong Password", "Password is not correct.", parent=dlg)
    btns = ttk.Frame(frm); btns.pack(fill="x", pady=14)
    ttk.Button(btns, text="Cancel", command=dlg.destroy).pack(side="right")
    ttk.Button(btns, text="Continue", command=check).pack(side="right", padx=(0, 8))
    ent.bind("<Return>", check); parent.wait_window(dlg); return result["ok"]

def _db_data_tables_v5():
    return ["ledger_rows", "charges", "salaries", "commissions", "credit_notes", "manual_adjustments", "incentive_payments"]

def open_wipe_database_v5(self):
    if not _verify_admin_password_dialog(self.root, "Wipe Database - Password Required"): return
    confirm = tk.Toplevel(self.root); confirm.title("Confirm Wipe Clean"); confirm.geometry("520x250"); confirm.transient(self.root); confirm.grab_set()
    frm = ttk.Frame(confirm, padding=16); frm.pack(fill="both", expand=True)
    ttk.Label(frm, text="This will permanently delete all business data from the local database.", font=("Segoe UI", 10, "bold"), foreground="red").pack(anchor="w")
    ttk.Label(frm, text="Deleted: Odoo rows, charges, salaries, commissions, credit note assignments, manual adjustments and payment records.\nKept: app password and incentive settings.", wraplength=480, padding=(0, 10, 0, 10)).pack(anchor="w")
    typed = tk.StringVar(); ttk.Label(frm, text="Type WIPE to confirm:").pack(anchor="w"); ttk.Entry(frm, textvariable=typed, width=18).pack(anchor="w", pady=(4, 10))
    def do_wipe():
        if typed.get().strip() != "WIPE":
            messagebox.showwarning("Confirmation Needed", "Type WIPE exactly to continue.", parent=confirm); return
        with _db_conn() as con:
            for table in _db_data_tables_v5(): con.execute(f"DELETE FROM {table}")
        self.df_raw = pd.DataFrame(columns=REQ_COLS + ["Salesperson"]); self.invoice_summary = pd.DataFrame(); self.bill_lines = pd.DataFrame(); self.charges_map = {}; self.salary_steps = {}; self.allowed_salespeople = []
        self.commission_db = pd.DataFrame(columns=["RowID","Date","YearMonth","Amount","Partner","Label","Number","Account","Salesperson"])
        self.creditnote_db = self.commission_db.copy()
        try:
            self._update_sp_listbox([]); self.populate_months_quarters(); self.refresh_view(); self.populate_customer_dropdown(); self.refresh_raw_data_view(); self.status_var.set(f"Database wiped clean. SQLite DB is ready for fresh import: {LOCAL_DB_FILE}")
        except Exception: pass
        confirm.destroy(); messagebox.showinfo("Done", "All business data has been wiped from the local database.")
    btns = ttk.Frame(frm); btns.pack(fill="x", pady=8)
    ttk.Button(btns, text="Cancel", command=confirm.destroy).pack(side="right")
    ttk.Button(btns, text="Wipe Clean Data", command=do_wipe).pack(side="right", padx=(0, 8))

def export_initial_database_v5(self):
    if not _verify_admin_password_dialog(self.root, "Export Database - Password Required"): return
    _db_init_schema()
    path = filedialog.asksaveasfilename(title="Save Initial Database for Another Computer", defaultextension=".db", filetypes=[("SQLite DB", "*.db"), ("All files", "*.*")], initialfile=f"Sunlectric_Incentives_Initial_{datetime.now().strftime('%Y%m%d_%H%M')}.db")
    if not path: return
    src_con = sqlite3.connect(LOCAL_DB_FILE); dst_con = sqlite3.connect(path)
    try:
        src_con.backup(dst_con)
        salt, digest = _password_hash("admin123")
        dst_con.execute("INSERT OR REPLACE INTO app_settings(key,value) VALUES('password_salt',?)", (salt,)); dst_con.execute("INSERT OR REPLACE INTO app_settings(key,value) VALUES('password_hash',?)", (digest,)); dst_con.commit()
    finally:
        src_con.close(); dst_con.close()
    messagebox.showinfo("Export Complete", f"Initial database exported.\n\nThe exported copy opens with password: admin123\nChange it on the other computer after first login.\n\n{path}")

def import_initial_database_v5(self):
    if not _verify_admin_password_dialog(self.root, "Import Database - Password Required"): return
    path = filedialog.askopenfilename(title="Select Initial Database From Another Computer", filetypes=[("SQLite DB", "*.db"), ("All files", "*.*")])
    if not path: return
    if not messagebox.askyesno("Replace Local Business Data?", "This will replace this computer's business data with the selected database.\n\nYour current login password will be kept. Continue?"): return
    backup = str(APP_DATA_DIR / f"SUNLECTRIC_INCENTIVES_LOCAL_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    if os.path.exists(LOCAL_DB_FILE):
        import shutil; shutil.copy2(LOCAL_DB_FILE, backup)
    _db_init_schema(); src = sqlite3.connect(path)
    try:
        src_tables = {r[0] for r in src.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
        missing = [t for t in _db_data_tables_v5() if t not in src_tables]
        if missing: raise ValueError("Selected file is not a valid Sunlectric incentive database. Missing tables: " + ", ".join(missing))
        with _db_conn() as dst:
            for table in _db_data_tables_v5():
                dst.execute(f"DELETE FROM {table}")
                cols = [r[1] for r in dst.execute(f"PRAGMA table_info({table})").fetchall()]
                src_cols = [r[1] for r in src.execute(f"PRAGMA table_info({table})").fetchall()]
                common = [c for c in cols if c in src_cols]
                if not common: continue
                col_sql = ",".join(common); placeholders = ",".join(["?"] * len(common))
                for row in src.execute(f"SELECT {col_sql} FROM {table}").fetchall():
                    dst.execute(f"INSERT OR REPLACE INTO {table} ({col_sql}) VALUES ({placeholders})", row)
    finally:
        src.close()
    self.load_salaries(); self.load_commissions_db(); self.load_creditnotes_db(); _db_rebuild_working_data_from_database_v4(self)
    try: self.status_var.set(f"Imported initial database. Current login password was preserved. Backup: {backup if os.path.exists(backup) else 'not needed'}")
    except Exception: pass
    messagebox.showinfo("Import Complete", "Initial database imported successfully. Current computer password was kept.")

def _make_tree_v5(self, parent, kind="txn"):
    frame = ttk.Frame(parent, padding=10); frame.pack(fill="both", expand=True); frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)
    if kind == "txn": cols = ["Type", "Date", "Number", "Partner/Info", "Sales", "COGS", "Transport", "Loading", "Adj GM", "GM %", "Emp Exp", "Commission", "Net"]
    elif kind == "inc": cols = ["Month", "Salary Used", "Base", "Threshold", "Carry In", "Met?", "10%", "New", "Repeat", "Total", "Notice"]
    elif kind == "comm": cols = ["Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "RowID"]
    elif kind == "cust": cols = ["Type", "Date", "Number", "Salesperson", "Partner", "Sales", "COGS", "Adj GM", "GM %"]
    else: cols = ["Date", "YearMonth", "Amount", "Partner", "Label", "Number", "Account", "RowID", "Salesperson"]
    tree = ttk.Treeview(frame, columns=cols, show="headings", height=22)
    for c in cols:
        tree.heading(c, text=c)
        if kind == "txn": width = 340 if c == "Partner/Info" else (180 if c == "Number" else (110 if c in ("Type", "Date") else 120)); anchor = "w" if c in ("Type", "Date", "Number", "Partner/Info") else "e"
        elif kind == "inc": width = 170 if c == "Notice" else 145; anchor = "w" if c in ("Month", "Met?", "Notice") else "e"
        elif kind == "cust": width = 360 if c == "Partner" else 140; anchor = "w" if c in ("Type", "Date", "Number", "Salesperson", "Partner") else "e"
        else: width = 280 if c in ("Partner", "Label", "Account") else (170 if c == "RowID" else 130); anchor = "e" if c == "Amount" else "w"
        tree.column(c, width=width, minwidth=80, anchor=anchor, stretch=False)
    vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview); hsb = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set); tree.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns"); hsb.grid(row=1, column=0, sticky="ew")
    return tree

_PREV_BUILD_UI_V5 = OdooSalesMarginApp._build_ui
OdooSalesMarginApp._make_tree = _make_tree_v5

def _add_database_admin_menu_v5(self):
    menubar = tk.Menu(self.root); dbmenu = tk.Menu(menubar, tearoff=0)
    dbmenu.add_command(label="Export Initial Database for Another Computer", command=lambda: export_initial_database_v5(self))
    dbmenu.add_command(label="Import Initial Database on This Computer", command=lambda: import_initial_database_v5(self))
    dbmenu.add_separator(); dbmenu.add_command(label="Wipe Clean Whole Data", command=lambda: open_wipe_database_v5(self))
    menubar.add_cascade(label="Database Admin", menu=dbmenu); self.root.config(menu=menubar)

def _build_ui_v5(self):
    _PREV_BUILD_UI_V5(self); _add_database_admin_menu_v5(self)
OdooSalesMarginApp._build_ui = _build_ui_v5

# Raw data tab: same functionality as v4, but uses grid so horizontal scroll is visible.
def _add_raw_data_tab_v5(self):
    self.tab_raw = ttk.Frame(self.nb); self.nb.add(self.tab_raw, text="Raw Data / Charges Verify")
    filters = ttk.LabelFrame(self.tab_raw, text="Raw DB Filters", padding=8); filters.pack(fill="x", padx=10, pady=(8, 4))
    self.raw_name_filter_var = tk.StringVar(value=""); self.raw_emp_filter_var = tk.StringVar(value="All"); self.raw_from_filter_var = tk.StringVar(value=""); self.raw_to_filter_var = tk.StringVar(value=""); self.raw_type_filter_var = tk.StringVar(value="All")
    ttk.Label(filters, text="Customer/Name/Number contains:").grid(row=0, column=0, sticky="w", padx=(0, 4), pady=3); ttk.Entry(filters, textvariable=self.raw_name_filter_var, width=34).grid(row=0, column=1, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="Employee:").grid(row=0, column=2, sticky="w", padx=(0, 4), pady=3); self.raw_emp_combo = ttk.Combobox(filters, textvariable=self.raw_emp_filter_var, width=28, state="readonly"); self.raw_emp_combo.grid(row=0, column=3, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="Type:").grid(row=0, column=4, sticky="w", padx=(0, 4), pady=3); ttk.Combobox(filters, textvariable=self.raw_type_filter_var, values=["All", "Invoice", "Credit Note", "Bill", "Charge", "Other"], width=14, state="readonly").grid(row=0, column=5, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="From YYYY-MM-DD:").grid(row=1, column=0, sticky="w", padx=(0, 4), pady=3); ttk.Entry(filters, textvariable=self.raw_from_filter_var, width=16).grid(row=1, column=1, sticky="w", padx=(0, 12), pady=3)
    ttk.Label(filters, text="To YYYY-MM-DD:").grid(row=1, column=2, sticky="w", padx=(0, 4), pady=3); ttk.Entry(filters, textvariable=self.raw_to_filter_var, width=16).grid(row=1, column=3, sticky="w", padx=(0, 12), pady=3)
    ttk.Button(filters, text="Apply Raw Filters", command=self.refresh_raw_data_view).grid(row=1, column=4, sticky="w", padx=(0, 8), pady=3); ttk.Button(filters, text="Clear", command=self.clear_raw_data_filters).grid(row=1, column=5, sticky="w", pady=3)
    self.raw_count_var = tk.StringVar(value=""); ttk.Label(self.tab_raw, textvariable=self.raw_count_var, padding=(10, 2, 10, 2), foreground="blue").pack(anchor="w")
    frame = ttk.Frame(self.tab_raw, padding=10); frame.pack(fill="both", expand=True); frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)
    cols = ["Type", "Date", "Number/Invoice", "Employee", "Partner/Name", "Account", "Label", "Debit", "Credit", "Transport", "Loading", "Source"]
    self.tree_raw = ttk.Treeview(frame, columns=cols, show="headings", height=22)
    for c in cols:
        self.tree_raw.heading(c, text=c); width = 240 if c in ("Partner/Name", "Account", "Label", "Source") else (120 if c in ("Debit", "Credit", "Transport", "Loading") else 140); anchor = "e" if c in ("Debit", "Credit", "Transport", "Loading") else "w"; self.tree_raw.column(c, width=width, minwidth=90, anchor=anchor, stretch=False)
    vsb = ttk.Scrollbar(frame, orient="vertical", command=self.tree_raw.yview); hsb = ttk.Scrollbar(frame, orient="horizontal", command=self.tree_raw.xview)
    self.tree_raw.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set); self.tree_raw.grid(row=0, column=0, sticky="nsew"); vsb.grid(row=0, column=1, sticky="ns"); hsb.grid(row=1, column=0, sticky="ew")
_add_raw_data_tab_v4 = _add_raw_data_tab_v5


# =========================
# App main
# =========================
def main():
    root = tk.Tk()
    root.withdraw()
    try:
        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
    except Exception:
        pass
    if not _login_then_start(root):
        root.destroy()
        return
    root.deiconify()
    OdooSalesMarginApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
